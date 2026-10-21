import io
import json
import os
import sqlite3
import tempfile
import threading
import time
import zipfile
import xml.etree.ElementTree as ET
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import patch

# This module imports the live FastAPI app whose data root is the tool package's
# data/ directory. Run it only from the isolated copied package used by CI/manual QA.
if os.environ.get("POINT_QUOTE_ALLOW_APP_TESTS") != "1":
    import pytest
    pytest.skip("run in an isolated tool copy with POINT_QUOTE_ALLOW_APP_TESTS=1", allow_module_level=True)

TEST_ROOT = tempfile.mkdtemp(prefix="point-quote-test-")
os.environ["POINT_QUOTE_DATA_DIR"] = os.path.join(TEST_ROOT, "data")
os.environ["POINT_QUOTE_VAULT_ROOT"] = TEST_ROOT

import fitz
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

from backend.app import app, APP_VERSION, DATA, DB, EXPORTS, LOG_FILE, UPLOADS, knowledge_repository, migrate_legacy_catalog, product_repository
from backend.pptx_export import export_quote_pptx
from backend.product_repository import write_note
from backend.product_repository import ObsidianProductRepository
import backend.product_repository as product_repository_module
from backend.file_writes import atomic_write_text

client = TestClient(app)

# Resolve fixtures from the isolated package root, not the deprecated env override.
TEST_ROOT = str(DATA)


def request(method, path, **kwargs):
    response = client.request(method, path, **kwargs)
    assert response.status_code < 400, response.text
    return response.json()


def test_v32_floor_layer_copy_and_independent_drawing():
    project=request("POST","/api/projects",json={"name":"v32 floor/layer regression","customer":"","address":"","status":"进行中"})
    pid=project["id"]
    current=request("GET",f"/api/projects/{pid}")
    assert len(current["floors"])==len(current["layers"])==1
    first_floor=current["floors"][0]["id"]
    first_layer=current["layers"][0]["id"]
    second=request("POST",f"/api/projects/{pid}/floors",json={"name":"二楼"})
    second_floor=next(floor["id"] for floor in second["floors"] if floor["name"]=="二楼")
    image=Image.new("RGB",(320,240),"white"); buffer=io.BytesIO(); image.save(buffer,"PNG")
    source=request("POST",f"/api/projects/{pid}/drawings",files={"file":("floor.png",buffer.getvalue(),"image/png")},data={"page":"1","floor_id":first_floor})
    copied_project=request("POST",f"/api/projects/{pid}/drawings/{source['id']}/copy",json={"floor_id":second_floor})
    target=next(drawing for drawing in copied_project["drawings"] if drawing["floor_id"]==second_floor)
    point=request("POST",f"/api/projects/{pid}/points?drawing_id={source['id']}",json={"type_id":"curtain","layer_id":first_layer,"x":.2,"y":.4,"x2":.6,"y2":.4,"line_width":5,"size":20})
    copied_project=request("POST",f"/api/projects/{pid}/layers/{first_layer}/copy",json={"floor_id":second_floor})
    copied=next(item for item in copied_project["points"] if item["drawing_id"]==target["id"])
    assert copied["id"]!=point["id"] and copied["layer_id"]!=first_layer
    assert (copied["x"],copied["y"],copied["x2"],copied["y2"],copied["line_width"],copied["size"]) == (.2,.4,.6,.4,5,20)
    request("PATCH",f"/api/projects/{pid}/drawings/{target['id']}",json={"name":"二楼独立图纸"})
    reopened=request("GET",f"/api/projects/{pid}")
    assert next(drawing for drawing in reopened["drawings"] if drawing["id"]==source["id"])["name"]!="二楼独立图纸"
    assert next(drawing for drawing in reopened["drawings"] if drawing["id"]==target["id"])["name"]=="二楼独立图纸"


def test_v32_export_only_selected_layer():
    project=request("POST","/api/projects",json={"name":"v32 selected layer export"})
    pid=project["id"]
    project=request("GET",f"/api/projects/{pid}")
    floor_id=project["floors"][0]["id"]
    first_layer=project["layers"][0]["id"]
    updated=request("POST",f"/api/projects/{pid}/floors/{floor_id}/layers",json={"name":"另一图层"})
    second_layer=next(layer["id"] for layer in updated["layers"] if layer["name"]=="另一图层")
    image=io.BytesIO(); Image.new("RGB",(320,240),"white").save(image,"PNG")
    drawing=request("POST",f"/api/projects/{pid}/drawings",files={"file":("layers.png",image.getvalue(),"image/png")},data={"page":"1","floor_id":floor_id})
    request("POST",f"/api/projects/{pid}/points?drawing_id={drawing['id']}",json={"type_id":"switch","layer_id":first_layer,"x":.2,"y":.3})
    request("POST",f"/api/projects/{pid}/points?drawing_id={drawing['id']}",json={"type_id":"water","layer_id":second_layer,"x":.7,"y":.6})
    first=client.get(f"/api/projects/{pid}/drawings/{drawing['id']}/export.png?layer_id={first_layer}")
    second=client.get(f"/api/projects/{pid}/drawings/{drawing['id']}/export.png?layer_id={second_layer}")
    assert first.status_code==second.status_code==200
    assert first.content!=second.content
    assert client.get(f"/api/projects/{pid}/drawings/{drawing['id']}/export.png?layer_id=missing").status_code==404


def test_rq017a_request_ids_structured_errors_and_scan_metrics():
    success = client.get("/api/version", headers={"X-Request-ID": "REQ-success-check"})
    assert success.status_code == 200
    assert success.headers["X-Request-ID"] == "REQ-success-check"

    locked = OSError("The process cannot access the file because it is being used by another process")
    # Read endpoints use the local index and must not fail because SMB is locked.
    with patch.object(knowledge_repository, "sync", side_effect=locked):
        local_read = client.get("/api/projects")
        assert local_read.status_code == 200
    with patch.object(product_repository, "list", side_effect=locked):
        failed = client.get("/api/products", headers={"X-Request-ID": "REQ-locked-check"})
    assert failed.status_code == 500
    assert failed.headers["X-Request-ID"] == "REQ-locked-check"
    assert failed.json() == {
        "detail": "本地保存失败，诊断日志已记录。",
        "error_code": "FILE_LOCKED",
        "request_id": "REQ-locked-check",
    }

    status = client.get("/api/knowledge-sync/status")
    assert status.status_code == 200
    assert status.json()["sync_status"] == "local"
    assert status.json()["conflicts"] == []

    entries = [json.loads(line) for line in LOG_FILE.read_text(encoding="utf-8").splitlines()]
    error = next(entry for entry in entries if entry.get("requestId") == "REQ-locked-check")
    assert error["errorType"] == "FILE_LOCKED"
    assert error["route"] == "/api/products"
    assert error["method"] == "GET"
    assert error["stackTrace"]


def test_rq017c_product_index_reuses_unchanged_notes_and_refreshes_changed_note(tmp_path):
    product_dir = tmp_path / "products"
    product_dir.mkdir()
    product_id = "a" * 32
    note = product_dir / "product.md"
    note.write_text(
        "---\nid: " + product_id + "\n类型: 报价产品\n产品名称: 索引测试\n品牌: 测试品牌\n分类: 智能\n点位类型: switch\n所属系统: lighting\n售价_分: 100\n启用: true\n更新时间: t1\n---\n",
        encoding="utf-8",
    )
    repo = ObsidianProductRepository()
    scan_count = 0
    original_rglob = Path.rglob

    def count_product_walk(path, pattern):
        nonlocal scan_count
        if path == product_dir:
            scan_count += 1
        return original_rglob(path, pattern)

    with patch.object(product_repository_module, "PRODUCT_DIR", product_dir), patch.object(product_repository_module, "VAULT_ROOT", tmp_path), patch.object(Path, "rglob", count_product_walk):
        first = repo.list(include_disabled=True)
        assert len(first) == 1 and first[0]["standard_cents"] == 100
        # Direct project/quotation lookups hit the indexed path and parsed row.
        for _ in range(8):
            assert repo.get(None, product_id)["name"] == "索引测试"
        # Canvas/quote lookups must use the warmed ID map, with no per-point SMB stat.
        with patch.object(Path, "stat", side_effect=AssertionError("cached lookup touched the filesystem")):
            for _ in range(8):
                assert repo.get_cached(product_id)["name"] == "索引测试"
        assert scan_count == 1
        # Expired catalog inventory is stale-while-revalidate: a list response
        # must return the warm cache while the SMB refresh is still in progress.
        release_refresh = threading.Event()
        repo._inventory_at = 0
        with patch.object(repo, "_refresh_inventory_background", side_effect=lambda: release_refresh.wait(1)):
            started = time.perf_counter()
            assert len(repo.list(include_disabled=True)) == 1
            assert time.perf_counter() - started < .1
            release_refresh.set()
        note.write_text(
            "---\nid: " + product_id + "\n类型: 报价产品\n产品名称: 索引测试\n品牌: 测试品牌\n分类: 智能\n点位类型: switch\n所属系统: lighting\n售价_分: 250\n启用: true\n更新时间: t2\n---\n",
            encoding="utf-8",
        )
        refreshed = repo.get(None, product_id)
        assert refreshed["standard_cents"] == 250
        assert scan_count == 2


def test_portable_point_save_never_starts_background_knowledge_sync():
    project = request("POST", "/api/projects", json={"name": "Portable 本地点位测试"})
    image = io.BytesIO(); Image.new("RGB", (320, 240), "white").save(image, "PNG")
    drawing = request("POST", f"/api/projects/{project['id']}/drawings",
                      files={"file": ("portable.png", image.getvalue(), "image/png")})
    with patch.object(knowledge_repository, "sync", side_effect=AssertionError("Portable mode must not invoke sync")) as sync:
        response = client.post(f"/api/projects/{project['id']}/points?drawing_id={drawing['id']}",
                               json={"type_id": "sensor", "x": .4, "y": .4})
        assert response.status_code == 201, response.text
        assert client.get(f"/api/projects/{project['id']}").json()["points"][0]["type_id"] == "sensor"
        assert sync.call_count == 0


def test_portable_sync_status_and_retry_are_local_only():
    with patch.object(knowledge_repository, "sync", side_effect=AssertionError("Portable sync endpoint invoked sync")) as sync:
        status = client.get("/api/knowledge-sync/status")
        retry = client.post("/api/knowledge-sync/retry")
    assert status.status_code == retry.status_code == 200
    assert status.json()["sync_status"] == retry.json()["sync_status"] == "local"
    assert status.json()["conflicts"] == retry.json()["conflicts"] == []
    assert sync.call_count == 0


def test_portable_project_create_edit_survives_restart_without_sync_retry():
    with patch.object(knowledge_repository, "sync", side_effect=AssertionError("Portable write invoked sync")) as sync:
        created = client.post("/api/projects", json={"name": "本地独立保存测试"})
        assert created.status_code == 201, created.text
        project_id = created.json()["id"]
        updated = client.patch(f"/api/projects/{project_id}", json={"name": "本地独立保存测试", "customer": "本地客户"})
        assert updated.status_code == 200, updated.text
        assert client.get(f"/api/projects/{project_id}").json()["customer"] == "本地客户"
        retried = client.post("/api/knowledge-sync/retry")
        assert retried.status_code == 200
        assert sync.call_count == 0
    path = knowledge_repository.project_path(project_id)
    assert path and path.is_file()
    snapshot = knowledge_repository.read_project(path)
    assert snapshot and snapshot["tables"]["projects"][0]["customer"] == "本地客户"


def test_v31_brand_logo_product_links_and_delete_guard():
    brand=request("POST","/api/brands",json={"name":"v31品牌"})
    svg=b'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1 1"/>'
    first=client.post(f"/api/brands/{brand['id']}/logo",files={"file":("logo.svg",svg,"image/svg+xml")})
    assert first.status_code==200
    first_name=first.json()["logo_path"]
    png=io.BytesIO(); Image.new("RGBA",(2,2),(0,0,0,0)).save(png,"PNG")
    second=client.post(f"/api/brands/{brand['id']}/logo",files={"file":("logo.png",png.getvalue(),"image/png")})
    assert second.status_code==200 and second.json()["logo_path"]!=first_name
    assert second.json()["logo_path"].endswith(".png")
    product=request("POST","/api/products",json={"brand_id":brand["id"],"name":"网关产品","category":"other","device_type":"cloud_gateway","system_category":"network","standard_price":"100.00"})
    linked=client.get(f"/api/brands/{brand['id']}/products")
    assert linked.status_code==200 and [item["id"] for item in linked.json()]==[product["id"]]
    blocked=client.delete(f"/api/brands/{brand['id']}")
    assert blocked.status_code==409
    assert client.delete(f"/api/brands/{brand['id']}/logo").json()["logo_path"] is None
    assert client.delete(f"/api/products/{product['id']}").status_code==200
    assert client.delete(f"/api/brands/{brand['id']}").status_code==200


def test_v31_cable_system_and_bulk_service_fee_remain_individually_editable():
    project=request("POST","/api/projects",json={"name":"v31报价批量费测试"})
    quote=request("GET",f"/api/projects/{project['id']}/quotation-v2")
    cable=next((system for system in quote["system_categories"] if system["id"]=="cable"),None)
    assert cable is not None
    applied=client.patch(f"/api/projects/{project['id']}/quotation-v2/service-fees",json={"service_type":"percentage","service_rate":"8","service_amount":"0","service_name":"统一服务费"})
    assert applied.status_code==200, applied.text
    assert all(system["service_type"]=="percentage" and system["service_rate_basis"]==800 for system in applied.json()["systems"])
    target=applied.json()["systems"][0]
    edited=client.patch(f"/api/quotation-systems/{target['id']}",json={"service_type":"fixed","service_rate":"0","service_amount":"250","service_name":"单独服务费"})
    assert edited.status_code==200, edited.text
    values={system["id"]:system for system in edited.json()["systems"]}
    assert values[target["id"]]["service_type"]=="fixed"
    assert all(system["service_type"]=="percentage" for id,system in values.items() if id!=target["id"])


def test_v31_point_paste_creates_new_id_and_preserves_point_fields():
    project=request("POST","/api/projects",json={"name":"v31复制粘贴测试"})
    image=io.BytesIO(); Image.new("RGB",(320,240),"white").save(image,"PNG")
    drawing=request("POST",f"/api/projects/{project['id']}/drawings",files={"file":("paste.png",image.getvalue(),"image/png")})
    source=request("POST",f"/api/projects/{project['id']}/points?drawing_id={drawing['id']}",json={"type_id":"sensor","x":.3,"y":.4,"quantity":3,"room":"客厅","note":"复制属性"})
    payload={**source,"x":.35,"y":.45}
    pasted=client.post("/api/points/paste",json=payload)
    assert pasted.status_code==200, pasted.text
    copy=pasted.json()
    assert copy["id"]!=source["id"] and copy["number"]!=source["number"]
    assert (copy["x"],copy["y"],copy["quantity"],copy["room"],copy["note"]) == (.35,.45,3,"客厅","复制属性")


def test_full_flow():
    project = request("POST", "/api/projects", json={"name":"中文住宅测试","customer":"测试客户","address":"上海","status":"进行中"})
    pid = project["id"]
    brand = request("POST", "/api/brands", json={"name":"真实测试品牌"})
    product = request("POST", "/api/products", json={"brand_id":brand["id"],"name":"测试开关","series":"测试系列","model":"S-1","category":"switch","standard_price":"199.90","unit":"个","params":"220V"})
    assert product["standard_price"] == "199.90"
    assert product["category"] == "智能" and product["device_type"] == "switch"
    product_image=Image.new("RGB",(200,200),"#ff0000"); ib=io.BytesIO(); product_image.save(ib,"PNG")
    request("POST",f"/api/products/{product['id']}/image",files={"file":("产品.png",ib.getvalue(),"image/png")})
    image = Image.new("RGB", (1200, 800), "white")
    b = io.BytesIO(); image.save(b, "PNG")
    drawing = request("POST", f"/api/projects/{pid}/drawings", files={"file":("户型图.png", b.getvalue(), "image/png")}, data={"page":"1"})
    assert (drawing["width"], drawing["height"]) == (1200,800)
    p1 = request("POST", f"/api/projects/{pid}/points?drawing_id={drawing['id']}", json={"type_id":"switch","x":0.25,"y":0.3})
    p2 = request("POST", f"/api/projects/{pid}/points?drawing_id={drawing['id']}", json={"type_id":"switch","x":0.5,"y":0.5})
    assert (p1["number"],p2["number"]) == ("开关01","开关02")
    request("DELETE", f"/api/points/{p2['id']}")
    p3 = request("POST", f"/api/projects/{pid}/points?drawing_id={drawing['id']}", json={"type_id":"switch","x":0.7,"y":0.5})
    assert p3["number"] == "开关03"
    restored=request("POST","/api/points/restore",json=p2)
    assert restored["id"]==p2["id"] and restored["number"]=="开关02"
    request("DELETE",f"/api/points/{p2['id']}")
    request("PATCH", f"/api/points/{p1['id']}", json={"room":"客厅","quantity":2,"name":"客厅主开关","x":0.33})
    request("POST", "/api/points/assign-product", json={"point_ids":[p1["id"],p3["id"]],"product_id":product["id"]})
    request("PATCH", f"/api/points/{p1['id']}", json={"actual_price":"180.05"})
    quote = request("GET",f"/api/projects/{pid}/quote")
    assert len(quote["detail"]) == 2 and len(quote["summary"]) == 2
    assert quote["line_total"] == "560.00"
    # 2 × 180.05 + 1 × 199.90 = 560.00
    request("PUT",f"/api/products/{product['id']}",json={"brand_id":brand["id"],"name":"测试开关","series":"测试系列","model":"S-1","category":"switch","standard_price":"299.90","unit":"个"})
    reloaded=request("GET",f"/api/projects/{pid}")
    old_point=next(p for p in reloaded["points"] if p["id"]==p1["id"])
    assert old_point["standard_price"]=="199.90" and old_point["actual_price"]=="180.05"
    assert request("GET",f"/api/products/{product['id']}")["standard_price"]=="299.90"
    request("POST",f"/api/projects/{pid}/fees",json={"name":"安装","kind":"安装费","amount":"80.10"})
    request("PATCH",f"/api/projects/{pid}/quote",json={"note":"中文备注"})
    quote=request("GET",f"/api/projects/{pid}/quote")
    assert quote["total"]=="640.10" and "discount" not in quote["quote"]
    png=client.get(f"/api/projects/{pid}/drawings/{drawing['id']}/export.png")
    assert png.status_code==200
    exported=Image.open(io.BytesIO(png.content)); assert exported.width>=1200 and exported.height>=800
    pdf=client.get(f"/api/projects/{pid}/drawings/{drawing['id']}/export.pdf")
    assert pdf.status_code==200 and len(fitz.open(stream=pdf.content,filetype="pdf"))==1
    xlsx=client.get(f"/api/projects/{pid}/quote/export.xlsx")
    assert xlsx.status_code==200
    wb=load_workbook(io.BytesIO(xlsx.content),data_only=True)
    assert wb["产品清单"]["A5"].value=="测试开关" and wb["产品清单"]["C5"].value=="S-1" and wb["产品清单"]["D5"].value=="220V"
    assert len(wb["产品清单"]._images)==2
    assert wb["项目汇总"].max_row>=2
    pdf_quote=client.get(f"/api/projects/{pid}/quote/export.pdf")
    assert pdf_quote.status_code==200
    quote_doc=fitz.open(stream=pdf_quote.content,filetype="pdf")
    assert len(quote_doc)==2 and quote_doc[0].rect.width>quote_doc[0].rect.height
    pptx_quote=client.get(f"/api/projects/{pid}/quote/export.pptx")
    assert pptx_quote.status_code==200
    assert pptx_quote.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.presentationml.presentation")
    assert ".pptx" in pptx_quote.headers["content-disposition"]
    with zipfile.ZipFile(io.BytesIO(pptx_quote.content)) as package:
        names=set(package.namelist())
        assert "ppt/slides/slide1.xml" in names and "ppt/slides/slide2.xml" in names
        assert any(name.startswith("ppt/media/") for name in names)
        assert b"PingFang SC" in package.read("ppt/slides/slide1.xml")
        for name in names:
            if name.endswith(".xml") or name.endswith(".rels"):
                ET.fromstring(package.read(name))
    assert Path(DB).exists() and any(EXPORTS.iterdir())


def test_pdf_second_page_upload():
    project=request("POST","/api/projects",json={"name":"PDF 测试"})
    source=fitz.open(); source.new_page(); second=source.new_page(width=300,height=200); second.draw_rect(fitz.Rect(20,20,100,100),color=(1,0,0),fill=(1,0,0))
    drawing=request("POST",f"/api/projects/{project['id']}/drawings",files={"file":("图纸.pdf",source.tobytes(),"application/pdf")},data={"page":"2"})
    assert drawing["page"]==2 and drawing["width"]==600 and drawing["height"]==400


def test_pptx_detail_rows_are_paginated():
    product = {"name": "分页产品", "model": "M-1", "params": "参数说明", "brand_name": "品牌", "origin": "中国", "unit": "台", "image_path": None, "category": "智能"}
    details = [{"product": product, "quantity": 1, "actual_price": "10.00", "subtotal": "10.00"} for _ in range(13)]
    data = {"project": {"name": "分页测试", "customer": "", "address": ""}, "detail": details, "summary": [{"product": product, "quantity": 13, "actual_price": "10.00", "subtotal": "130.00"}], "fees": [], "total": "130.00"}
    output = Path(TEST_ROOT) / "分页测试.pptx"
    export_quote_pptx(data, output, lambda _: None)
    with zipfile.ZipFile(output) as package:
        slide_names = [name for name in package.namelist() if name.startswith("ppt/slides/slide") and name.endswith(".xml")]
        assert len(slide_names) == 4  # 6 + 6 + 1 detail rows, then summary


def test_csv_import_atomic_validation():
    content="品牌,产品名称,分类,标准价格,型号\n品牌乙,窗帘电机,窗帘,1234.56,M2\n".encode("utf-8-sig")
    result=request("POST","/api/products/import",files={"file":("products.csv",content,"text/csv")})
    assert result["imported"]==1
    assert any(p["name"]=="窗帘电机" and p["model"]=="M2" for p in request("GET","/api/products"))


def test_summary_merges_only_same_product_and_project_price():
    project=request("POST","/api/projects",json={"name":"汇总测试"})
    brand=request("POST","/api/brands",json={"name":"汇总品牌"})
    product=request("POST","/api/products",json={"brand_id":brand["id"],"name":"产品甲","category":"switch","standard_price":"0.10"})
    b=io.BytesIO(); Image.new("RGB",(300,200),"white").save(b,"PNG")
    drawing=request("POST",f"/api/projects/{project['id']}/drawings",files={"file":("a.png",b.getvalue(),"image/png")})
    ids=[]
    for x in (.1,.3,.5):
        ids.append(request("POST",f"/api/projects/{project['id']}/points?drawing_id={drawing['id']}",json={"type_id":"switch","x":x,"y":.5})["id"])
    request("POST","/api/points/assign-product",json={"point_ids":ids,"product_id":product["id"]})
    quote=request("GET",f"/api/projects/{project['id']}/quote")
    assert len(quote["summary"])==1 and quote["summary"][0]["quantity"]==3 and quote["total"]=="0.30"
    request("PATCH",f"/api/points/{ids[0]}",json={"actual_price":"0.11"})
    quote=request("GET",f"/api/projects/{project['id']}/quote")
    assert len(quote["summary"])==2 and quote["total"]=="0.31"


def test_product_can_be_added_directly_to_quote_without_drawing_point():
    project=request("POST","/api/projects",json={"name":"报价直接加项"})
    product=request("POST","/api/products",json={"brand_name":"清单品牌","name":"后增设备","category":"其他设备","standard_price":"88.80","unit":"台"})
    quote=request("POST",f"/api/projects/{project['id']}/quote-items",json={"product_id":product["id"],"quantity":2,"actual_price":"80.05","note":"不修改图纸"})
    assert len(quote["detail"])==1 and quote["detail"][0]["source"]=="manual"
    assert quote["detail"][0]["point_id"] is None and quote["total"]=="160.10"
    line_id=quote["detail"][0]["line_id"]
    quote=request("PATCH",f"/api/quote-items/{line_id}",json={"quantity":3,"actual_price":"81.00"})
    assert quote["total"]=="243.00" and quote["summary"][0]["quantity"]==3
    quote=request("POST",f"/api/quote-items/{line_id}/restore-price",json={})
    assert quote["total"]=="266.40"
    exported=client.get(f"/api/projects/{project['id']}/quote/export.xlsx")
    assert exported.status_code==200
    sheet=load_workbook(io.BytesIO(exported.content),data_only=True)["产品清单"]
    assert sheet["A5"].value=="后增设备" and sheet["G5"].value==3
    request("DELETE",f"/api/products/{product['id']}")
    assert request("GET",f"/api/projects/{project['id']}/quote")["detail"][0]["product"]["name"]=="后增设备"
    request("DELETE",f"/api/quote-items/{line_id}")
    quote=request("GET",f"/api/projects/{project['id']}/quote")
    assert quote["detail"]==[] and quote["total"]=="0.00"


def test_multiple_products_can_be_added_to_quote_in_one_request():
    project=request("POST","/api/projects",json={"name":"批量报价加项"})
    first=request("POST","/api/products",json={"brand_name":"批量品牌","name":"产品一","category":"智能","standard_price":"10.10","unit":"台"})
    second=request("POST","/api/products",json={"brand_name":"批量品牌","name":"产品二","category":"网络","standard_price":"20.20","unit":"个"})
    quote=request("POST",f"/api/projects/{project['id']}/quote-items/batch",json={"product_ids":[first["id"],second["id"],first["id"]],"quantity":2,"note":"批量加入"})
    assert len(quote["detail"])==2
    assert {line["product"]["name"] for line in quote["detail"]}=={"产品一","产品二"}
    assert all(line["quantity"]==2 and line["note"]=="批量加入" for line in quote["detail"])
    assert quote["total"]=="60.60"


def test_obsidian_note_is_catalog_source():
    product=request("POST","/api/products",json={"brand_name":"真实品牌","name":"真实产品","model":"TP-100","category":"智能电箱模块","device_type":"other","standard_price":"100.00","cost_price":"40.00","origin":"中国","unit":"台","description":"原始说明"})
    path=Path(TEST_ROOT)/product["note_path"]
    assert path.exists() and path.parent.name=="未分类"
    content=path.read_text(encoding="utf-8")
    assert "审核人: 待审核" in content and "型号: TP-100" in content and "| 型号 | TP-100 |" in content and "<!-- pointquote:start -->" in content
    assert "点位类型: system" in content
    assert "margin_percent" not in product and product["category"]=="智能" and product["device_type"]=="system"
    picture=Image.new("RGBA",(80,80),(255,0,0,0)); picture.putpixel((40,40),(255,0,0,255)); buffer=io.BytesIO(); picture.save(buffer,"PNG")
    updated=request("POST",f"/api/products/{product['id']}/image",files={"file":("photo.png",buffer.getvalue(),"image/png")})
    assert updated["image_path"] and (Path(TEST_ROOT)/"09_附件"/"产品图片"/updated["image_path"]).is_file()
    assert updated["image_path"].endswith(".png") and (Path(TEST_ROOT)/"09_附件"/"产品图片"/updated["image_path"]).read_bytes()==buffer.getvalue()
    with Image.open(Path(TEST_ROOT)/"09_附件"/"产品图片"/updated["image_path"]) as stored:
        assert stored.mode=="RGBA" and stored.getpixel((0,0))[3]==0
    assert f"![[{updated['image_path']}]]" in path.read_text(encoding="utf-8")
    text=path.read_text(encoding="utf-8")
    path.write_text(text.replace("售价_分: 10000","售价_分: 12345")+"\n## 手工维护\n\n保留这段内容。\n",encoding="utf-8")
    fresh=request("GET",f"/api/products/{product['id']}")
    assert fresh["standard_price"]=="123.45"
    request("PUT",f"/api/products/{product['id']}",json={"brand_name":"真实品牌","name":"真实产品新版","category":"智能电箱模块","device_type":"other","standard_price":"123.45","unit":"台"})
    assert "保留这段内容" in path.read_text(encoding="utf-8")
    request("DELETE",f"/api/products/{product['id']}")
    assert not path.exists() and client.get(f"/api/products/{product['id']}").status_code==404
    assert product_repository.get_cached(product["id"]) is None
    assert not any(p["id"]==product["id"] for p in request("GET","/api/products"))


def test_windows_legacy_catalog_migrates_to_shared_notes():
    legacy_dir=Path(TEST_ROOT)/"windows-local"; legacy_dir.mkdir()
    legacy_db=legacy_dir/"app.sqlite3"
    db=sqlite3.connect(legacy_db)
    db.executescript("""
    CREATE TABLE brands(id TEXT PRIMARY KEY,name TEXT NOT NULL);
    CREATE TABLE products(id TEXT PRIMARY KEY,brand_id TEXT NOT NULL,name TEXT NOT NULL,series TEXT,model TEXT,category TEXT,standard_cents INTEGER,unit TEXT,params TEXT,description TEXT,image_path TEXT,enabled INTEGER,updated_at TEXT);
    """)
    db.execute("INSERT INTO brands VALUES(?,?)",("legacy-brand","Windows旧品牌"))
    db.execute("INSERT INTO products VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",("a"*32,"legacy-brand","Windows旧产品","","","network",68000,"台","千兆","旧库迁移",None,1,"2026-09-27T10:00:00+08:00"))
    db.commit(); db.close()
    report=migrate_legacy_catalog(legacy_db)
    assert report["imported"]==1 and not report["error"]
    migrated=request("GET","/api/products/"+"a"*32)
    assert migrated["category"]=="网络" and migrated["device_type"]=="network"
    assert (Path(TEST_ROOT)/migrated["note_path"]).parent.name=="未分类"


def test_windows_smb_existing_note_replace_fallback(monkeypatch, tmp_path):
    note=tmp_path/"产品.md"
    note.write_text("旧内容",encoding="utf-8")
    real_replace=os.replace

    def windows_share_replace(source, destination):
        if Path(destination)==note and note.exists():
            raise PermissionError("simulated Windows SMB overwrite restriction")
        return real_replace(source,destination)

    monkeypatch.setattr(os,"replace",windows_share_replace)
    monkeypatch.setattr("backend.file_writes.time.sleep",lambda _:None)
    write_note(note,"新内容\n")
    assert note.read_text(encoding="utf-8")=="新内容\n"


def test_rq017b_atomic_writes_deduplicate_and_serialize(tmp_path):
    note=tmp_path/"并发保存.md"
    assert atomic_write_text(note,"相同内容\n")=="written"
    assert atomic_write_text(note,"相同内容\n")=="unchanged"
    values=[("甲"*20000)+"\n",("乙"*20000)+"\n"]
    with ThreadPoolExecutor(max_workers=2) as pool:
        results=list(pool.map(lambda value:atomic_write_text(note,value),values))
    assert results==["written","written"]
    assert note.read_text(encoding="utf-8") in values
    assert not list(tmp_path.glob("*.tmp")) and not list(tmp_path.glob("*.bak"))


def test_geometry_drawings_batch_and_feedback():
    project=request("POST","/api/projects",json={"name":"几何与反馈"})
    assert request("GET","/api/version")["version"]==APP_VERSION
    storage=request("GET","/api/storage-status")
    assert storage["categories"]==["智能","影音","灯光","网络"] and storage["product_dir"].endswith("01_产品库")
    product_response=client.get("/api/products?include_disabled=true")
    assert product_response.status_code==200 and product_response.headers["cache-control"].startswith("no-store")
    project=request("PATCH",f"/api/projects/{project['id']}",json={"name":"几何与反馈新版","customer":"客户","address":"地址","status":"进行中"})
    assert project["name"]=="几何与反馈新版" and project["customer"]=="客户"
    image=io.BytesIO(); Image.new("RGB",(800,600),"white").save(image,"PNG")
    drawing=request("POST",f"/api/projects/{project['id']}/drawings",files={"file":("一层.png",image.getvalue(),"image/png")})
    assert drawing["name"]=="一层"
    renamed=request("PATCH",f"/api/projects/{project['id']}/drawings/{drawing['id']}",json={"name":"首层平面"})
    assert renamed["drawings"][0]["name"]=="首层平面"
    curtain=request("POST",f"/api/projects/{project['id']}/points?drawing_id={drawing['id']}",json={"type_id":"curtain","x":.1,"y":.2,"x2":.6,"y2":.2})
    area=request("POST",f"/api/projects/{project['id']}/points?drawing_id={drawing['id']}",json={"type_id":"light","x":.2,"y":.3,"x2":.7,"y2":.8})
    assert curtain["shape"]=="line" and curtain["x2"]==.6
    assert area["shape"]=="area" and area["y2"]==.8
    area=request("PATCH",f"/api/points/{area['id']}",json={"x":.15,"y":.25,"x2":.75,"y2":.85})
    assert (area["x"],area["y"],area["x2"],area["y2"])==(.15,.25,.75,.85)
    rotated=request("POST",f"/api/projects/{project['id']}/drawings/{drawing['id']}/rotate",json={})
    rotated_drawing=rotated["drawings"][0]; rotated_area=next(p for p in rotated["points"] if p["id"]==area["id"])
    assert (rotated_drawing["width"],rotated_drawing["height"],rotated_drawing["rotation"])==(600,800,90)
    assert tuple(round(rotated_area[k],4) for k in ("x","y","x2","y2"))==(.75,.15,.15,.75)
    exported=client.get(f"/api/projects/{project['id']}/drawings/{drawing['id']}/export.pdf?show_name=false&show_number=false")
    assert exported.status_code==200 and exported.headers["cache-control"].startswith("no-store")
    product=request("POST","/api/products",json={"brand_name":"批量品牌","name":"智能主机","category":"智能系统","standard_price":"123.00"})
    result=request("PATCH","/api/products/batch",json={"ids":[product["id"]],"action":"category","category":"网络设备"})
    changed=request("GET",f"/api/products/{product['id']}")
    assert result["updated"]==1 and changed["category"]=="网络" and changed["device_type"]=="network"
    result=request("PATCH","/api/products/batch",json={"ids":[product["id"]],"action":"delete"})
    assert result["updated"]==1 and client.get(f"/api/products/{product['id']}").status_code==404
    feedback=request("POST","/api/feedback",json={"title":"建议","description":"增加快捷操作","module":"图纸编辑"})
    screenshot=io.BytesIO(); Image.new("RGBA",(20,20),(255,0,0,0)).save(screenshot,"PNG")
    uploaded=request("POST",f"/api/feedback/{feedback['id']}/screenshot",files={"file":("截图.png",screenshot.getvalue(),"image/png")})
    assert uploaded["screenshot_path"].endswith(".png") and request("GET","/api/feedback")[0]["status"]=="待处理"


def test_rq002_ap_types_non_smart_area_and_legacy_network_compatibility():
    types={item["id"]:item for item in request("GET","/api/types")}
    assert types["ceiling_ap"]["name"]=="吸顶 AP" and types["ceiling_ap"]["product_selectable"] is True
    assert types["wall_ap"]["name"]=="墙面 AP" and types["wall_ap"]["product_selectable"] is True
    assert types["non_smart_area"]["kind"]=="area" and types["non_smart_area"]["product_selectable"] is False
    assert types["network"]["legacy"] is True and types["network"]["editor_selectable"] is False

    ceiling_product=request("POST","/api/products",json={
        "brand_name":"AP 测试品牌","name":"吸顶无线接入点","category":"网络","device_type":"ceiling_ap",
        "system_category":"network","standard_price":"680.00","unit":"台",
    })
    wall_product=request("POST","/api/products",json={
        "brand_name":"AP 测试品牌","name":"墙面无线接入点","category":"网络","device_type":"wall_ap",
        "system_category":"lighting-control","standard_price":"580.00","unit":"台",
    })
    assert (ceiling_product["category"],ceiling_product["device_type"],ceiling_product["system_category"])==("网络","ceiling_ap","network")
    assert (wall_product["category"],wall_product["device_type"],wall_product["system_category"])==("网络","wall_ap","network")

    project=request("POST","/api/projects",json={"name":"RQ-002 测试"})
    image=io.BytesIO(); Image.new("RGB",(600,400),"white").save(image,"PNG")
    drawing=request("POST",f"/api/projects/{project['id']}/drawings",files={"file":("rq002.png",image.getvalue(),"image/png")})
    ceiling=request("POST",f"/api/projects/{project['id']}/points?drawing_id={drawing['id']}",json={"type_id":"ceiling_ap","x":.2,"y":.25})
    wall=request("POST",f"/api/projects/{project['id']}/points?drawing_id={drawing['id']}",json={"type_id":"wall_ap","x":.7,"y":.25})
    area=request("POST",f"/api/projects/{project['id']}/points?drawing_id={drawing['id']}",json={"type_id":"non_smart_area","x":.1,"y":.55,"x2":.8,"y2":.9})
    legacy=request("POST",f"/api/projects/{project['id']}/points?drawing_id={drawing['id']}",json={"type_id":"network","x":.5,"y":.45})
    assert ceiling["number"]=="吸顶AP01" and wall["number"]=="墙面AP01"
    assert area["shape"]=="area" and (area["x2"],area["y2"])==(.8,.9)
    assert legacy["type_id"]=="network"  # 旧项目中的 network 仍能读取、保存和恢复

    request("POST","/api/points/assign-product",json={"point_ids":[ceiling["id"]],"product_id":ceiling_product["id"]})
    request("POST","/api/points/assign-product",json={"point_ids":[wall["id"]],"product_id":wall_product["id"]})
    rejected=client.post("/api/points/assign-product",json={"point_ids":[area["id"]],"product_id":ceiling_product["id"]})
    assert rejected.status_code==422 and "不能配置产品" in rejected.text

    quote=request("GET",f"/api/projects/{project['id']}/quote")
    assert len(quote["detail"])==2 and area["number"] not in quote["missing"]
    listed=next(item for item in request("GET","/api/projects") if item["id"]==project["id"])
    assert listed["point_count"]==3  # 两种 AP + 历史 network；非智能区域不计入点位数量
    exported=client.get(f"/api/projects/{project['id']}/drawings/{drawing['id']}/export.png")
    assert exported.status_code==200

    moved=request("PATCH",f"/api/points/{area['id']}",json={"x":.15,"y":.6,"x2":.75,"y2":.88})
    assert (moved["x"],moved["y"],moved["x2"],moved["y2"])==(.15,.6,.75,.88)
    request("DELETE",f"/api/points/{area['id']}")
    reloaded=request("GET",f"/api/projects/{project['id']}")
    assert all(point["id"]!=area["id"] for point in reloaded["points"])


def test_rq003c_curtain_direction_motor_visibility_and_geometry_lifecycle():
    project=request("POST","/api/projects",json={"name":"RQ-003C 窗帘方向测试"})
    image=io.BytesIO(); Image.new("RGB",(800,600),"white").save(image,"PNG")
    drawing=request("POST",f"/api/projects/{project['id']}/drawings",files={"file":("curtain.png",image.getvalue(),"image/png")})
    curtain=request("POST",f"/api/projects/{project['id']}/points?drawing_id={drawing['id']}",json={"type_id":"curtain","x":.15,"y":.2,"x2":.65,"y2":.4})
    assert curtain["shape"]=="line" and curtain["show_motor"]==1
    assert (curtain["x"],curtain["y"],curtain["x2"],curtain["y2"])==(.15,.2,.65,.4)

    hidden=request("PATCH",f"/api/points/{curtain['id']}",json={"show_motor":False})
    assert hidden["show_motor"]==0
    moved=request("PATCH",f"/api/points/{curtain['id']}",json={"x":.2,"y":.25,"x2":.7,"y2":.45})
    assert (moved["x"],moved["y"],moved["x2"],moved["y2"],moved["show_motor"])==(.2,.25,.7,.45,0)

    reloaded=request("GET",f"/api/projects/{project['id']}")
    saved=next(point for point in reloaded["points"] if point["id"]==curtain["id"])
    assert saved["show_motor"]==0 and (saved["x"],saved["y"],saved["x2"],saved["y2"])==(.2,.25,.7,.45)

    rotated=request("POST",f"/api/projects/{project['id']}/drawings/{drawing['id']}/rotate",json={})
    rotated_curtain=next(point for point in rotated["points"] if point["id"]==curtain["id"])
    assert tuple(round(rotated_curtain[key],4) for key in ("x","y","x2","y2"))==(.75,.2,.55,.7)
    assert rotated_curtain["show_motor"]==0

    shown=request("PATCH",f"/api/points/{curtain['id']}",json={"show_motor":True})
    duplicate=request("POST",f"/api/points/{curtain['id']}/duplicate",json={})
    assert shown["show_motor"]==1 and duplicate["show_motor"]==1
    visible_export=client.get(f"/api/projects/{project['id']}/drawings/{drawing['id']}/export.png?show_name=false&show_number=false")
    restored=request("POST","/api/points/restore",json={**shown,"show_motor":False})
    assert restored["show_motor"]==0

    exported=client.get(f"/api/projects/{project['id']}/drawings/{drawing['id']}/export.png?show_name=false&show_number=false")
    assert visible_export.status_code==200 and exported.status_code==200
    assert visible_export.content!=exported.content and Image.open(io.BytesIO(exported.content)).width>=800
    request("DELETE",f"/api/points/{duplicate['id']}")
    assert all(point["id"]!=duplicate["id"] for point in request("GET",f"/api/projects/{project['id']}")["points"])


def test_rq003d_camera_sector_direction_angle_distance_and_lifecycle():
    project=request("POST","/api/projects",json={"name":"RQ-003D 摄像头范围测试"})
    image=io.BytesIO(); Image.new("RGB",(900,600),"white").save(image,"PNG")
    drawing=request("POST",f"/api/projects/{project['id']}/drawings",files={"file":("camera.png",image.getvalue(),"image/png")})
    camera=request("POST",f"/api/projects/{project['id']}/points?drawing_id={drawing['id']}",json={"type_id":"camera","x":.2,"y":.3,"x2":.6,"y2":.5,"fov_angle":75})
    assert camera["shape"]=="sector" and camera["fov_angle"]==75
    assert (camera["x"],camera["y"],camera["x2"],camera["y2"])==(.2,.3,.6,.5)

    adjusted=request("PATCH",f"/api/points/{camera['id']}",json={"x":.25,"y":.35,"x2":.65,"y2":.55,"fov_angle":120})
    assert (adjusted["x"],adjusted["y"],adjusted["x2"],adjusted["y2"],adjusted["fov_angle"])==(.25,.35,.65,.55,120)
    rejected=client.patch(f"/api/points/{camera['id']}",json={"fov_angle":181})
    assert rejected.status_code==422

    reloaded=request("GET",f"/api/projects/{project['id']}")
    saved=next(point for point in reloaded["points"] if point["id"]==camera["id"])
    assert saved["shape"]=="sector" and saved["fov_angle"]==120

    rotated=request("POST",f"/api/projects/{project['id']}/drawings/{drawing['id']}/rotate",json={})
    rotated_camera=next(point for point in rotated["points"] if point["id"]==camera["id"])
    assert tuple(round(rotated_camera[key],4) for key in ("x","y","x2","y2"))==(.65,.25,.45,.65)
    assert rotated_camera["fov_angle"]==120

    duplicate=request("POST",f"/api/points/{camera['id']}/duplicate",json={})
    assert duplicate["shape"]=="sector" and duplicate["fov_angle"]==120
    restored=request("POST","/api/points/restore",json={**rotated_camera,"fov_angle":45})
    assert restored["fov_angle"]==45 and restored["shape"]=="sector"

    exported=client.get(f"/api/projects/{project['id']}/drawings/{drawing['id']}/export.png?show_name=false&show_number=false")
    assert exported.status_code==200 and Image.open(io.BytesIO(exported.content)).width>=600
    request("DELETE",f"/api/points/{duplicate['id']}")
    assert all(point["id"]!=duplicate["id"] for point in request("GET",f"/api/projects/{project['id']}")["points"])


def test_rq004_project_signed_status_persists_and_does_not_replace_workflow_status():
    project=request("POST","/api/projects",json={"name":"RQ-004 已签状态测试","customer":"签约客户","status":"进行中"})
    assert project["signed"]==0 and project["status"]=="进行中"

    signed=request("PATCH",f"/api/projects/{project['id']}/signed",json={"signed":True})
    assert signed["signed"]==1 and signed["status"]=="进行中"
    listed=next(item for item in request("GET","/api/projects") if item["id"]==project["id"])
    assert listed["signed"]==1
    reopened=request("GET",f"/api/projects/{project['id']}")
    assert reopened["signed"]==1

    edited=request("PATCH",f"/api/projects/{project['id']}",json={"name":"RQ-004 已签状态测试新版","customer":"签约客户","address":"杭州","status":"已完成"})
    assert edited["signed"]==1 and edited["status"]=="已完成"

    unsigned=request("PATCH",f"/api/projects/{project['id']}/signed",json={"signed":False})
    assert unsigned["signed"]==0 and unsigned["status"]=="已完成"
    assert client.patch("/api/projects/not-found/signed",json={"signed":True}).status_code==404


def test_project_delete_cascades_but_keeps_catalog_product():
    product=request("POST","/api/products",json={"brand_name":"保留品牌","name":"保留产品","category":"智能开关","standard_price":"10.00"})
    project=request("POST","/api/projects",json={"name":"待删除项目"})
    image=io.BytesIO(); Image.new("RGB",(320,200),"white").save(image,"PNG")
    drawing=request("POST",f"/api/projects/{project['id']}/drawings",files={"file":("待删除.png",image.getvalue(),"image/png")})
    point=request("POST",f"/api/projects/{project['id']}/points?drawing_id={drawing['id']}",json={"type_id":"switch","x":.5,"y":.5})
    request("POST","/api/points/assign-product",json={"point_ids":[point["id"]],"product_id":product["id"]})
    background=UPLOADS/drawing["background_path"]
    assert background.exists()
    request("DELETE",f"/api/projects/{project['id']}")
    assert client.get(f"/api/projects/{project['id']}").status_code==404 and not background.exists()
    assert request("GET",f"/api/products/{product['id']}")["name"]=="保留产品"


def test_product_brand_case_variants_share_shadow_brand():
    # Match the failure: lowercase brand exists, uppercase product is added.
    for first, second in (("catalog-case", "CATALOG-CASE"), ("Straße", "STRASSE")):
        payload = {"name":"case regression", "brand_name":first, "category":"network", "standard_price":"100"}
        a = request("POST", "/api/products", json=payload)
        b = request("POST", "/api/products", json={**payload, "brand_name":second})
        assert a["brand_id"] == b["brand_id"]
        request("PUT", f"/api/products/{b['id']}", json={**payload, "brand_name":second, "standard_price":"120"})
        with sqlite3.connect(DB) as db:
            ids = db.execute("SELECT DISTINCT brand_id FROM products WHERE id IN (?,?)", (a["id"], b["id"])).fetchall()
        assert len(ids) == 1


def test_product_reuses_legacy_brand_case_variant():
    brand = request("POST", "/api/brands", json={"name":"legacy-case-brand"})
    product = request("POST", "/api/products", json={"name":"legacy regression", "brand_name":"LEGACY-CASE-BRAND", "category":"network", "standard_price":"100"})
    with sqlite3.connect(DB) as db:
        assert db.execute("SELECT brand_id FROM products WHERE id=?", (product["id"],)).fetchone()[0] == brand["id"]


def test_quotation_v2_systems_calculation_override_and_versions():
    project=request("POST","/api/projects",json={"name":"V2完整测试","customer":"测试客户","address":"测试地址"})
    product=request("POST","/api/products",json={"brand_name":"V2品牌","name":"DALI调光模块","model":"DL-8","category":"灯光","device_type":"light","system_category":"lighting-control","standard_price":"2500.00","unit":"台","params":"8路 · 0–10V · 24V · 0°C"})
    image=io.BytesIO(); Image.new("RGB",(600,400),"white").save(image,"PNG")
    drawing=request("POST",f"/api/projects/{project['id']}/drawings",files={"file":("V2图纸.png",image.getvalue(),"image/png")})
    point_ids=[]
    for index,quantity in enumerate((3,4)):
        point=request("POST",f"/api/projects/{project['id']}/points?drawing_id={drawing['id']}",json={"type_id":"light","x":.2+index*.2,"y":.4,"quantity":quantity,"room":"客厅" if index==0 else "卧室"})
        point_ids.append(point["id"])
    request("POST","/api/points/assign-product",json={"point_ids":point_ids,"product_id":product["id"]})
    quote=request("GET",f"/api/projects/{project['id']}/quotation-v2")
    lighting=next(system for system in quote["systems"] if system["category_id"]=="lighting-control")
    assert len(lighting["items"])==1 and lighting["items"][0]["quantity"]==7
    item=lighting["items"][0]
    assert item["product_code"]=="DL-8" and item["params"].startswith("8路")
    quote=request("PATCH",f"/api/quotation-items/{item['id']}",json={"quantity":9,"unit_price":"2999.99","remark":"手动覆盖"})
    item=next(system for system in quote["systems"] if system["category_id"]=="lighting-control")["items"][0]
    assert item["quantity_override"] is True and item["quantity"]==9 and item["unit_price_cents"]==299999
    request("PATCH",f"/api/points/{point_ids[0]}",json={"quantity":8})
    quote=request("GET",f"/api/projects/{project['id']}/quotation-v2")
    item=next(system for system in quote["systems"] if system["category_id"]=="lighting-control")["items"][0]
    assert item["source_quantity"]==12 and item["quantity"]==9
    quote=request("POST",f"/api/quotation-items/{item['id']}/restore-source-quantity",json={})
    lighting=next(system for system in quote["systems"] if system["category_id"]=="lighting-control")
    assert lighting["items"][0]["quantity"]==12 and lighting["items"][0]["quantity_override"] is False
    quote=request("PATCH",f"/api/quotation-systems/{lighting['id']}",json={"service_type":"percentage","service_rate":"10","service_amount":"0","service_name":"系统安装调试"})
    lighting=next(system for system in quote["systems"] if system["category_id"]=="lighting-control")
    assert lighting["subtotal_cents"]==12*299999 and lighting["service_calculated_cents"]==round(lighting["subtotal_cents"]*.1)
    quote=request("PATCH",f"/api/projects/{project['id']}/quotation-v2",json={"tax_rate":"6","discount_type":"fixed","discount_value":"100.00","show_empty_categories":True})
    assert quote["settings"]["show_empty_categories"] is True and quote["totals"]["discount_cents"]==10000 and quote["totals"]["tax_cents"]>0
    quote=request("POST",f"/api/projects/{project['id']}/quotation-v2/items",json={"system_category":"network","product_name":"临时交换机","product_code":"TEMP-24","brand":"临时品牌","origin":"中国","quantity":2,"unit_price":"999.50","unit":"台","params":"24口","remark":"项目临时增加"})
    network=next(system for system in quote["systems"] if system["category_id"]=="network")
    assert network["items"][0]["source_type"]=="temp" and network["items"][0]["item_total_cents"]==199900
    quote=request("POST",f"/api/projects/{project['id']}/quotation-v2/versions",json={"note":"客户确认前版本"})
    assert quote["version"]==2 and quote["versions"][0]["note"]=="客户确认前版本"


def test_v33_quotation_order_replace_cross_page_reorder_and_construction_fees():
    from backend.quotation_v2 import SYSTEM_CATEGORIES

    expected_order=["equipment-cabinet","central-control","touch-control","curtain","hvac","cable","intercom","network","surveillance","music","lighting-control"]
    assert [category["id"] for category in SYSTEM_CATEGORIES]==expected_order
    project=request("POST","/api/projects",json={"name":"v33报价回归"})
    pid=project["id"]
    quote=request("GET",f"/api/projects/{pid}/quotation-v2")
    assert quote["construction_fees"] and all(fee["amount_cents"]==0 for fee in quote["construction_fees"])
    assert [system["category_id"] for system in quote["systems"]]==expected_order+["uncategorized"]

    existing_target=request("POST",f"/api/projects/{pid}/quotation-v2/items",json={"system_category":"central-control","product_name":"保留顺序的中控产品","quantity":1,"unit_price":"80.00","unit":"台"})
    preserved_id=next(system for system in existing_target["systems"] if system["category_id"]=="central-control")["items"][0]["id"]
    old=request("POST",f"/api/projects/{pid}/quotation-v2/items",json={"system_category":"network","product_name":"旧报价产品","quantity":3,"unit_price":"100.00","unit":"台"})
    item=next(system for system in old["systems"] if system["category_id"]=="network")["items"][0]
    replacement=request("POST","/api/products",json={"brand_name":"替换品牌","name":"新产品","model":"NEW-1","category":"智能","device_type":"other","system_category":"central-control","standard_price":"25.50","unit":"套","params":"新参数"})
    changed=request("POST",f"/api/quotation-items/{item['id']}/replace",json={"product_id":replacement["id"]})
    new_item=next(row for system in changed["systems"] for row in system["items"] if row["id"]==item["id"])
    assert (new_item["product_id"],new_item["product_name"],new_item["product_code"],new_item["quantity"],new_item["unit_price_cents"],new_item["unit"],new_item["params"]) == (replacement["id"],"新产品","NEW-1",1,2550,"套","新参数")
    assert new_item["system_category"]=="central-control"
    central_items=next(system for system in changed["systems"] if system["category_id"]=="central-control")["items"]
    assert [item["id"] for item in central_items]==[preserved_id,item["id"]]

    # Reordering submits all rows in the system, including rows that will be split across pages.
    ids=[]
    for index in range(8):
        snapshot=request("POST",f"/api/projects/{pid}/quotation-v2/items",json={"system_category":"network","product_name":f"跨页产品{index}","quantity":1,"unit_price":"1.00"})
        ids=[item["id"] for item in next(system for system in snapshot["systems"] if system["category_id"]=="network")["items"]]
    reordered=request("POST",f"/api/projects/{pid}/quotation-v2/reorder",json={"system_category":"network","item_ids":list(reversed(ids))})
    network=next(system for system in reordered["systems"] if system["category_id"]=="network")
    assert [item["id"] for item in network["items"]]==list(reversed(ids))
    assert [item["id"] for item in request("GET",f"/api/projects/{pid}/quotation-v2")["systems"][7]["items"]]==list(reversed(ids))
    invalid=client.post(f"/api/projects/{pid}/quotation-v2/reorder",json={"system_category":"network","item_ids":[new_item["id"],ids[0]]})
    assert invalid.status_code==422

    fee_result=request("PATCH",f"/api/projects/{pid}/quotation-v2/construction-fees",json={"fees":{"strong_power":"20.00","weak_power":"","equipment_installation":"0","line_material":"0","system_training":"0"}})
    quote=request("PATCH",f"/api/projects/{pid}/quotation-v2",json={"tax_rate":"10"})
    totals=quote["totals"]
    assert totals["construction_fee_total_cents"]==2000
    assert totals["tax_cents"]==round(totals["systems_total_cents"]*.1)
    assert totals["grand_total_cents"]==totals["systems_total_cents"]+totals["tax_cents"]+2000
    reopened=request("GET",f"/api/projects/{pid}/quotation-v2")
    assert reopened["construction_fees"][0]["amount_cents"]==2000


def test_z_v33_existing_32_quote_migrates_without_loss_when_fee_table_is_absent():
    from backend.app import init_db

    project=request("POST","/api/projects",json={"name":"v32旧报价兼容测试"})
    original=request("POST",f"/api/projects/{project['id']}/quotation-v2/items",json={"system_category":"network","product_name":"旧系统产品","quantity":2,"unit_price":"123.45"})
    old_network=next(system for system in original["systems"] if system["category_id"]=="network")
    old_item_id=old_network["items"][0]["id"]
    with sqlite3.connect(DB) as db:
        db.execute("DROP TABLE quotation_construction_fees")
    init_db()
    migrated=request("GET",f"/api/projects/{project['id']}/quotation-v2")
    migrated_item=next(item for system in migrated["systems"] for item in system["items"] if item["id"]==old_item_id)
    assert migrated_item["product_name"]=="旧系统产品" and migrated_item["quantity"]==2
    assert all(fee["amount_cents"]==0 for fee in migrated["construction_fees"])


def test_product_system_category_markdown_and_batch():
    product=request("POST","/api/products",json={"brand_name":"系统分类品牌","name":"触控面板","category":"智能","system_category":"touch-control","standard_price":"88.00"})
    assert Path(product["note_path"]).parent.name == "触控系统"
    assert any(p["id"] == product["id"] for p in request("GET", "/api/products?include_disabled=true"))
    note=(Path(TEST_ROOT)/product["note_path"]).read_text(encoding="utf-8")
    assert "所属系统: touch-control" in note and "| 所属系统 | 触控系统 |" in note
    request("PATCH","/api/products/batch",json={"ids":[product["id"]],"action":"system_category","system_category":"central-control"})
    changed=request("GET",f"/api/products/{product['id']}")
    assert changed["system_category"]=="central-control"
    assert Path(changed["note_path"]).parent.name == "中控系统"
    assert not (Path(TEST_ROOT)/product["note_path"]).exists()
    listed=next(p for p in request("GET", "/api/products?include_disabled=true") if p["id"] == product["id"])
    assert listed["system_category"] == "central-control" and listed["note_path"] == changed["note_path"]
    # Obsidian-side YAML edits become visible on explicit sync, with a matching folder.
    current_note=Path(TEST_ROOT)/changed["note_path"]
    current_note.write_text(current_note.read_text(encoding="utf-8").replace("所属系统: central-control", "所属系统: curtain"),encoding="utf-8")
    synced=next(p for p in request("GET", "/api/products?include_disabled=true&refresh=true") if p["id"] == product["id"])
    assert synced["system_category"]=="curtain"
    assert Path(synced["note_path"]).parent.name=="电动窗帘系统"
    assert not current_note.exists()


def test_product_layout_migration_preserves_note_and_uncategorized(tmp_path):
    product_dir=tmp_path/"01_产品库"
    legacy=product_dir/"智能"
    legacy.mkdir(parents=True)
    product_id="b"*32
    note=legacy/"legacy.md"
    original=f"---\nid: {product_id}\n类型: 报价产品\n产品名称: 原始产品\n品牌: 旧品牌\n分类: 智能\n点位类型: system\n所属系统: uncategorized\n售价_分: 100\n---\n用户笔记正文\n"
    note.write_text(original,encoding="utf-8")
    repo=ObsidianProductRepository()
    with patch.object(product_repository_module,"PRODUCT_DIR",product_dir),patch.object(product_repository_module,"VAULT_ROOT",tmp_path):
        assert repo.normalize_layout()==1
        target=product_dir/"未分类"/"legacy.md"
        assert target.read_text(encoding="utf-8")==original
        assert repo.normalize_layout()==0
        assert repo.list(include_disabled=True)[0]["note_path"]=="01_产品库/未分类/legacy.md"


def test_product_index_prefers_system_folder_over_recreated_legacy_copy(tmp_path):
    product_dir=tmp_path/"01_产品库"
    canonical=product_dir/"中控系统"/"product.md"
    legacy=product_dir/"智能"/"product.md"
    canonical.parent.mkdir(parents=True)
    legacy.parent.mkdir(parents=True)
    note="---\nid: " + "c"*32 + "\n类型: 报价产品\n产品名称: 同步测试\n品牌: 测试品牌\n分类: 智能\n点位类型: system\n所属系统: central-control\n售价_分: 100\n---\n"
    canonical.write_text(note,encoding="utf-8")
    legacy.write_text(note.replace("售价_分: 100","售价_分: 200"),encoding="utf-8")
    repo=ObsidianProductRepository()
    with patch.object(product_repository_module,"PRODUCT_DIR",product_dir),patch.object(product_repository_module,"VAULT_ROOT",tmp_path):
        product=repo.list(include_disabled=True,refresh=True)[0]
        assert product["standard_cents"]==100
        assert product["note_path"]=="01_产品库/中控系统/product.md"


def test_rq003a_switch_points_bind_product_during_creation_and_reload():
    switch_product=request("POST","/api/products",json={
        "brand_name":"RQ003A品牌","name":"三键智能开关","model":"SW-3",
        "category":"智能","device_type":"switch","system_category":"lighting-control",
        "standard_price":"368.00","unit":"个"
    })
    light_product=request("POST","/api/products",json={
        "brand_name":"RQ003A品牌","name":"调光模块","category":"灯光","device_type":"light",
        "standard_price":"500.00"
    })
    project=request("POST","/api/projects",json={"name":"RQ-003A 开关布点测试"})
    image=io.BytesIO(); Image.new("RGB",(600,400),"white").save(image,"PNG")
    drawing=request("POST",f"/api/projects/{project['id']}/drawings",files={"file":("开关测试.png",image.getvalue(),"image/png")})

    created=[]
    for x in (.25,.55):
        created.append(request("POST",f"/api/projects/{project['id']}/points?drawing_id={drawing['id']}",json={
            "type_id":"switch","product_id":switch_product["id"],"x":x,"y":.4
        }))
    assert all(point["product_id"]==switch_product["id"] for point in created)
    assert all(point["product_snapshot"]["id"]==switch_product["id"] for point in created)
    assert all(point["standard_price"]=="368.00" and point["actual_price"]=="368.00" for point in created)

    moved=request("PATCH",f"/api/points/{created[0]['id']}",json={"x":.35,"y":.6})
    assert moved["product_id"]==switch_product["id"] and (moved["x"],moved["y"])==(.35,.6)
    reopened=request("GET",f"/api/projects/{project['id']}")
    reopened_points={point["id"]:point for point in reopened["points"]}
    assert reopened_points[created[0]["id"]]["product_snapshot"]["id"]==switch_product["id"]
    assert reopened_points[created[1]["id"]]["product_id"]==switch_product["id"]

    mismatch=client.post(f"/api/projects/{project['id']}/points?drawing_id={drawing['id']}",json={
        "type_id":"switch","product_id":light_product["id"],"x":.75,"y":.4
    })
    assert mismatch.status_code==422

    legacy=request("POST",f"/api/projects/{project['id']}/points?drawing_id={drawing['id']}",json={"type_id":"switch","x":.8,"y":.7})
    assert legacy["product_id"] is None
    request("DELETE",f"/api/points/{created[0]['id']}")
    reopened=request("GET",f"/api/projects/{project['id']}")
    assert all(point["id"]!=created[0]["id"] for point in reopened["points"])
    assert any(point["id"]==created[1]["id"] and point["product_id"]==switch_product["id"] for point in reopened["points"])


def test_rq006_brand_library_logo_product_and_quotation_link():
    brand=request("POST","/api/brands",json={
        "name":"RQ006 品牌","english_name":"RQ006 Brand","aliases":["六号牌","RQ6"],
        "enabled":True,"note":"品牌库专项测试"
    })
    assert brand["english_name"]=="RQ006 Brand" and brand["aliases"]==["六号牌","RQ6"]
    assert brand["logo_url"] is None

    png=io.BytesIO(); Image.new("RGBA",(120,40),(0,90,180,0)).save(png,"PNG")
    brand=request("POST",f"/api/brands/{brand['id']}/logo",files={"file":("brand.png",png.getvalue(),"image/png")})
    assert brand["logo_url"].endswith(".png")
    svg=b'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 120 40"><rect width="120" height="40" fill="#1262a3"/></svg>'
    brand=request("POST",f"/api/brands/{brand['id']}/logo",files={"file":("brand.svg",svg,"image/svg+xml")})
    assert brand["logo_url"].endswith(".svg")  # SVG 始终优先于已存在的 PNG
    logo=client.get(brand["logo_url"])
    assert logo.status_code==200 and logo.headers["content-type"].startswith("image/svg+xml")

    product=request("POST","/api/products",json={
        "brand_id":brand["id"],"name":"品牌关联产品","category":"智能","device_type":"switch",
        "system_category":"lighting-control","standard_price":"188.00"
    })
    assert product["brand_id"]==brand["id"] and product["brand_logo_url"].endswith(".svg")
    note=(Path(TEST_ROOT)/product["note_path"]).read_text(encoding="utf-8")
    assert f"品牌ID: {brand['id']}" in note

    project=request("POST","/api/projects",json={"name":"RQ006 报价测试"})
    quote=request("POST",f"/api/projects/{project['id']}/quotation-v2/items",json={
        "system_category":"lighting-control","product_id":product["id"],"product_name":product["name"],
        "brand":product["brand_name"],"quantity":1,"unit_price":"188.00","unit":"个"
    })
    item=next(system for system in quote["systems"] if system["category_id"]=="lighting-control")["items"][0]
    assert item["brand_id"]==brand["id"] and item["brand_logo_url"].endswith(".svg")

    updated=request("PUT",f"/api/brands/{brand['id']}",json={
        "name":"RQ006 品牌新名","english_name":"RQ006 Brand","aliases":["旧名称"],"enabled":False,"note":"已停用但历史关联保留"
    })
    assert updated["id"]==brand["id"] and updated["enabled"] is False
    reloaded=request("GET",f"/api/products/{product['id']}")
    assert reloaded["brand_id"]==brand["id"] and reloaded["brand_name"]=="RQ006 品牌新名" and reloaded["brand_logo_url"].endswith(".svg")


def test_rq007_project_status_uses_markdown_documents():
    response=client.get("/api/project-status")
    assert response.status_code==200
    assert response.headers["cache-control"]=="no-store"
    data=response.json()
    assert data["source"]=={
        "requirements":"docs/REQUIREMENTS.md",
        "issues":"docs/ISSUES.md",
        "changelog":"docs/CHANGELOG.md",
    }
    assert [status["label"] for status in data["status_definitions"]]==["已解决","开发中","待优化","待开发","有问题"]
    assert sum(data["counts"]["requirements"].values())==len(data["requirements"])
    assert sum(data["counts"]["issues"].values())==len(data["issues"])
    assert next(item for item in data["requirements"] if item["id"]=="RQ-007")["status"]["key"]=="resolved"
    assert next(item for item in data["issues"] if item["id"]=="ISSUE-006")["status"]["key"]=="resolved"
    assert any(item["title"]=="RQ-007 问题与版本管理体系" for item in data["changelog"])


def test_local_markdown_is_authoritative_and_rebuilds_runtime_cache_without_sync():
    brand=request("POST","/api/brands",json={"name":"RQ012知识库品牌","english_name":"Vault Brand","aliases":["知识牌"],"note":"双向同步测试"})
    product=request("POST","/api/products",json={
        "brand_id":brand["id"],"name":"RQ012知识库产品","model":"KB-12","category":"智能",
        "device_type":"switch","system_category":"lighting-control","standard_price":"1200.00","cost_price":"800.00"
    })
    project=request("POST","/api/projects",json={"name":"RQ012知识库项目","customer":"数据恢复客户","address":"知识库","status":"进行中"})
    image=io.BytesIO(); Image.new("RGB",(500,300),"white").save(image,"PNG")
    drawing=request("POST",f"/api/projects/{project['id']}/drawings",files={"file":("知识库图纸.png",image.getvalue(),"image/png")})
    point=request("POST",f"/api/projects/{project['id']}/points?drawing_id={drawing['id']}",json={"type_id":"switch","product_id":product["id"],"x":.31,"y":.42})
    request("POST", "/api/knowledge-sync/retry")

    brand_note=knowledge_repository.brand_path(brand["id"])
    project_note=knowledge_repository.project_path(project["id"])
    assert brand_note and brand_note.is_file() and "Vault Brand" in brand_note.read_text(encoding="utf-8")
    assert project_note and project_note.is_file() and point["id"] in project_note.read_text(encoding="utf-8")

    # Retry is retained as a compatibility endpoint but reports local mode and never syncs.
    retried=request("POST", "/api/knowledge-sync/retry")
    restored=request("GET",f"/api/projects/{project['id']}")
    assert restored["name"]=="RQ012知识库项目"
    assert any(item["id"]==point["id"] and item["product_id"]==product["id"] for item in restored["points"])
    assert retried["sync_status"] == "local"

    status=request("GET","/api/knowledge-sync/status")
    assert status["brand_notes"]>=1 and status["project_notes"]>=1 and status["conflicts"]==[]


def test_rq013_legend_position_persists_to_vault_and_exports():
    project=request("POST","/api/projects",json={"name":"RQ013图例测试"})
    image=io.BytesIO(); Image.new("RGB",(900,600),"white").save(image,"PNG")
    drawing=request("POST",f"/api/projects/{project['id']}/drawings",files={"file":("图例.png",image.getvalue(),"image/png")})
    assert drawing["legend_x"]==.025 and drawing["legend_y"]==.025
    request("POST",f"/api/projects/{project['id']}/points?drawing_id={drawing['id']}",json={"type_id":"sensor","x":.3,"y":.35})
    request("POST",f"/api/projects/{project['id']}/points?drawing_id={drawing['id']}",json={"type_id":"sensor","x":.55,"y":.5})
    curtain=request("POST",f"/api/projects/{project['id']}/points?drawing_id={drawing['id']}",json={"type_id":"curtain","x":.2,"y":.75,"x2":.65,"y2":.75})

    updated=request("PATCH",f"/api/projects/{project['id']}/drawings/{drawing['id']}",json={"legend_x":.58,"legend_y":.12})
    saved=updated["drawings"][0]
    assert abs(saved["legend_x"]-.58)<1e-9 and abs(saved["legend_y"]-.12)<1e-9
    request("POST", "/api/knowledge-sync/retry")
    note=knowledge_repository.project_path(project["id"])
    assert note and '"legend_x": 0.58' in note.read_text(encoding="utf-8")
    history=DATA/"06_客户档案"/"智能图纸与报价"/"_备份"/"项目历史"/project["id"]
    assert history.is_dir() and list(history.glob("*.md"))

    png=client.get(f"/api/projects/{project['id']}/drawings/{drawing['id']}/export.png?show_name=false&show_number=false")
    pdf=client.get(f"/api/projects/{project['id']}/drawings/{drawing['id']}/export.pdf?show_name=false&show_number=false")
    assert png.status_code==200 and png.headers["content-type"].startswith("image/png") and len(png.content)>10000
    assert pdf.status_code==200 and pdf.content.startswith(b"%PDF") and len(pdf.content)>10000

    request("DELETE",f"/api/points/{curtain['id']}")
    reopened=request("GET",f"/api/projects/{project['id']}")
    assert len([point for point in reopened["points"] if point["type_id"]=="sensor"])==2
    assert all(point["type_id"]!="curtain" for point in reopened["points"])
    assert abs(reopened["drawings"][0]["legend_x"]-.58)<1e-9

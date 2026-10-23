from __future__ import annotations

import csv
import io
import json
import math
import os
import re
import shutil
import sqlite3
import subprocess
import textwrap
import time
import uuid
import xml.etree.ElementTree as ET
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Optional, Union

import fitz
from fastapi import FastAPI, File, Form, HTTPException, Response, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from PIL import Image, ImageChops, ImageDraw, ImageFilter, ImageFont, ImageOps
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XImage
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, Field
from backend.path_resolver import get_paths

PATHS = get_paths(__file__)
if PATHS.local_data_mode:
    required_data_dirs = (
        "01_产品库", "01_产品库/_品牌", "06_客户档案/智能图纸与报价",
        "09_附件", "uploads",
    )
    missing_data_dirs = [name for name in required_data_dirs if not (PATHS.data_root / name).is_dir()]
    if missing_data_dirs:
        raise FileNotFoundError(
            "本地 data/ 知识库不完整，请将完整工具目录复制到本地；缺少：" + ", ".join(missing_data_dirs)
        )

from backend.product_repository import (
    BRAND_LOGO_DIR, IMAGE_DIR, PRODUCT_CATEGORIES, PRODUCT_DIR, VAULT_ROOT,
    brand_id as stable_brand_id, canonical_device_type, canonical_product_category, product_repository,
)
from backend.diagnostics import LOG_FILE, record_event, record_exception, reset_request_id, set_request_id
from backend.knowledge_repository import BRAND_DIR, PROJECT_DIR, KnowledgeRepository
from backend.pptx_export import export_quote_pptx
from backend.quotation_v2 import (
    SYSTEM_BY_ID, SYSTEM_CATEGORIES, SYSTEM_IDS, UNCATEGORIZED,
    calculate_quotation, canonical_system_category, normalize_quotation_text,
)

ROOT = PATHS.app_root
APP_VERSION = (ROOT / "VERSION").read_text(encoding="utf-8").strip() if (ROOT / "VERSION").is_file() else "development"
DATA = PATHS.data_root
UPLOADS = PATHS.uploads_root
EXPORTS = PATHS.exports_root
DB = DATA / "app.sqlite3"
knowledge_repository = KnowledgeRepository(DB)
for folder in (DATA, EXPORTS):
    folder.mkdir(parents=True, exist_ok=True)
try:
    UPLOADS.mkdir(parents=True, exist_ok=True)
except OSError:
    pass

TYPES = [
    ("switch", "智能开关", "开关", "#e51f4b"), ("curtain", "窗帘", "窗帘", "#2459b8"),
    ("sensor", "人体传感器", "传感器", "#079447"), ("thermostat", "温控器", "温控器", "#e59144"),
    ("light", "调光", "调光", "#e0a420"),
    ("non_smart_area", "非智能区域", "非智区", "#6f7786"),
    ("lock", "智能门禁", "门禁", "#555e75"),
    ("camera", "摄像头", "摄像头", "#dd6b65"),
    ("ceiling_ap", "吸顶 AP", "吸顶AP", "#2389b9"), ("wall_ap", "墙面 AP", "墙面AP", "#3971b7"),
    ("water", "水浸传感器", "水浸", "#008fb3"), ("av", "影音设备", "影音", "#ad6c9e"), ("system", "智能系统", "系统", "#293f71"),
    ("other", "其他设备", "其他", "#808080"),
    ("network", "网络设备（历史兼容）", "网络", "#2c9bce"),
]
TYPE_MAP = {v[0]: v for v in TYPES}
NON_PRODUCT_TYPES = {"non_smart_area"}
LEGACY_TYPES = {"network"}


def point_shape(type_id: str) -> str:
    if type_id == "curtain":
        return "line"
    if type_id == "camera":
        return "sector"
    if type_id in {"light", "non_smart_area"}:
        return "area"
    return "point"


def canonical_category(value: str) -> str:
    return canonical_product_category(value)


def category_type(value: str) -> str:
    return canonical_device_type(value, canonical_category(value))
def find_chinese_font() -> Path | None:
    configured = os.environ.get("POINT_QUOTE_FONT_PATH")
    windows_dir = os.environ.get("WINDIR")
    windows_fonts = Path(windows_dir) / "Fonts" if windows_dir else None
    candidates = [
        Path(configured) if configured else None,
        ROOT / "assets" / "fonts" / "NotoSansSC-Regular.ttf",
        windows_fonts / "msyh.ttc" if windows_fonts else None,
        windows_fonts / "msyh.ttf" if windows_fonts else None,
        windows_fonts / "simhei.ttf" if windows_fonts else None,
        windows_fonts / "simsun.ttc" if windows_fonts else None,
        Path("/System/Library/Fonts/STHeiti Medium.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"),
    ]
    return next((path for path in candidates if path and path.is_file()), None)


FONT_PATH = find_chinese_font()


def now():
    return datetime.now().astimezone().isoformat(timespec="seconds")


def uid():
    return uuid.uuid4().hex


def money(value):
    try:
        d = Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    except (InvalidOperation, TypeError):
        raise HTTPException(422, "金额格式错误")
    if d < 0 or d > Decimal("999999999.99"):
        raise HTTPException(422, "金额超出允许范围")
    return int(d * 100)


def connect():
    db = sqlite3.connect(DB)
    db.row_factory = sqlite3.Row
    db.execute("PRAGMA foreign_keys=ON")
    # SMB/NAS volumes often do not support SQLite shared-memory WAL files.
    db.execute("PRAGMA journal_mode=DELETE")
    return db


def rows(cursor):
    return [dict(r) for r in cursor.fetchall()]


def one(db, sql, args=()):
    r = db.execute(sql, args).fetchone()
    if not r:
        raise HTTPException(404, "记录不存在")
    return dict(r)


def init_db():
    with connect() as db:
        db.executescript("""
        CREATE TABLE IF NOT EXISTS device_types(id TEXT PRIMARY KEY,name TEXT NOT NULL,prefix TEXT NOT NULL,color TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS projects(id TEXT PRIMARY KEY,name TEXT NOT NULL,customer TEXT NOT NULL DEFAULT '',address TEXT NOT NULL DEFAULT '',status TEXT NOT NULL DEFAULT '进行中',signed INTEGER NOT NULL DEFAULT 0,created_at TEXT NOT NULL,updated_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS project_folders(id TEXT PRIMARY KEY,name TEXT NOT NULL UNIQUE,sort_order INTEGER NOT NULL DEFAULT 0,created_at TEXT NOT NULL,updated_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS drawings(id TEXT PRIMARY KEY,project_id TEXT NOT NULL REFERENCES projects(id) ON DELETE CASCADE,source_path TEXT NOT NULL,background_path TEXT NOT NULL,filename TEXT NOT NULL,page INTEGER NOT NULL DEFAULT 1,width INTEGER NOT NULL,height INTEGER NOT NULL,created_at TEXT NOT NULL,rotation INTEGER NOT NULL DEFAULT 0,legend_x REAL NOT NULL DEFAULT 0.025,legend_y REAL NOT NULL DEFAULT 0.025);
        CREATE TABLE IF NOT EXISTS brands(id TEXT PRIMARY KEY,name TEXT NOT NULL UNIQUE,created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS products(id TEXT PRIMARY KEY,brand_id TEXT NOT NULL REFERENCES brands(id),name TEXT NOT NULL,series TEXT NOT NULL DEFAULT '',model TEXT NOT NULL DEFAULT '',category TEXT NOT NULL REFERENCES device_types(id),standard_cents INTEGER NOT NULL,unit TEXT NOT NULL DEFAULT '个',params TEXT NOT NULL DEFAULT '',description TEXT NOT NULL DEFAULT '',image_path TEXT,enabled INTEGER NOT NULL DEFAULT 1,updated_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS points(id TEXT PRIMARY KEY,project_id TEXT NOT NULL REFERENCES projects(id) ON DELETE CASCADE,drawing_id TEXT NOT NULL REFERENCES drawings(id) ON DELETE CASCADE,type_id TEXT NOT NULL REFERENCES device_types(id),number TEXT NOT NULL,name TEXT NOT NULL,room TEXT NOT NULL DEFAULT '',quantity INTEGER NOT NULL DEFAULT 1,note TEXT NOT NULL DEFAULT '',x REAL NOT NULL,y REAL NOT NULL,size INTEGER NOT NULL DEFAULT 28,color TEXT NOT NULL,locked INTEGER NOT NULL DEFAULT 0,product_id TEXT REFERENCES products(id) ON DELETE SET NULL,product_snapshot TEXT,standard_cents INTEGER,actual_cents INTEGER,created_at TEXT NOT NULL,updated_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS quotes(id TEXT PRIMARY KEY,project_id TEXT NOT NULL UNIQUE REFERENCES projects(id) ON DELETE CASCADE,discount_cents INTEGER NOT NULL DEFAULT 0,note TEXT NOT NULL DEFAULT '',updated_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS extra_fees(id TEXT PRIMARY KEY,quote_id TEXT NOT NULL REFERENCES quotes(id) ON DELETE CASCADE,name TEXT NOT NULL,amount_cents INTEGER NOT NULL,kind TEXT NOT NULL DEFAULT '额外费用');
        CREATE TABLE IF NOT EXISTS quote_lines(id TEXT PRIMARY KEY,quote_id TEXT NOT NULL REFERENCES quotes(id) ON DELETE CASCADE,point_id TEXT NOT NULL REFERENCES points(id) ON DELETE CASCADE,quantity INTEGER NOT NULL,actual_cents INTEGER NOT NULL,subtotal_cents INTEGER NOT NULL,snapshot TEXT NOT NULL,updated_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS manual_quote_lines(id TEXT PRIMARY KEY,quote_id TEXT NOT NULL REFERENCES quotes(id) ON DELETE CASCADE,product_id TEXT NOT NULL,quantity INTEGER NOT NULL,standard_cents INTEGER NOT NULL,actual_cents INTEGER NOT NULL,subtotal_cents INTEGER NOT NULL,snapshot TEXT NOT NULL,note TEXT NOT NULL DEFAULT '',created_at TEXT NOT NULL,updated_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS quotation_systems(id TEXT PRIMARY KEY,quote_id TEXT NOT NULL REFERENCES quotes(id) ON DELETE CASCADE,category_id TEXT NOT NULL,sort_order INTEGER NOT NULL,service_type TEXT NOT NULL DEFAULT 'none',service_rate_basis INTEGER NOT NULL DEFAULT 0,service_amount_cents INTEGER NOT NULL DEFAULT 0,service_name TEXT NOT NULL DEFAULT '（服务）硬件安装、设备编程、场景调试、售后维保',created_at TEXT NOT NULL,updated_at TEXT NOT NULL,UNIQUE(quote_id,category_id));
        CREATE TABLE IF NOT EXISTS quotation_items(id TEXT PRIMARY KEY,quote_id TEXT NOT NULL REFERENCES quotes(id) ON DELETE CASCADE,source_type TEXT NOT NULL,source_key TEXT,product_id TEXT,system_category TEXT NOT NULL DEFAULT 'uncategorized',product_name TEXT NOT NULL,product_code TEXT NOT NULL DEFAULT '',image_path TEXT,brand TEXT NOT NULL DEFAULT '',origin TEXT NOT NULL DEFAULT '',quantity INTEGER NOT NULL DEFAULT 1,source_quantity INTEGER NOT NULL DEFAULT 0,quantity_override INTEGER NOT NULL DEFAULT 0,unit_price_cents INTEGER NOT NULL DEFAULT 0,standard_cents INTEGER NOT NULL DEFAULT 0,unit TEXT NOT NULL DEFAULT '个',params TEXT NOT NULL DEFAULT '',remark TEXT NOT NULL DEFAULT '',sort_order INTEGER NOT NULL DEFAULT 0,visible INTEGER NOT NULL DEFAULT 1,created_at TEXT NOT NULL,updated_at TEXT NOT NULL,UNIQUE(quote_id,source_type,source_key));
        CREATE TABLE IF NOT EXISTS quotation_construction_fees(quote_id TEXT NOT NULL REFERENCES quotes(id) ON DELETE CASCADE,fee_key TEXT NOT NULL,amount_cents INTEGER NOT NULL DEFAULT 0,updated_at TEXT NOT NULL,PRIMARY KEY(quote_id,fee_key));
        CREATE TABLE IF NOT EXISTS quotation_versions(id TEXT PRIMARY KEY,quote_id TEXT NOT NULL REFERENCES quotes(id) ON DELETE CASCADE,version INTEGER NOT NULL,total_cents INTEGER NOT NULL,snapshot_json TEXT NOT NULL,created_by TEXT NOT NULL DEFAULT '本机用户',note TEXT NOT NULL DEFAULT '',created_at TEXT NOT NULL,UNIQUE(quote_id,version));
        CREATE TABLE IF NOT EXISTS point_counters(project_id TEXT NOT NULL REFERENCES projects(id) ON DELETE CASCADE,type_id TEXT NOT NULL REFERENCES device_types(id),last_number INTEGER NOT NULL,PRIMARY KEY(project_id,type_id));
        CREATE INDEX IF NOT EXISTS idx_points_project ON points(project_id);
        CREATE INDEX IF NOT EXISTS idx_products_filters ON products(category,brand_id,series);
        CREATE INDEX IF NOT EXISTS idx_quotation_items_quote ON quotation_items(quote_id,system_category,sort_order);
        CREATE TABLE IF NOT EXISTS feedback(id TEXT PRIMARY KEY,title TEXT NOT NULL,description TEXT NOT NULL,module TEXT NOT NULL DEFAULT '',screenshot_path TEXT,created_at TEXT NOT NULL,status TEXT NOT NULL DEFAULT '待处理');
        """)
        columns = {r[1] for r in db.execute("PRAGMA table_info(drawings)")}
        if "name" not in columns:
            db.execute("ALTER TABLE drawings ADD COLUMN name TEXT NOT NULL DEFAULT ''")
        if "rotation" not in columns:
            db.execute("ALTER TABLE drawings ADD COLUMN rotation INTEGER NOT NULL DEFAULT 0")
        if "legend_x" not in columns:
            db.execute("ALTER TABLE drawings ADD COLUMN legend_x REAL NOT NULL DEFAULT 0.025")
        if "legend_y" not in columns:
            db.execute("ALTER TABLE drawings ADD COLUMN legend_y REAL NOT NULL DEFAULT 0.025")
        if "floor_id" not in columns:
            db.execute("ALTER TABLE drawings ADD COLUMN floor_id TEXT")
        if "legend_items" not in columns:
            db.execute("ALTER TABLE drawings ADD COLUMN legend_items TEXT NOT NULL DEFAULT '{}'")
        project_columns = {r[1] for r in db.execute("PRAGMA table_info(projects)")}
        if "folder_id" not in project_columns:
            db.execute("ALTER TABLE projects ADD COLUMN folder_id TEXT")
        if "signed" not in project_columns:
            db.execute("ALTER TABLE projects ADD COLUMN signed INTEGER NOT NULL DEFAULT 0")
        if "shadow_strength" not in project_columns:
            db.execute("ALTER TABLE projects ADD COLUMN shadow_strength INTEGER NOT NULL DEFAULT 25")
        db.execute("CREATE TABLE IF NOT EXISTS floors(id TEXT PRIMARY KEY,project_id TEXT NOT NULL REFERENCES projects(id) ON DELETE CASCADE,name TEXT NOT NULL,sort_order INTEGER NOT NULL DEFAULT 0)")
        db.execute("CREATE TABLE IF NOT EXISTS layers(id TEXT PRIMARY KEY,floor_id TEXT NOT NULL REFERENCES floors(id) ON DELETE CASCADE,name TEXT NOT NULL,sort_order INTEGER NOT NULL DEFAULT 0,visible INTEGER NOT NULL DEFAULT 1,locked INTEGER NOT NULL DEFAULT 0)")
        floor_columns={r[1] for r in db.execute("PRAGMA table_info(floors)")}
        if "visibility_initialized" not in floor_columns:
            db.execute("ALTER TABLE floors ADD COLUMN visibility_initialized INTEGER NOT NULL DEFAULT 0")
        db.execute("CREATE INDEX IF NOT EXISTS idx_floors_project ON floors(project_id,sort_order)")
        db.execute("CREATE INDEX IF NOT EXISTS idx_layers_floor ON layers(floor_id,sort_order)")
        brand_columns = {r[1] for r in db.execute("PRAGMA table_info(brands)")}
        for column, definition in (
            ("english_name", "TEXT NOT NULL DEFAULT ''"), ("aliases", "TEXT NOT NULL DEFAULT '[]'"),
            ("logo_svg_path", "TEXT"), ("logo_png_path", "TEXT"),
            ("enabled", "INTEGER NOT NULL DEFAULT 1"), ("note", "TEXT NOT NULL DEFAULT ''"),
            ("updated_at", "TEXT NOT NULL DEFAULT ''"),
        ):
            if column not in brand_columns:
                db.execute(f"ALTER TABLE brands ADD COLUMN {column} {definition}")
        db.execute("UPDATE brands SET updated_at=created_at WHERE updated_at='' OR updated_at IS NULL")
        point_columns = {r[1] for r in db.execute("PRAGMA table_info(points)")}
        for column, definition in (("shape", "TEXT NOT NULL DEFAULT 'point'"), ("x2", "REAL"), ("y2", "REAL"), ("show_motor", "INTEGER NOT NULL DEFAULT 1"), ("fov_angle", "REAL NOT NULL DEFAULT 60"), ("layer_id", "TEXT"), ("line_width", "REAL NOT NULL DEFAULT 3"), ("curtain_style", "TEXT NOT NULL DEFAULT 'single'"), ("curtain_length", "REAL NOT NULL DEFAULT 0")):
            if column not in point_columns:
                db.execute(f"ALTER TABLE points ADD COLUMN {column} {definition}")
        drawing_columns = {r[1] for r in db.execute("PRAGMA table_info(drawings)")}
        for column, definition in (
            ("canvas_mode", "TEXT NOT NULL DEFAULT 'legacy'"),
            ("canvas_width", "INTEGER NOT NULL DEFAULT 0"),
            ("canvas_height", "INTEGER NOT NULL DEFAULT 0"),
            ("image_x", "REAL NOT NULL DEFAULT 0"),
            ("image_y", "REAL NOT NULL DEFAULT 0"),
            ("image_width", "REAL NOT NULL DEFAULT 0"),
            ("image_height", "REAL NOT NULL DEFAULT 0"),
            ("image_rotation", "REAL NOT NULL DEFAULT 0"),
            ("guide_lines", "TEXT NOT NULL DEFAULT '[]'"),
            ("is_blank", "INTEGER NOT NULL DEFAULT 0"),
            ("crop_source_path", "TEXT"),
            ("crop_source_width", "INTEGER"),
            ("crop_source_height", "INTEGER"),
            ("crop_frame_x", "REAL"),
            ("crop_frame_y", "REAL"),
            ("crop_frame_width", "REAL"),
            ("crop_frame_height", "REAL"),
            ("crop_left", "REAL"),
            ("crop_top", "REAL"),
            ("crop_right", "REAL"),
            ("crop_bottom", "REAL"),
        ):
            if column not in drawing_columns:
                db.execute(f"ALTER TABLE drawings ADD COLUMN {column} {definition}")
        for project in db.execute("SELECT id FROM projects").fetchall():
            floor = db.execute("SELECT id FROM floors WHERE project_id=? ORDER BY sort_order LIMIT 1", (project[0],)).fetchone()
            if not floor:
                floor_id=uid()
                db.execute("INSERT INTO floors(id,project_id,name,sort_order,visibility_initialized) VALUES(?,?,?,0,1)",(floor_id,project[0],"默认楼层"))
            else: floor_id=floor[0]
            layer = db.execute("SELECT id FROM layers WHERE floor_id=? ORDER BY sort_order LIMIT 1",(floor_id,)).fetchone()
            if not layer:
                layer_id=uid()
                db.execute("INSERT INTO layers(id,floor_id,name,sort_order) VALUES(?,?,?,0)",(layer_id,floor_id,"默认图层"))
            else: layer_id=layer[0]
            db.execute("UPDATE drawings SET floor_id=? WHERE project_id=? AND floor_id IS NULL",(floor_id,project[0]))
            db.execute("UPDATE points SET layer_id=? WHERE project_id=? AND layer_id IS NULL",(layer_id,project[0]))
        # Floors with no plan get a persistent blank 16:9 canvas. A legacy floor
        # with exactly one drawing can be upgraded without changing point data;
        # floors with multiple drawings retain their old page-per-drawing model.
        for floor_row in db.execute("SELECT id,project_id FROM floors").fetchall():
            floor_id,project_id=floor_row
            initialized=db.execute("SELECT visibility_initialized FROM floors WHERE id=?",(floor_id,)).fetchone()[0]
            if not initialized:
                first_layer=db.execute("SELECT id FROM layers WHERE floor_id=? ORDER BY sort_order,id LIMIT 1",(floor_id,)).fetchone()
                if first_layer:
                    db.execute("UPDATE layers SET visible=CASE WHEN id=? THEN 1 ELSE 0 END WHERE floor_id=?",(first_layer[0],floor_id))
                db.execute("UPDATE floors SET visibility_initialized=1 WHERE id=?",(floor_id,))
            floor_drawings=rows(db.execute("SELECT * FROM drawings WHERE floor_id=? ORDER BY created_at,id",(floor_id,)))
            if not floor_drawings:
                drawing_id=uid();blank_path=UPLOADS/f"{drawing_id}_blank.png"
                Image.new("RGB",(1600,900),"white").save(blank_path,"PNG")
                db.execute("INSERT INTO drawings(id,project_id,floor_id,source_path,background_path,filename,page,width,height,created_at,name,canvas_mode,canvas_width,canvas_height,image_x,image_y,image_width,image_height,image_rotation,guide_lines,is_blank) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",(drawing_id,project_id,floor_id,blank_path.name,blank_path.name,"未导入主图纸.png",1,1600,900,now(),"主图纸","fixed",1600,900,0,0,1600,900,0,"[]",1))
            elif len(floor_drawings)==1 and floor_drawings[0].get("canvas_mode")=="legacy":
                drawing=floor_drawings[0];cw,ch=1600,900;fit_scale=min(cw/drawing["width"],ch/drawing["height"]);iw,ih=drawing["width"]*fit_scale,drawing["height"]*fit_scale
                db.execute("UPDATE drawings SET canvas_mode='fixed',canvas_width=?,canvas_height=?,image_x=?,image_y=?,image_width=?,image_height=?,is_blank=0 WHERE id=?",(cw,ch,(cw-iw)/2,(ch-ih)/2,iw,ih,drawing["id"]))
        for t in TYPES:
            db.execute("INSERT INTO device_types VALUES(?,?,?,?) ON CONFLICT(id) DO UPDATE SET name=excluded.name,prefix=excluded.prefix,color=excluded.color", t)
        # Old drawings retain their native image-sized canvas and point coordinates.
        db.execute("UPDATE drawings SET canvas_width=width,canvas_height=height,image_width=width,image_height=height WHERE canvas_mode='legacy' AND canvas_width=0")
        db.execute("UPDATE drawings SET name=filename WHERE name='' OR name IS NULL")
        db.execute("UPDATE points SET shape='line',x2=MIN(1,x+0.12),y2=y WHERE type_id='curtain' AND (shape='point' OR x2 IS NULL)")
        db.execute("UPDATE points SET shape='sector',x2=MIN(1,x+0.12),y2=y,fov_angle=COALESCE(fov_angle,60) WHERE type_id='camera' AND (shape='point' OR x2 IS NULL)")
        db.execute("UPDATE points SET shape='area',x2=MIN(1,x+0.12),y2=MIN(1,y+0.08) WHERE type_id='light' AND (shape='point' OR x2 IS NULL)")
        db.execute("UPDATE points SET shape='area',x2=COALESCE(x2,MIN(1,x+0.12)),y2=COALESCE(y2,MIN(1,y+0.08)),product_id=NULL,product_snapshot=NULL,standard_cents=NULL,actual_cents=NULL WHERE type_id='non_smart_area'")
        db.execute("DELETE FROM quote_lines WHERE point_id IN (SELECT id FROM points WHERE type_id='non_smart_area')")
        # 折扣功能已取消，历史折扣不再影响报价总额。
        db.execute("UPDATE quotes SET discount_cents=0 WHERE discount_cents<>0")
        quote_columns = {r[1] for r in db.execute("PRAGMA table_info(quotes)")}
        quote_additions = (
            ("version", "INTEGER NOT NULL DEFAULT 1"),
            ("project_name", "TEXT NOT NULL DEFAULT ''"),
            ("customer_name", "TEXT NOT NULL DEFAULT ''"),
            ("show_empty_categories", "INTEGER NOT NULL DEFAULT 0"),
            ("decimal_places", "INTEGER NOT NULL DEFAULT 2"),
            ("tax_rate_basis", "INTEGER NOT NULL DEFAULT 0"),
            ("discount_type", "TEXT NOT NULL DEFAULT 'none'"),
            ("discount_value", "INTEGER NOT NULL DEFAULT 0"),
            ("template_id", "TEXT NOT NULL DEFAULT 'taiyuan-default-v2'"),
            ("company_info", "TEXT NOT NULL DEFAULT '态元智控'"),
            ("quote_date", "TEXT NOT NULL DEFAULT ''"),
            ("logo_path", "TEXT"),
        )
        for column, definition in quote_additions:
            if column not in quote_columns:
                db.execute(f"ALTER TABLE quotes ADD COLUMN {column} {definition}")
        db.execute("UPDATE quotes SET project_name=(SELECT name FROM projects WHERE projects.id=quotes.project_id) WHERE project_name='' OR project_name IS NULL")
        db.execute("UPDATE quotes SET customer_name=(SELECT customer FROM projects WHERE projects.id=quotes.project_id) WHERE customer_name='' OR customer_name IS NULL")
        db.execute("UPDATE quotes SET quote_date=substr(updated_at,1,10) WHERE quote_date='' OR quote_date IS NULL")
        knowledge_repository.ensure_tables(db)


init_db()
app = FastAPI(title="智能图纸点位与报价工具")
app.add_middleware(CORSMiddleware, allow_origins=["http://localhost:5173", "http://127.0.0.1:5173"], allow_methods=["*"], allow_headers=["*"])


class ProjectIn(BaseModel):
    name: str = Field(min_length=1, max_length=120)
    customer: str = Field(default="", max_length=120)
    address: str = Field(default="", max_length=240)
    status: str = Field(default="进行中", max_length=30)
    folder_id: Optional[str] = None


class ProjectFolderIn(BaseModel):
    name: str = Field(min_length=1, max_length=80)


class ProjectDuplicateIn(BaseModel):
    name: Optional[str] = Field(default=None, max_length=120)


class ProjectSignedIn(BaseModel):
    signed: bool


class ShadowStrengthIn(BaseModel):
    shadow_strength: int = Field(ge=0,le=100)


class BrandIn(BaseModel):
    name: str = Field(min_length=1, max_length=100)
    english_name: str = Field(default="", max_length=120)
    aliases: list[str] = Field(default_factory=list, max_length=30)
    enabled: bool = True
    note: str = Field(default="", max_length=500)


class ProductIn(BaseModel):
    brand_id: str = ""
    brand_name: str = ""
    name: str = Field(min_length=1, max_length=160)
    series: str = Field(default="", max_length=100)
    model: str = Field(default="", max_length=100)
    category: str = Field(min_length=1, max_length=100)
    device_type: str = "other"  # 兼容旧客户端；实际关联由 category 统一推导
    system_category: str = "uncategorized"
    standard_price: str
    cost_price: Optional[str] = None
    origin: str = ""
    unit: str = Field(default="个", min_length=1, max_length=20)
    params: str = ""
    description: str = ""
    enabled: bool = True


class PointIn(BaseModel):
    type_id: str
    product_id: Optional[str] = None
    layer_id: Optional[str] = None
    x: float = Field(ge=0, le=1)
    y: float = Field(ge=0, le=1)
    name: Optional[str] = None
    room: str = ""
    quantity: int = Field(default=1, ge=1, le=100000)
    note: str = ""
    size: int = Field(default=20, ge=12, le=96)
    color: Optional[str] = None
    locked: bool = False
    shape: str = "point"
    x2: Optional[float] = Field(default=None, ge=0, le=1)
    y2: Optional[float] = Field(default=None, ge=0, le=1)
    show_motor: bool = True
    fov_angle: float = Field(default=60, ge=10, le=180)
    line_width: float = Field(default=3, ge=1, le=12)
    curtain_style: str = Field(default="single", pattern="^(single|double|double_layer_single|double_layer_double)$")
    curtain_length: float = Field(default=0, ge=0, le=5000)


class PointPatch(BaseModel):
    type_id: Optional[str] = None
    x: Optional[float] = Field(default=None, ge=0, le=1)
    y: Optional[float] = Field(default=None, ge=0, le=1)
    name: Optional[str] = None
    room: Optional[str] = None
    quantity: Optional[int] = Field(default=None, ge=1, le=100000)
    note: Optional[str] = None
    size: Optional[int] = Field(default=None, ge=12, le=96)
    color: Optional[str] = None
    locked: Optional[bool] = None
    actual_price: Optional[str] = None
    shape: Optional[str] = None
    x2: Optional[float] = Field(default=None, ge=0, le=1)
    y2: Optional[float] = Field(default=None, ge=0, le=1)
    show_motor: Optional[bool] = None
    fov_angle: Optional[float] = Field(default=None, ge=10, le=180)
    line_width: Optional[float] = Field(default=None, ge=1, le=12)
    curtain_style: Optional[str] = Field(default=None, pattern="^(single|double|double_layer_single|double_layer_double)$")
    curtain_length: Optional[float] = Field(default=None, ge=0, le=5000)
    layer_id: Optional[str] = None


class AssignIn(BaseModel):
    point_ids: list[str] = Field(min_length=1)
    product_id: str


class PointRestore(BaseModel):
    id: str
    project_id: str
    drawing_id: str
    layer_id: Optional[str] = None
    type_id: str
    number: str
    name: str
    room: str = ""
    quantity: int = Field(ge=1)
    note: str = ""
    x: float = Field(ge=0, le=1)
    y: float = Field(ge=0, le=1)
    size: int = Field(ge=12, le=96)
    color: str
    locked: int = 0
    product_id: Optional[str] = None
    product_snapshot: Optional[dict] = None
    standard_price: Optional[str] = None
    actual_price: Optional[str] = None
    shape: str = "point"
    x2: Optional[float] = Field(default=None, ge=0, le=1)
    y2: Optional[float] = Field(default=None, ge=0, le=1)
    show_motor: bool = True
    fov_angle: float = Field(default=60, ge=10, le=180)
    line_width: float = Field(default=3, ge=1, le=12)
    curtain_style: str = Field(default="single", pattern="^(single|double|double_layer_single|double_layer_double)$")
    curtain_length: float = Field(default=0, ge=0, le=5000)


class FloorIn(BaseModel):
    name: str = Field(min_length=1,max_length=80)


class FloorPatch(BaseModel):
    name: Optional[str] = Field(default=None,min_length=1,max_length=80)
    sort_order: Optional[int] = Field(default=None,ge=0)


class LayerIn(BaseModel):
    name: str = Field(min_length=1,max_length=80)


class LayerPatch(BaseModel):
    name: Optional[str] = Field(default=None,min_length=1,max_length=80)
    sort_order: Optional[int] = Field(default=None,ge=0)
    visible: Optional[bool] = None
    locked: Optional[bool] = None


class DrawingPatch(BaseModel):
    name: Optional[str] = Field(default=None, min_length=1, max_length=160)
    legend_x: Optional[float] = Field(default=None, ge=0, le=1)
    legend_y: Optional[float] = Field(default=None, ge=0, le=1)
    legend_items: Optional[dict[str,list[float]]] = None
    image_x: Optional[float] = Field(default=None, ge=-3200, le=3200)
    image_y: Optional[float] = Field(default=None, ge=-1800, le=1800)
    image_width: Optional[float] = Field(default=None, gt=0, le=6400)
    image_height: Optional[float] = Field(default=None, gt=0, le=3600)
    image_rotation: Optional[float] = Field(default=None, ge=-3600, le=3600)
    guide_lines: Optional[list[dict[str, Union[float, str]]]] = None


class DrawingCropIn(BaseModel):
    left: float = Field(ge=0, le=0.95)
    top: float = Field(ge=0, le=0.95)
    right: float = Field(ge=0.05, le=1)
    bottom: float = Field(ge=0.05, le=1)


class ProductBatchIn(BaseModel):
    ids: list[str] = Field(min_length=1, max_length=1000)
    action: str
    category: Optional[str] = None
    system_category: Optional[str] = None


class FeedbackIn(BaseModel):
    title: str = Field(min_length=1, max_length=160)
    description: str = Field(min_length=1, max_length=3000)
    module: str = Field(default="", max_length=80)


class QuotePatch(BaseModel):
    note: str = ""


class FeeIn(BaseModel):
    name: str = Field(min_length=1, max_length=120)
    amount: str
    kind: str = "额外费用"


class ManualQuoteLineIn(BaseModel):
    product_id: str
    quantity: int = Field(default=1, ge=1, le=100000)
    actual_price: Optional[str] = None
    note: str = Field(default="", max_length=500)


class ManualQuoteBatchIn(BaseModel):
    product_ids: list[str] = Field(min_length=1, max_length=100)
    quantity: int = Field(default=1, ge=1, le=100000)
    note: str = Field(default="", max_length=500)


class ManualQuoteLinePatch(BaseModel):
    quantity: Optional[int] = Field(default=None, ge=1, le=100000)
    actual_price: Optional[str] = None
    note: Optional[str] = Field(default=None, max_length=500)


class QuotationSettingsPatch(BaseModel):
    project_name: Optional[str] = Field(default=None, min_length=1, max_length=120)
    customer_name: Optional[str] = Field(default=None, max_length=120)
    show_empty_categories: Optional[bool] = None
    decimal_places: Optional[int] = Field(default=None, ge=0, le=2)
    tax_rate: Optional[str] = None
    discount_type: Optional[str] = None
    discount_value: Optional[str] = None
    company_info: Optional[str] = Field(default=None, max_length=500)
    quote_date: Optional[str] = Field(default=None, max_length=20)


class QuotationSystemPatch(BaseModel):
    service_type: str
    service_rate: str = "0"
    service_amount: str = "0"
    service_name: str = Field(default="（服务）硬件安装、设备编程、场景调试、售后维保", max_length=200)


class QuotationItemPatch(BaseModel):
    system_category: Optional[str] = None
    product_name: Optional[str] = Field(default=None, min_length=1, max_length=160)
    product_code: Optional[str] = Field(default=None, max_length=120)
    brand: Optional[str] = Field(default=None, max_length=120)
    origin: Optional[str] = Field(default=None, max_length=120)
    quantity: Optional[int] = Field(default=None, ge=0, le=100000)
    unit_price: Optional[str] = None
    unit: Optional[str] = Field(default=None, min_length=1, max_length=20)
    params: Optional[str] = Field(default=None, max_length=3000)
    remark: Optional[str] = Field(default=None, max_length=1000)
    sort_order: Optional[int] = Field(default=None, ge=0)
    visible: Optional[bool] = None


class TemporaryQuotationItemIn(BaseModel):
    system_category: str
    product_name: str = Field(min_length=1, max_length=160)
    product_code: str = Field(default="", max_length=120)
    brand: str = Field(default="", max_length=120)
    origin: str = Field(default="", max_length=120)
    quantity: int = Field(default=1, ge=1, le=100000)
    unit_price: str
    unit: str = Field(default="台", min_length=1, max_length=20)
    params: str = Field(default="", max_length=3000)
    remark: str = Field(default="", max_length=1000)
    product_id: Optional[str] = None
    save_to_library: bool = False


class QuotationReorderIn(BaseModel):
    item_ids: list[str] = Field(min_length=1, max_length=1000)
    system_category: Optional[str] = None


class QuotationReplaceIn(BaseModel):
    product_id: str = Field(min_length=1, max_length=120)


class ConstructionFeesPatch(BaseModel):
    fees: dict[str, str]


class QuotationVersionIn(BaseModel):
    note: str = Field(default="", max_length=500)
    created_by: str = Field(default="本机用户", max_length=100)


@app.get("/api/version")
def version():
    return {"version": APP_VERSION, "local_data": PATHS.local_data_mode}


MIGRATION_REPORT: list[dict] = []


@app.get("/api/storage-status")
def storage_status(response: Response):
    response.headers["Cache-Control"] = "no-store, max-age=0"
    return {**product_repository.status(), "data_dir": str(DATA), "database": str(DB), "migrations": MIGRATION_REPORT}


def brand_out(brand):
    item = dict(brand)
    try:
        item["aliases"] = json.loads(item.get("aliases") or "[]")
    except (TypeError, json.JSONDecodeError):
        item["aliases"] = []
    item["enabled"] = bool(item.get("enabled", 1))
    item["logo_path"] = item.get("logo_svg_path") or item.get("logo_png_path")
    item["logo_url"] = f'/api/brand-files/{item["logo_path"]}' if item["logo_path"] else None
    return item


def find_brand(db, *, brand_id_value="", name=""):
    if brand_id_value:
        row = db.execute("SELECT * FROM brands WHERE id=?", (brand_id_value,)).fetchone()
        if row:
            return dict(row)
    clean = name.strip().casefold()
    if clean:
        for row in db.execute("SELECT * FROM brands"):
            aliases = []
            try: aliases = json.loads(row["aliases"] or "[]")
            except (TypeError, json.JSONDecodeError): pass
            candidates = [row["name"], row["english_name"], *aliases]
            if any(str(value or "").strip().casefold() == clean for value in candidates):
                return dict(row)
    return None


def product_out(p, db=None):
    p = dict(p)
    p["standard_price"] = f'{p["standard_cents"] / 100:.2f}'
    p["cost_price"] = f'{p["cost_cents"] / 100:.2f}' if p.get("cost_cents") is not None else None
    p["image_url"] = f'/api/files/{p["image_path"]}' if p["image_path"] else None
    brand = find_brand(db, brand_id_value=p.get("brand_id", ""), name=p.get("brand_name", "")) if db else None
    if brand:
        # 旧版本曾生成随机品牌 ID；API 统一回传实际品牌库 ID，编辑后会写回产品笔记。
        p["brand_id"] = brand["id"]
    p["brand_logo_url"] = brand_out(brand)["logo_url"] if brand else None
    return p


def point_out(p):
    p = dict(p)
    p["standard_price"] = f'{p["standard_cents"] / 100:.2f}' if p["standard_cents"] is not None else None
    p["actual_price"] = f'{p["actual_cents"] / 100:.2f}' if p["actual_cents"] is not None else None
    p["product_snapshot"] = json.loads(p["product_snapshot"]) if p["product_snapshot"] else None
    return p


@app.get("/api/types")
def types():
    return [{
        "id": t[0], "name": t[1], "prefix": t[2], "color": t[3],
        "editor_selectable": t[0] not in LEGACY_TYPES,
        "product_selectable": t[0] not in LEGACY_TYPES | NON_PRODUCT_TYPES,
        "legacy": t[0] in LEGACY_TYPES,
        "kind": "area" if t[0] in NON_PRODUCT_TYPES else "device",
    } for t in TYPES]


@app.get("/api/projects")
def projects():
    started = time.perf_counter()
    with connect() as db:
        scan_started = time.perf_counter()
        result = rows(db.execute("SELECT p.*,pf.name AS folder_name,(SELECT count(*) FROM points WHERE project_id=p.id AND type_id<>'non_smart_area') point_count,COALESCE((SELECT sum(subtotal_cents) FROM quote_lines ql JOIN quotes q ON q.id=ql.quote_id WHERE q.project_id=p.id),0)+COALESCE((SELECT sum(subtotal_cents) FROM manual_quote_lines ml JOIN quotes q ON q.id=ml.quote_id WHERE q.project_id=p.id),0) line_total_cents FROM projects p LEFT JOIN project_folders pf ON pf.id=p.folder_id ORDER BY updated_at DESC"))
        record_event("projectList.scan", operation="local_sqlite_index", projectCount=len(result),
            durationMs=round((time.perf_counter()-scan_started)*1000, 2), success=True)
        parse_started = time.perf_counter()
        for p in result:
            quotation=quotation_v2_data(db,p["id"])
            p["total_price"] = f'{quotation["totals"]["grand_total_cents"]/100:.2f}'
        record_event("projectList.parse", operation="summary_and_quote_totals", projectCount=len(result),
            durationMs=round((time.perf_counter()-parse_started)*1000, 2), success=True)
        record_event("projectList.load", operation="local_index", projectCount=len(result),
            durationMs=round((time.perf_counter()-started)*1000, 2), success=True)
        return result


def persist_all_projects():
    with connect() as db:
        project_ids = [row[0] for row in db.execute("SELECT id FROM projects")]
    for project_id in project_ids:
        _persist_local_project(project_id)


@app.get("/api/project-folders")
def list_project_folders():
    with connect() as db:
        return rows(db.execute("SELECT * FROM project_folders ORDER BY sort_order,name"))


@app.post("/api/project-folders", status_code=201)
def create_project_folder(data: ProjectFolderIn):
    name = data.name.strip()
    if not name:
        raise HTTPException(422, "文件夹名称不能为空")
    with connect() as db:
        if db.execute("SELECT 1 FROM project_folders WHERE name=?", (name,)).fetchone():
            raise HTTPException(409, "已存在同名文件夹")
        order = db.execute("SELECT COALESCE(MAX(sort_order),-1)+1 FROM project_folders").fetchone()[0]
        db.execute("INSERT INTO project_folders VALUES(?,?,?,?,?)", (uid(), name, order, now(), now()))
    persist_all_projects()
    return list_project_folders()


@app.patch("/api/project-folders/{folder_id}")
def rename_project_folder(folder_id: str, data: ProjectFolderIn):
    name = data.name.strip()
    if not name:
        raise HTTPException(422, "文件夹名称不能为空")
    with connect() as db:
        one(db, "SELECT id FROM project_folders WHERE id=?", (folder_id,))
        if db.execute("SELECT 1 FROM project_folders WHERE name=? AND id<>?", (name, folder_id)).fetchone():
            raise HTTPException(409, "已存在同名文件夹")
        db.execute("UPDATE project_folders SET name=?,updated_at=? WHERE id=?", (name, now(), folder_id))
    persist_all_projects()
    return list_project_folders()


@app.delete("/api/project-folders/{folder_id}")
def delete_project_folder(folder_id: str):
    with connect() as db:
        one(db, "SELECT id FROM project_folders WHERE id=?", (folder_id,))
        db.execute("UPDATE projects SET folder_id=NULL,updated_at=? WHERE folder_id=?", (now(), folder_id))
        db.execute("DELETE FROM project_folders WHERE id=?", (folder_id,))
    persist_all_projects()
    return list_project_folders()


@app.post("/api/projects", status_code=201)
def create_project(data: ProjectIn):
    p = {"id": uid(), **data.model_dump(), "signed": 0, "created_at": now(), "updated_at": now()}
    floor_id,layer_id,drawing_id=uid(),uid(),uid()
    blank_path=UPLOADS/f"{drawing_id}_blank.png"
    Image.new("RGB",(1600,900),"white").save(blank_path,"PNG")
    with connect() as db:
        if p["folder_id"]:
            one(db, "SELECT id FROM project_folders WHERE id=?", (p["folder_id"],))
        db.execute("INSERT INTO projects(id,name,customer,address,status,signed,created_at,updated_at,folder_id) VALUES(:id,:name,:customer,:address,:status,:signed,:created_at,:updated_at,:folder_id)", p)
        db.execute("INSERT INTO quotes(id,project_id,discount_cents,note,updated_at,project_name,customer_name,quote_date) VALUES(?,?,?,?,?,?,?,?)", (uid(), p["id"], 0, "", now(), p["name"], p["customer"], now()[:10]))
        db.execute("INSERT INTO floors(id,project_id,name,sort_order,visibility_initialized) VALUES(?,?,?,0,1)",(floor_id,p["id"],"默认楼层"))
        db.execute("INSERT INTO layers(id,floor_id,name,sort_order) VALUES(?,?,?,0)",(layer_id,floor_id,"默认图层"))
        db.execute("""INSERT INTO drawings(id,project_id,floor_id,source_path,background_path,filename,page,width,height,created_at,name,canvas_mode,canvas_width,canvas_height,image_x,image_y,image_width,image_height,image_rotation,guide_lines,is_blank)
            VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",(drawing_id,p["id"],floor_id,blank_path.name,blank_path.name,"未导入主图纸.png",1,1600,900,now(),"主图纸","fixed",1600,900,0,0,1600,900,0,"[]",1))
    if PATHS.local_data_mode:
        _persist_local_project(p["id"])
    return p


@app.post("/api/projects/{pid}/duplicate", status_code=201)
def duplicate_project(pid: str, data: ProjectDuplicateIn = ProjectDuplicateIn()):
    copied_files: list[Path] = []
    file_map: dict[str, str] = {}
    try:
        with connect() as db:
            source = one(db, "SELECT * FROM projects WHERE id=?", (pid,))
            source_quote = one(db, "SELECT * FROM quotes WHERE project_id=?", (pid,))
            project_id, quote_id = uid(), uid()
            name = (data.name or "").strip() or f"{source['name']} 副本"
            timestamp = now()
            clone = {**source, "id": project_id, "name": name, "signed": 0,
                     "created_at": timestamp, "updated_at": timestamp}
            db.execute(
                "INSERT INTO projects(id,name,customer,address,status,signed,created_at,updated_at,shadow_strength,folder_id) VALUES(?,?,?,?,?,?,?,?,?,?)",
                (clone["id"], clone["name"], clone["customer"], clone["address"], clone["status"], 0,
                 timestamp, timestamp, clone.get("shadow_strength", 25), clone.get("folder_id")),
            )

            floor_map: dict[str, str] = {}
            for row in rows(db.execute("SELECT * FROM floors WHERE project_id=? ORDER BY sort_order,id", (pid,))):
                new_id = uid(); floor_map[row["id"]] = new_id
                values = {**row, "id": new_id, "project_id": project_id}
                cols = list(values)
                db.execute(f"INSERT INTO floors({','.join(cols)}) VALUES({','.join('?' for _ in cols)})", [values[c] for c in cols])

            layer_map: dict[str, str] = {}
            for row in rows(db.execute("SELECT l.* FROM layers l JOIN floors f ON f.id=l.floor_id WHERE f.project_id=? ORDER BY l.sort_order,l.id", (pid,))):
                new_id = uid(); layer_map[row["id"]] = new_id
                values = {**row, "id": new_id, "floor_id": floor_map[row["floor_id"]]}
                cols = list(values)
                db.execute(f"INSERT INTO layers({','.join(cols)}) VALUES({','.join('?' for _ in cols)})", [values[c] for c in cols])

            drawing_rows = rows(db.execute("SELECT * FROM drawings WHERE project_id=? ORDER BY created_at,id", (pid,)))
            drawing_map: dict[str, str] = {}
            for row in drawing_rows:
                new_id = uid(); drawing_map[row["id"]] = new_id
                values = {**row, "id": new_id, "project_id": project_id,
                          "floor_id": floor_map.get(row.get("floor_id"))}
                for field in ("source_path", "background_path"):
                    original = row.get(field)
                    if not original:
                        continue
                    if original not in file_map:
                        suffix = Path(original).suffix
                        target_name = f"{uid()}{suffix}"
                        source_path, target_path = UPLOADS / Path(original).name, UPLOADS / target_name
                        if source_path.is_file():
                            shutil.copy2(source_path, target_path)
                            copied_files.append(target_path)
                        file_map[original] = target_name
                    values[field] = file_map[original]
                cols = list(values)
                db.execute(f"INSERT INTO drawings({','.join(cols)}) VALUES({','.join('?' for _ in cols)})", [values[c] for c in cols])

            point_map: dict[str, str] = {}
            for row in rows(db.execute("SELECT * FROM points WHERE project_id=? ORDER BY created_at,id", (pid,))):
                new_id = uid(); point_map[row["id"]] = new_id
                values = {**row, "id": new_id, "project_id": project_id,
                          "drawing_id": drawing_map[row["drawing_id"]],
                          "layer_id": layer_map.get(row.get("layer_id"))}
                cols = list(values)
                db.execute(f"INSERT INTO points({','.join(cols)}) VALUES({','.join('?' for _ in cols)})", [values[c] for c in cols])

            for row in rows(db.execute("SELECT * FROM point_counters WHERE project_id=?", (pid,))):
                db.execute("INSERT INTO point_counters(project_id,type_id,last_number) VALUES(?,?,?)",
                           (project_id, row["type_id"], row["last_number"]))

            cloned_quote = {**source_quote, "id": quote_id, "project_id": project_id,
                            "project_name": name, "updated_at": timestamp}
            cols = list(cloned_quote)
            db.execute(f"INSERT INTO quotes({','.join(cols)}) VALUES({','.join('?' for _ in cols)})", [cloned_quote[c] for c in cols])

            manual_map: dict[str, str] = {}
            for table in ("extra_fees", "manual_quote_lines", "quote_lines", "quotation_systems", "quotation_items", "quotation_construction_fees", "quotation_versions"):
                table_rows = rows(db.execute(f"SELECT * FROM {table} WHERE quote_id=?", (source_quote["id"],)))
                for row in table_rows:
                    values = dict(row); old_id = values.get("id")
                    values["quote_id"] = quote_id
                    if old_id:
                        values["id"] = uid()
                        if table == "manual_quote_lines":
                            manual_map[old_id] = values["id"]
                    if table == "quote_lines":
                        values["point_id"] = point_map.get(values["point_id"], values["point_id"])
                    if table == "quotation_versions":
                        try:
                            snapshot = json.loads(values["snapshot_json"])
                            snapshot["quote_id"] = quote_id
                            if isinstance(snapshot.get("project"), dict):
                                snapshot["project"]["id"] = project_id
                                snapshot["project"]["name"] = name
                            if isinstance(snapshot.get("settings"), dict):
                                snapshot["settings"]["project_name"] = name
                            values["snapshot_json"] = json.dumps(snapshot, ensure_ascii=False)
                        except (ValueError, TypeError):
                            pass
                    if table == "quotation_items":
                        if values.get("source_type") == "points":
                            values["source_key"] = point_map.get(values.get("source_key"), values.get("source_key"))
                        elif values.get("source_type") == "manual":
                            values["source_key"] = manual_map.get(values.get("source_key"), values.get("source_key"))
                        elif values.get("source_type") == "temp":
                            values["source_key"] = uid()
                    cols = list(values)
                    db.execute(f"INSERT INTO {table}({','.join(cols)}) VALUES({','.join('?' for _ in cols)})", [values[c] for c in cols])
        if PATHS.local_data_mode:
            _persist_local_project(project_id)
        return get_project(project_id)
    except Exception:
        for path in copied_files:
            try: path.unlink(missing_ok=True)
            except OSError: pass
        raise


@app.get("/api/projects/{pid}")
def get_project(pid: str):
    with connect() as db:
        p = one(db, "SELECT * FROM projects WHERE id=?", (pid,))
        p["drawings"] = rows(db.execute("SELECT * FROM drawings WHERE project_id=? ORDER BY created_at", (pid,)))
        for drawing in p["drawings"]:
            drawing["legend_items"]=json.loads(drawing.get("legend_items") or "{}")
            try: drawing["guide_lines"]=json.loads(drawing.get("guide_lines") or "[]")
            except Exception: drawing["guide_lines"]=[]
        p["floors"] = rows(db.execute("SELECT * FROM floors WHERE project_id=? ORDER BY sort_order,id",(pid,)))
        p["layers"] = rows(db.execute("SELECT l.* FROM layers l JOIN floors f ON f.id=l.floor_id WHERE f.project_id=? ORDER BY l.sort_order,l.id",(pid,)))
        p["points"] = [point_out(x) for x in rows(db.execute("SELECT * FROM points WHERE project_id=? ORDER BY created_at", (pid,)))]
        return p


@app.post("/api/projects/{pid}/floors",status_code=201)
def create_floor(pid: str,data: FloorIn):
    floor_id,layer_id,drawing_id=uid(),uid(),uid()
    blank_path=UPLOADS/f"{drawing_id}_blank.png"
    Image.new("RGB",(1600,900),"white").save(blank_path,"PNG")
    with connect() as db:
        one(db,"SELECT id FROM projects WHERE id=?",(pid,))
        order=db.execute("SELECT COALESCE(MAX(sort_order),-1)+1 FROM floors WHERE project_id=?",(pid,)).fetchone()[0]
        db.execute("INSERT INTO floors(id,project_id,name,sort_order,visibility_initialized) VALUES(?,?,?,?,1)",(floor_id,pid,data.name.strip(),order))
        db.execute("INSERT INTO layers VALUES(?,?,?,?,?,?)",(layer_id,floor_id,"默认图层",0,1,0))
        db.execute("""INSERT INTO drawings(id,project_id,floor_id,source_path,background_path,filename,page,width,height,created_at,name,canvas_mode,canvas_width,canvas_height,image_x,image_y,image_width,image_height,image_rotation,guide_lines,is_blank)
            VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",(drawing_id,pid,floor_id,blank_path.name,blank_path.name,"未导入主图纸.png",1,1600,900,now(),"主图纸","fixed",1600,900,0,0,1600,900,0,"[]",1))
    return get_project(pid)


@app.patch("/api/projects/{pid}/floors/{floor_id}")
def update_floor(pid: str,floor_id: str,data: FloorPatch):
    with connect() as db:
        one(db,"SELECT id FROM floors WHERE id=? AND project_id=?",(floor_id,pid))
        if data.name is not None: db.execute("UPDATE floors SET name=? WHERE id=?",(data.name.strip(),floor_id))
        if data.sort_order is not None:
            ids=[row[0] for row in db.execute("SELECT id FROM floors WHERE project_id=? ORDER BY sort_order,id",(pid,))]
            ids.remove(floor_id); ids.insert(min(data.sort_order,len(ids)),floor_id)
            for index,item_id in enumerate(ids): db.execute("UPDATE floors SET sort_order=? WHERE id=?",(index,item_id))
    return get_project(pid)


@app.delete("/api/projects/{pid}/floors/{floor_id}")
def delete_floor(pid: str,floor_id: str):
    with connect() as db:
        one(db,"SELECT id FROM floors WHERE id=? AND project_id=?",(floor_id,pid))
        if db.execute("SELECT COUNT(*) FROM floors WHERE project_id=?",(pid,)).fetchone()[0]<=1:
            raise HTTPException(409,"项目至少保留一个楼层")
        files=rows(db.execute("SELECT source_path,background_path FROM drawings WHERE floor_id=?",(floor_id,)))
        db.execute("DELETE FROM points WHERE layer_id IN (SELECT id FROM layers WHERE floor_id=?)",(floor_id,))
        db.execute("DELETE FROM drawings WHERE floor_id=?",(floor_id,))
        db.execute("DELETE FROM floors WHERE id=?",(floor_id,))
    for drawing in files:
        for key in ("source_path","background_path"):
            path=UPLOADS/drawing[key]
            if path.is_file(): path.unlink()
    return get_project(pid)


@app.post("/api/projects/{pid}/floors/{floor_id}/layers",status_code=201)
def create_layer(pid: str,floor_id: str,data: LayerIn):
    with connect() as db:
        one(db,"SELECT id FROM floors WHERE id=? AND project_id=?",(floor_id,pid))
        order=db.execute("SELECT COALESCE(MAX(sort_order),-1)+1 FROM layers WHERE floor_id=?",(floor_id,)).fetchone()[0]
        db.execute("INSERT INTO layers VALUES(?,?,?,?,?,?)",(uid(),floor_id,data.name.strip(),order,1,0))
    return get_project(pid)


@app.patch("/api/projects/{pid}/layers/{layer_id}")
def update_layer(pid: str,layer_id: str,data: LayerPatch):
    with connect() as db:
        one(db,"SELECT l.id FROM layers l JOIN floors f ON f.id=l.floor_id WHERE l.id=? AND f.project_id=?",(layer_id,pid))
        changes=data.model_dump(exclude_unset=True)
        if "name" in changes: changes["name"]=changes["name"].strip()
        if changes: db.execute("UPDATE layers SET "+",".join(f"{key}=?" for key in changes)+" WHERE id=?",(*[int(value) if isinstance(value,bool) else value for value in changes.values()],layer_id))
    return get_project(pid)


@app.delete("/api/projects/{pid}/layers/{layer_id}")
def delete_layer(pid: str,layer_id: str):
    with connect() as db:
        layer=one(db,"SELECT l.* FROM layers l JOIN floors f ON f.id=l.floor_id WHERE l.id=? AND f.project_id=?",(layer_id,pid))
        if db.execute("SELECT COUNT(*) FROM layers WHERE floor_id=?",(layer["floor_id"],)).fetchone()[0]<=1:
            raise HTTPException(409,"楼层至少保留一个图层")
        db.execute("DELETE FROM points WHERE layer_id=?",(layer_id,))
        db.execute("DELETE FROM layers WHERE id=?",(layer_id,))
    return get_project(pid)


class CopyToFloor(BaseModel):
    floor_id: str


def _copy_point_to(db,point: dict,drawing_id: str,layer_id: str):
    point_id=uid(); number=next_number(db,point["project_id"],point["type_id"])
    fields=("id","project_id","drawing_id","layer_id","type_id","number","name","room","quantity","note","x","y","size","color","locked","product_id","product_snapshot","standard_cents","actual_cents","created_at","updated_at","shape","x2","y2","show_motor","fov_angle","line_width","curtain_style","curtain_length")
    values={**point,"id":point_id,"drawing_id":drawing_id,"layer_id":layer_id,"number":number,"name":number if point["name"]==point["number"] else point["name"],"created_at":now(),"updated_at":now()}
    db.execute("INSERT INTO points("+",".join(fields)+") VALUES("+",".join("?" for _ in fields)+")",tuple(values[field] for field in fields))
    sync_line(db,point_id)


@app.post("/api/projects/{pid}/layers/{layer_id}/copy",status_code=201)
def copy_layer(pid: str,layer_id: str,data: CopyToFloor):
    with connect() as db:
        source=one(db,"SELECT l.* FROM layers l JOIN floors f ON f.id=l.floor_id WHERE l.id=? AND f.project_id=?",(layer_id,pid))
        one(db,"SELECT id FROM floors WHERE id=? AND project_id=?",(data.floor_id,pid))
        points=rows(db.execute("SELECT * FROM points WHERE layer_id=?",(layer_id,)))
        source_drawings=[row[0] for row in db.execute("SELECT id FROM drawings WHERE floor_id=? ORDER BY created_at,id",(source["floor_id"],))]
        target_drawings=[row[0] for row in db.execute("SELECT id FROM drawings WHERE floor_id=? ORDER BY created_at,id",(data.floor_id,))]
        if points and len(target_drawings)<len(source_drawings): raise HTTPException(409,"目标楼层请先复制对应数量的图纸")
        drawing_map=dict(zip(source_drawings,target_drawings))
        new_layer=uid(); order=db.execute("SELECT COALESCE(MAX(sort_order),-1)+1 FROM layers WHERE floor_id=?",(data.floor_id,)).fetchone()[0]
        db.execute("INSERT INTO layers VALUES(?,?,?,?,?,?)",(new_layer,data.floor_id,source["name"]+" 副本",order,source["visible"],source["locked"]))
        for point in points: _copy_point_to(db,point,drawing_map[point["drawing_id"]],new_layer)
    return get_project(pid)


@app.patch("/api/projects/{pid}")
def update_project(pid: str, data: ProjectIn):
    with connect() as db:
        old=one(db, "SELECT * FROM projects WHERE id=?", (pid,))
        folder_id = old.get("folder_id")
        if "folder_id" in data.model_fields_set:
            folder_id = data.folder_id
            if folder_id:
                one(db, "SELECT id FROM project_folders WHERE id=?", (folder_id,))
        db.execute("UPDATE projects SET name=?,customer=?,address=?,status=?,folder_id=?,updated_at=? WHERE id=?", (data.name,data.customer,data.address,data.status,folder_id,now(),pid))
        db.execute("UPDATE quotes SET project_name=CASE WHEN project_name=? OR project_name='' THEN ? ELSE project_name END,customer_name=CASE WHEN customer_name=? OR customer_name='' THEN ? ELSE customer_name END WHERE project_id=?",(old["name"],data.name,old["customer"],data.customer,pid))
    return get_project(pid)


@app.patch("/api/projects/{pid}/signed")
def update_project_signed(pid: str, data: ProjectSignedIn):
    with connect() as db:
        one(db, "SELECT id FROM projects WHERE id=?", (pid,))
        db.execute("UPDATE projects SET signed=?,updated_at=? WHERE id=?", (int(data.signed), now(), pid))
    return get_project(pid)


@app.patch("/api/projects/{pid}/shadow")
def update_project_shadow(pid: str,data: ShadowStrengthIn):
    with connect() as db:
        one(db,"SELECT id FROM projects WHERE id=?",(pid,))
        db.execute("UPDATE projects SET shadow_strength=? WHERE id=?",(data.shadow_strength,pid))
    return get_project(pid)


@app.delete("/api/projects/{pid}")
def delete_project(pid: str):
    with connect() as db:
        one(db,"SELECT id FROM projects WHERE id=?",(pid,))
        drawing_files=rows(db.execute("SELECT source_path,background_path FROM drawings WHERE project_id=?",(pid,)))
        db.execute("DELETE FROM projects WHERE id=?",(pid,))
    with connect() as db:
        knowledge_repository.delete_project_note(db, pid)
    cleanup_warnings=[]
    for drawing in drawing_files:
        for key in ("source_path","background_path"):
            path=UPLOADS/drawing[key]
            if path.is_file():
                try: path.unlink()
                except OSError: cleanup_warnings.append(path.name)
    return {"ok":True,"cleanup_warnings":cleanup_warnings}


@app.post("/api/projects/{pid}/drawings", status_code=201)
async def upload_drawing(pid: str, file: UploadFile = File(...), page: int = Form(1), floor_id: Optional[str] = Form(None), drawing_id: Optional[str] = Form(None)):
    with connect() as db:
        one(db, "SELECT id FROM projects WHERE id=?", (pid,))
        if floor_id: one(db,"SELECT id FROM floors WHERE id=? AND project_id=?",(floor_id,pid))
        else: floor_id=one(db,"SELECT id FROM floors WHERE project_id=? ORDER BY sort_order LIMIT 1",(pid,))["id"]
        target=one(db,"SELECT * FROM drawings WHERE id=? AND project_id=? AND floor_id=?",(drawing_id,pid,floor_id)) if drawing_id else None
    ext = Path(file.filename or "").suffix.lower()
    if ext not in (".png", ".jpg", ".jpeg", ".pdf"):
        raise HTTPException(422, "仅支持 JPG、PNG、PDF")
    content = await file.read()
    if len(content) > 50 * 1024 * 1024:
        raise HTTPException(413, "文件不能超过 50 MB")
    did = target["id"] if target else uid()
    file_token=uid()
    source = UPLOADS / f"{file_token}{ext}"
    background = UPLOADS / f"{file_token}_background.png"
    try:
        if ext == ".pdf":
            doc = fitz.open(stream=content, filetype="pdf")
            if not 1 <= page <= len(doc):
                raise HTTPException(422, f"PDF 页码须在 1 到 {len(doc)} 之间")
            pix = doc[page-1].get_pixmap(matrix=fitz.Matrix(2,2), alpha=False)
            Image.frombytes("RGB", (pix.width,pix.height), pix.samples).save(background)
            width,height = pix.width,pix.height
        else:
            im = ImageOps.exif_transpose(Image.open(io.BytesIO(content))).convert("RGB")
            width,height = im.size
            if width * height > 80000000:
                raise HTTPException(422, "图纸像素尺寸过大")
            im.save(background)
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422, "图纸文件无法读取")
    source.write_bytes(content)
    drawing_name = Path(file.filename or "图纸").stem + (f" · 第 {page} 页" if ext == ".pdf" else "")
    d = {"id":did,"project_id":pid,"floor_id":floor_id,"source_path":source.name,"background_path":background.name,"filename":file.filename,"page":page,"width":width,"height":height,"created_at":target["created_at"] if target else now(),"name":drawing_name,"rotation":0,"legend_x":0.025,"legend_y":0.025}
    with connect() as db:
        if target and target.get("canvas_mode")=="fixed":
            cw,ch=target["canvas_width"] or 1600,target["canvas_height"] or 900
            fit_scale=min(cw/width,ch/height); image_width,image_height=width*fit_scale,height*fit_scale
            d.update(canvas_mode="fixed",canvas_width=cw,canvas_height=ch,image_x=(cw-image_width)/2,image_y=(ch-image_height)/2,image_width=image_width,image_height=image_height,image_rotation=0,guide_lines=target.get("guide_lines","[]"),is_blank=0)
            db.execute("""UPDATE drawings SET source_path=:source_path,background_path=:background_path,filename=:filename,page=:page,width=:width,height=:height,name=:name,image_x=:image_x,image_y=:image_y,image_width=:image_width,image_height=:image_height,image_rotation=0,is_blank=0,crop_source_path=NULL,crop_source_width=NULL,crop_source_height=NULL,crop_frame_x=NULL,crop_frame_y=NULL,crop_frame_width=NULL,crop_frame_height=NULL,crop_left=NULL,crop_top=NULL,crop_right=NULL,crop_bottom=NULL WHERE id=:id""",d)
        else:
            db.execute("INSERT INTO drawings(id,project_id,floor_id,source_path,background_path,filename,page,width,height,created_at,name,canvas_mode,canvas_width,canvas_height,image_x,image_y,image_width,image_height,image_rotation,guide_lines,is_blank) VALUES(:id,:project_id,:floor_id,:source_path,:background_path,:filename,:page,:width,:height,:created_at,:name,'legacy',:width,:height,0,0,:width,:height,0,'[]',0)",d)
        db.execute("UPDATE projects SET updated_at=? WHERE id=?", (now(),pid))
        d.update(canvas_mode=(target.get("canvas_mode") if target else "legacy"),canvas_width=target.get("canvas_width") if target else width,canvas_height=target.get("canvas_height") if target else height,image_x=target.get("image_x",0) if target else 0,image_y=target.get("image_y",0) if target else 0,image_width=target.get("image_width") if target else width,image_height=target.get("image_height") if target else height,image_rotation=target.get("image_rotation",0) if target else 0,guide_lines=target.get("guide_lines","[]") if target else "[]",is_blank=0)
    return d


@app.post("/api/projects/{pid}/drawings/{drawing_id}/copy",status_code=201)
def copy_drawing(pid: str,drawing_id: str,data: CopyToFloor):
    with connect() as db:
        source=one(db,"SELECT * FROM drawings WHERE id=? AND project_id=?",(drawing_id,pid))
        one(db,"SELECT id FROM floors WHERE id=? AND project_id=?",(data.floor_id,pid))
        target_canvas=db.execute("SELECT * FROM drawings WHERE floor_id=? AND canvas_mode='fixed' ORDER BY created_at,id LIMIT 1",(data.floor_id,)).fetchone() if source.get("canvas_mode")=="fixed" else None
        new_id=dict(target_canvas)["id"] if target_canvas else uid()
        file_token=uid()
        for key in ("source_path","background_path"):
            original=UPLOADS/source[key]
            if not original.is_file(): raise HTTPException(404,"图纸源文件不存在")
            destination=UPLOADS/(file_token+("_background" if key=="background_path" else "")+original.suffix)
            shutil.copy2(original,destination)
        copied_source=file_token+Path(source["source_path"]).suffix
        copied_background=file_token+"_background"+Path(source["background_path"]).suffix
        copied_crop_source=None
        if source.get("crop_source_path"):
            crop_original=UPLOADS/source["crop_source_path"]
            if not crop_original.is_file(): raise HTTPException(404,"裁剪原图文件不存在")
            copied_crop_source=file_token+"_crop_source"+crop_original.suffix
            shutil.copy2(crop_original,UPLOADS/copied_crop_source)
        if target_canvas:
            db.execute("UPDATE drawings SET source_path=?,background_path=?,filename=?,page=?,width=?,height=?,name=?,rotation=?,legend_x=?,legend_y=?,legend_items=?,image_x=?,image_y=?,image_width=?,image_height=?,image_rotation=?,guide_lines=?,is_blank=0,crop_source_path=?,crop_source_width=?,crop_source_height=?,crop_frame_x=?,crop_frame_y=?,crop_frame_width=?,crop_frame_height=?,crop_left=?,crop_top=?,crop_right=?,crop_bottom=? WHERE id=?",(copied_source,copied_background,source["filename"],source["page"],source["width"],source["height"],source["name"]+" 副本",source["rotation"],source["legend_x"],source["legend_y"],source["legend_items"],source.get("image_x",0),source.get("image_y",0),source.get("image_width",source["width"]),source.get("image_height",source["height"]),source.get("image_rotation",0),source.get("guide_lines","[]"),copied_crop_source,source.get("crop_source_width"),source.get("crop_source_height"),source.get("crop_frame_x"),source.get("crop_frame_y"),source.get("crop_frame_width"),source.get("crop_frame_height"),source.get("crop_left"),source.get("crop_top"),source.get("crop_right"),source.get("crop_bottom"),new_id))
        else:
            db.execute("INSERT INTO drawings(id,project_id,floor_id,source_path,background_path,filename,page,width,height,created_at,name,rotation,legend_x,legend_y,legend_items,canvas_mode,canvas_width,canvas_height,image_x,image_y,image_width,image_height,image_rotation,guide_lines,is_blank,crop_source_path,crop_source_width,crop_source_height,crop_frame_x,crop_frame_y,crop_frame_width,crop_frame_height,crop_left,crop_top,crop_right,crop_bottom) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",(new_id,pid,data.floor_id,copied_source,copied_background,source["filename"],source["page"],source["width"],source["height"],now(),source["name"]+" 副本",source["rotation"],source["legend_x"],source["legend_y"],source["legend_items"],source.get("canvas_mode","legacy"),source.get("canvas_width",source["width"]),source.get("canvas_height",source["height"]),source.get("image_x",0),source.get("image_y",0),source.get("image_width",source["width"]),source.get("image_height",source["height"]),source.get("image_rotation",0),source.get("guide_lines","[]"),source.get("is_blank",0),copied_crop_source,source.get("crop_source_width"),source.get("crop_source_height"),source.get("crop_frame_x"),source.get("crop_frame_y"),source.get("crop_frame_width"),source.get("crop_frame_height"),source.get("crop_left"),source.get("crop_top"),source.get("crop_right"),source.get("crop_bottom")))
    return get_project(pid)


@app.patch("/api/projects/{pid}/drawings/{drawing_id}")
def rename_drawing(pid: str, drawing_id: str, data: DrawingPatch):
    with connect() as db:
        drawing=one(db, "SELECT * FROM drawings WHERE id=? AND project_id=?", (drawing_id, pid))
        changes=[]; values=[]
        if data.name is not None: changes.append("name=?"); values.append(data.name.strip())
        if data.legend_x is not None: changes.append("legend_x=?"); values.append(data.legend_x)
        if data.legend_y is not None: changes.append("legend_y=?"); values.append(data.legend_y)
        if data.legend_items is not None:
            if any((key not in TYPE_MAP and not (key.startswith("other::") and len(key)<=110)) or len(value)!=2 or any(not isinstance(v,(int,float)) or not 0<=v<=1 for v in value) for key,value in data.legend_items.items()):
                raise HTTPException(422,"图例坐标无效")
            changes.append("legend_items=?"); values.append(json.dumps(data.legend_items,ensure_ascii=False))
        for key in ("image_x","image_y","image_width","image_height","image_rotation"):
            value=getattr(data,key)
            if value is not None:
                changes.append(f"{key}=?"); values.append(value)
                if drawing.get("crop_source_path"):
                    if key=="image_x" and drawing.get("image_x") is not None:
                        changes.append("crop_frame_x=?"); values.append((drawing.get("crop_frame_x") if drawing.get("crop_frame_x") is not None else drawing["image_x"])+(value-drawing["image_x"]))
                    elif key=="image_y" and drawing.get("image_y") is not None:
                        changes.append("crop_frame_y=?"); values.append((drawing.get("crop_frame_y") if drawing.get("crop_frame_y") is not None else drawing["image_y"])+(value-drawing["image_y"]))
                    elif key=="image_width" and drawing.get("crop_left") is not None and drawing.get("crop_right",1)>drawing["crop_left"]:
                        changes.append("crop_frame_width=?"); values.append(value/(drawing["crop_right"]-drawing["crop_left"]))
                    elif key=="image_height" and drawing.get("crop_top") is not None and drawing.get("crop_bottom",1)>drawing["crop_top"]:
                        changes.append("crop_frame_height=?"); values.append(value/(drawing["crop_bottom"]-drawing["crop_top"]))
        if data.guide_lines is not None:
            if len(data.guide_lines)>100 or any(item.get("axis") not in {"x","y"} or not isinstance(item.get("position"),(int,float)) or not 0<=item["position"]<=1 for item in data.guide_lines):
                raise HTTPException(422,"参考线数据无效")
            changes.append("guide_lines=?"); values.append(json.dumps(data.guide_lines,ensure_ascii=False))
        if not changes: raise HTTPException(422,"没有可更新的图纸字段")
        db.execute(f"UPDATE drawings SET {','.join(changes)} WHERE id=?", (*values,drawing_id))
        db.execute("UPDATE projects SET updated_at=? WHERE id=?", (now(), pid))
    return get_project(pid)


@app.post("/api/projects/{pid}/drawings/{drawing_id}/make-main")
def make_drawing_main(pid: str, drawing_id: str):
    """Promote a legacy page to a 16:9 floor canvas without changing point coordinates."""
    with connect() as db:
        drawing = one(db, "SELECT * FROM drawings WHERE id=? AND project_id=?", (drawing_id, pid))
        if drawing.get("canvas_mode") == "fixed":
            return get_project(pid)
        existing = db.execute(
            "SELECT id FROM drawings WHERE floor_id=? AND canvas_mode='fixed' AND id<>? LIMIT 1",
            (drawing["floor_id"], drawing_id),
        ).fetchone()
        if existing:
            raise HTTPException(409, "本楼层已有主图纸，请先切换到该主图纸")
        canvas_width, canvas_height = 1600, 900
        scale = min(canvas_width / drawing["width"], canvas_height / drawing["height"])
        image_width, image_height = drawing["width"] * scale, drawing["height"] * scale
        db.execute(
            """UPDATE drawings SET canvas_mode='fixed',canvas_width=?,canvas_height=?,
               image_x=?,image_y=?,image_width=?,image_height=?,image_rotation=0,
               guide_lines='[]',is_blank=0 WHERE id=?""",
            (canvas_width, canvas_height, (canvas_width-image_width)/2,
             (canvas_height-image_height)/2, image_width, image_height, drawing_id),
        )
        db.execute("UPDATE projects SET updated_at=? WHERE id=?", (now(), pid))
    return get_project(pid)


@app.post("/api/projects/{pid}/drawings/{drawing_id}/rotate")
def rotate_drawing(pid: str, drawing_id: str):
    with connect() as db:
        drawing=one(db,"SELECT * FROM drawings WHERE id=? AND project_id=?",(drawing_id,pid))
        background=UPLOADS/drawing["background_path"]
        try:
            with Image.open(background) as source:
                rotated=source.convert("RGB").transpose(Image.Transpose.ROTATE_270)
                temporary=background.with_name(background.stem+"_rotating.png")
                rotated.save(temporary,"PNG")
            temporary.replace(background)
        except Exception:
            raise HTTPException(500,"图纸旋转失败")
        for point in rows(db.execute("SELECT id,x,y,x2,y2 FROM points WHERE drawing_id=?",(drawing_id,))):
            x,y=1-point["y"],point["x"]
            if point["x2"] is None or point["y2"] is None:
                x2=y2=None
            else:
                x2,y2=1-point["y2"],point["x2"]
            db.execute("UPDATE points SET x=?,y=?,x2=?,y2=?,updated_at=? WHERE id=?",(x,y,x2,y2,now(),point["id"]))
        db.execute("UPDATE drawings SET width=?,height=?,rotation=? WHERE id=?",(drawing["height"],drawing["width"],(drawing.get("rotation",0)+90)%360,drawing_id))
        db.execute("UPDATE projects SET updated_at=? WHERE id=?",(now(),pid))
    return get_project(pid)


@app.post("/api/projects/{pid}/drawings/{drawing_id}/crop")
def crop_drawing(pid: str, drawing_id: str, data: DrawingCropIn):
    if data.right-data.left<0.05 or data.bottom-data.top<0.05: raise HTTPException(422,"裁剪区域至少保留原图的 5%")
    with connect() as db:
        drawing=one(db,"SELECT * FROM drawings WHERE id=? AND project_id=?",(drawing_id,pid))
        if drawing.get("canvas_mode")!="fixed": raise HTTPException(409,"只有 16:9 楼层画布支持裁剪主图纸")
        source_name=drawing.get("crop_source_path") or drawing["background_path"]
        source_path=UPLOADS/source_name
        if not source_path.is_file(): raise HTTPException(404,"主图纸原图文件不存在")
        with Image.open(source_path) as opened:
            source=opened.convert("RGB")
        width,height=source.size
        old_left=drawing.get("crop_left") if drawing.get("crop_left") is not None else 0.0
        old_top=drawing.get("crop_top") if drawing.get("crop_top") is not None else 0.0
        old_right=drawing.get("crop_right") if drawing.get("crop_right") is not None else 1.0
        old_bottom=drawing.get("crop_bottom") if drawing.get("crop_bottom") is not None else 1.0
        old_xspan,old_yspan=old_right-old_left,old_bottom-old_top
        if old_xspan<=0 or old_yspan<=0: raise HTTPException(409,"已有裁剪范围无效，请先恢复原图")
        base_x=drawing.get("crop_frame_x") if drawing.get("crop_frame_x") is not None else drawing.get("image_x",0)
        base_y=drawing.get("crop_frame_y") if drawing.get("crop_frame_y") is not None else drawing.get("image_y",0)
        base_width=drawing.get("crop_frame_width") if drawing.get("crop_frame_width") is not None else drawing.get("image_width",width)
        base_height=drawing.get("crop_frame_height") if drawing.get("crop_frame_height") is not None else drawing.get("image_height",height)
        points=rows(db.execute("SELECT id,x,y,x2,y2 FROM points WHERE drawing_id=?",(drawing_id,)))
        for point in points:
            for x,y in ((point["x"],point["y"]),(point.get("x2"),point.get("y2"))):
                if x is not None and y is not None:
                    original_x=old_left+x*old_xspan; original_y=old_top+y*old_yspan
                    if not (data.left<=original_x<=data.right and data.top<=original_y<=data.bottom):
                        raise HTTPException(409,"裁剪选区未完整包含所有点位，请先移动点位或扩大裁剪范围")
        x1=round(data.left*width);y1=round(data.top*height);x2=max(x1+1,round(data.right*width));y2=max(y1+1,round(data.bottom*height));box=(x1,y1,min(width,x2),min(height,y2))
        cropped=source.crop(box); new_path=UPLOADS/f"{uid()}_cropped.png";cropped.save(new_path,"PNG")
        xspan,yspan=data.right-data.left,data.bottom-data.top
        db.execute("UPDATE drawings SET background_path=?,width=?,height=?,image_x=?,image_y=?,image_width=?,image_height=?,crop_source_path=?,crop_source_width=?,crop_source_height=?,crop_frame_x=?,crop_frame_y=?,crop_frame_width=?,crop_frame_height=?,crop_left=?,crop_top=?,crop_right=?,crop_bottom=? WHERE id=?",(new_path.name,cropped.width,cropped.height,base_x+base_width*data.left,base_y+base_height*data.top,base_width*xspan,base_height*yspan,source_name,width,height,base_x,base_y,base_width,base_height,data.left,data.top,data.right,data.bottom,drawing_id))
        for point in points:
            original_x=old_left+point["x"]*old_xspan; original_y=old_top+point["y"]*old_yspan
            x=(original_x-data.left)/xspan;y=(original_y-data.top)/yspan
            original_x2=old_left+point["x2"]*old_xspan if point.get("x2") is not None else None
            original_y2=old_top+point["y2"]*old_yspan if point.get("y2") is not None else None
            x2=(original_x2-data.left)/xspan if original_x2 is not None else None
            y2=(original_y2-data.top)/yspan if original_y2 is not None else None
            db.execute("UPDATE points SET x=?,y=?,x2=?,y2=?,updated_at=? WHERE id=?",(x,y,x2,y2,now(),point["id"]))
        db.execute("UPDATE projects SET updated_at=? WHERE id=?",(now(),pid))
    return get_project(pid)


@app.delete("/api/projects/{pid}/drawings/{drawing_id}")
def delete_drawing(pid: str, drawing_id: str):
    with connect() as db:
        drawing = one(db, "SELECT * FROM drawings WHERE id=? AND project_id=?", (drawing_id, pid))
        db.execute("DELETE FROM drawings WHERE id=?", (drawing_id,))
        db.execute("UPDATE projects SET updated_at=? WHERE id=?", (now(), pid))
    for filename in {drawing.get("source_path"),drawing.get("background_path"),drawing.get("crop_source_path")}:
        if filename:
            path=UPLOADS/filename
            if path.is_file(): path.unlink()
    return {"ok": True}


@app.get("/api/files/{filename}")
def file(filename: str):
    if not re.fullmatch(r"[a-f0-9]{32}(?:_background|_blank|_cropped|_crop_source)?\.(?:png|jpg|jpeg|webp|pdf)", filename):
        raise HTTPException(404)
    path = UPLOADS / filename
    if not path.exists():
        path = IMAGE_DIR / filename
    if not path.exists():
        raise HTTPException(404)
    return FileResponse(path)


@app.get("/api/brand-files/{filename}")
def brand_file(filename: str):
    if not re.fullmatch(r"[a-f0-9]{32}(?:-[a-f0-9]{32})?\.(?:svg|png)", filename):
        raise HTTPException(404)
    path = BRAND_LOGO_DIR / filename
    if not path.is_file():
        raise HTTPException(404)
    media_type = "image/svg+xml" if path.suffix == ".svg" else "image/png"
    return FileResponse(path, media_type=media_type, headers={"X-Content-Type-Options": "nosniff"})


def product_image_file(filename: str | None) -> Path | None:
    if not filename or not re.fullmatch(r"[a-f0-9]{32}\.(?:png|jpg|jpeg|webp)", filename):
        return None
    for root in (IMAGE_DIR, UPLOADS):
        path = root / filename
        if path.is_file():
            return path
    return None


@app.get("/api/brands")
def brands():
    with connect() as db:
        result = rows(db.execute("SELECT * FROM brands ORDER BY name"))
        by_name = {b["name"]: b for b in result}
        for p in product_repository.list(db, include_disabled=True):
            if p["brand_name"] not in by_name:
                timestamp = now()
                db.execute("INSERT INTO brands(id,name,created_at,updated_at) VALUES(?,?,?,?) ON CONFLICT DO NOTHING", (p["brand_id"], p["brand_name"], timestamp, timestamp))
        return [brand_out(row) for row in rows(db.execute("SELECT * FROM brands ORDER BY name"))]


@app.post("/api/brands", status_code=201)
def create_brand(data: BrandIn):
    name = data.name.strip(); timestamp = now()
    b = {"id":stable_brand_id(name),"name":name,"english_name":data.english_name.strip(),"aliases":json.dumps([x.strip() for x in data.aliases if x.strip()],ensure_ascii=False),"enabled":int(data.enabled),"note":data.note.strip(),"created_at":timestamp,"updated_at":timestamp}
    with connect() as db:
        try:
            db.execute("INSERT INTO brands(id,name,english_name,aliases,enabled,note,created_at,updated_at) VALUES(:id,:name,:english_name,:aliases,:enabled,:note,:created_at,:updated_at)", b)
        except sqlite3.IntegrityError:
            raise HTTPException(409, "品牌已存在")
        if PATHS.local_data_mode:
            knowledge_repository.write_brand(knowledge_repository.brand_from_db(db, b["id"]))
    return brand_out(b)


@app.put("/api/brands/{brand_id_value}")
def update_brand(brand_id_value: str, data: BrandIn):
    with connect() as db:
        old = one(db, "SELECT * FROM brands WHERE id=?", (brand_id_value,))
        name = data.name.strip()
        try:
            db.execute("UPDATE brands SET name=?,english_name=?,aliases=?,enabled=?,note=?,updated_at=? WHERE id=?", (name,data.english_name.strip(),json.dumps([x.strip() for x in data.aliases if x.strip()],ensure_ascii=False),int(data.enabled),data.note.strip(),now(),brand_id_value))
        except sqlite3.IntegrityError:
            raise HTTPException(409, "品牌名称已存在")
        # 品牌 ID 是产品的稳定关联；改名只更新展示快照，不改变关联 ID。
        for product in product_repository.list(db, include_disabled=True):
            if product["brand_id"] == brand_id_value or product["brand_name"] == old["name"]:
                saved = product_repository.save({**product, "brand_id": brand_id_value, "brand_name": name, "updated_at": now()})
                ensure_shadow_product(db, saved)
        if PATHS.local_data_mode:
            knowledge_repository.write_brand(knowledge_repository.brand_from_db(db, brand_id_value))
        return brand_out(one(db, "SELECT * FROM brands WHERE id=?", (brand_id_value,)))


@app.post("/api/brands/{brand_id_value}/logo")
async def upload_brand_logo(brand_id_value: str, file: UploadFile = File(...)):
    with connect() as db:
        brand = one(db, "SELECT * FROM brands WHERE id=?", (brand_id_value,))
    content = await file.read()
    if len(content) > 2 * 1024 * 1024:
        raise HTTPException(413, "品牌 Logo 不能超过 2 MB")
    ext = Path(file.filename or "").suffix.lower()
    if ext not in (".svg", ".png"):
        raise HTTPException(422, "品牌 Logo 仅支持 SVG 或透明 PNG")
    try:
        if ext == ".svg":
            root = ET.fromstring(content)
            if root.tag.split("}")[-1].lower() != "svg" or any(node.tag.split("}")[-1].lower() in {"script", "foreignobject"} for node in root.iter()):
                raise ValueError()
        else:
            with Image.open(io.BytesIO(content)) as image:
                image.verify()
            with Image.open(io.BytesIO(content)) as image:
                if image.format != "PNG" or image.width * image.height > 16_000_000:
                    raise ValueError()
    except Exception:
        raise HTTPException(422, "品牌 Logo 文件无法读取或包含不安全内容")
    filename = f"{brand_id_value}-{uid()}{ext}"
    path = BRAND_LOGO_DIR / filename
    temporary = BRAND_LOGO_DIR / f".{filename}.uploading"
    temporary.write_bytes(content); temporary.replace(path)
    column = "logo_svg_path" if ext == ".svg" else "logo_png_path"
    with connect() as db:
        db.execute("UPDATE brands SET logo_svg_path=?,logo_png_path=?,updated_at=? WHERE id=?", (filename if ext == ".svg" else None, filename if ext == ".png" else None, now(),brand_id_value))
        if PATHS.local_data_mode:
            knowledge_repository.write_brand(knowledge_repository.brand_from_db(db, brand_id_value))
        return brand_out(one(db, "SELECT * FROM brands WHERE id=?", (brand_id_value,)))


@app.get("/api/brands/{brand_id_value}/products")
def brand_products(brand_id_value: str):
    with connect() as db:
        brand = one(db, "SELECT * FROM brands WHERE id=?", (brand_id_value,))
        return [product_out(p, db) for p in product_repository.list(db, include_disabled=True)
                if p["brand_id"] == brand_id_value or (find_brand(db, name=p["brand_name"]) or {}).get("id") == brand_id_value]


@app.delete("/api/brands/{brand_id_value}/logo")
def delete_brand_logo(brand_id_value: str):
    with connect() as db:
        one(db, "SELECT * FROM brands WHERE id=?", (brand_id_value,))
        db.execute("UPDATE brands SET logo_svg_path=NULL,logo_png_path=NULL,updated_at=? WHERE id=?", (now(),brand_id_value))
        knowledge_repository.write_brand(knowledge_repository.brand_from_db(db, brand_id_value))
        return brand_out(one(db, "SELECT * FROM brands WHERE id=?", (brand_id_value,)))


@app.delete("/api/brands/{brand_id_value}")
def delete_brand(brand_id_value: str):
    linked = brand_products(brand_id_value)
    if linked:
        raise HTTPException(409, f"该品牌仍有 {len(linked)} 个产品正在使用。请先处理产品关联。")
    with connect() as db:
        one(db, "SELECT * FROM brands WHERE id=?", (brand_id_value,))
        # Product notes are authoritative; stale index rows are disposable.
        db.execute("DELETE FROM products WHERE brand_id=?", (brand_id_value,))
        note = knowledge_repository.brand_path(brand_id_value)
        if note and note.exists():
            target = DATA / "backups" / "deleted-brands"
            target.mkdir(parents=True, exist_ok=True)
            shutil.copy2(note, target / f"{brand_id_value}-{uid()}.md")
            note.unlink()
        db.execute("DELETE FROM brands WHERE id=?", (brand_id_value,))
    return {"ok": True}


@app.get("/api/products")
def products(response: Response, q: str = "", brand_id: str = "", series: str = "", category: str = "", all_categories: bool = False, include_disabled: bool = False, refresh: bool = False):
    response.headers["Cache-Control"] = "no-store, max-age=0"
    if refresh:
        product_repository.normalize_layout()
    with connect() as db:
        return [product_out(p,db) for p in product_repository.list(db,q=q,brand_id=brand_id,series=series,category="" if all_categories else category,include_disabled=include_disabled)]


def resolve_brand(db, data: ProductIn):
    if data.brand_id:
        row = find_brand(db, brand_id_value=data.brand_id)
        if row:
            return row
    if data.brand_name.strip():
        row = find_brand(db, name=data.brand_name)
        if row:
            return row
        name = data.brand_name.strip(); timestamp = now(); new_id = stable_brand_id(name)
        db.execute("INSERT INTO brands(id,name,created_at,updated_at) VALUES(?,?,?,?) ON CONFLICT DO NOTHING", (new_id,name,timestamp,timestamp))
        return find_brand(db, brand_id_value=new_id, name=name)
    if data.brand_id:
        for p in product_repository.list(db, include_disabled=True):
            if p["brand_id"] == data.brand_id:
                return {"id": data.brand_id, "name": p["brand_name"]}
    raise HTTPException(422, "请填写品牌")


def ensure_shadow_product(db, p):
    """Keep old point FK valid; all catalog values are read from Markdown."""
    # Stable brand IDs use casefold(), so ABB and abb share an ID.
    # Reuse that row before checking legacy brands with non-stable IDs.
    existing_brand = db.execute("SELECT id FROM brands WHERE id=?", (p["brand_id"],)).fetchone()
    if existing_brand is None:
        existing_brand = next((row for row in db.execute("SELECT id,name FROM brands")
                               if row["name"].strip().casefold() == p["brand_name"].strip().casefold()), None)
    if existing_brand is None:
        db.execute("INSERT INTO brands(id,name,created_at) VALUES(?,?,?) ON CONFLICT DO NOTHING",
                   (p["brand_id"], p["brand_name"], now()))
        existing_brand = db.execute("SELECT id FROM brands WHERE id=? OR name=?",
                                    (p["brand_id"], p["brand_name"])).fetchone()
    shadow_brand_id = existing_brand[0]
    shadow_category = p.get("device_type") or category_type(p["category"])
    if shadow_category not in TYPE_MAP:
        shadow_category = "other"
    db.execute("""INSERT INTO products(id,brand_id,name,series,model,category,standard_cents,unit,params,description,image_path,enabled,updated_at)
      VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(id) DO UPDATE SET brand_id=excluded.brand_id,name=excluded.name,series=excluded.series,model=excluded.model,category=excluded.category,standard_cents=excluded.standard_cents,unit=excluded.unit,params=excluded.params,description=excluded.description,image_path=excluded.image_path,enabled=excluded.enabled,updated_at=excluded.updated_at""",
      (p["id"], shadow_brand_id, p["name"], p["series"], p["model"], shadow_category, p["standard_cents"], p["unit"], p["params"], p["description"], p["image_path"], p["enabled"], p["updated_at"]))


def migrate_legacy_catalog(db_path: Path, image_root: Path | None = None) -> dict:
    """Import products from an earlier SQLite-only catalog into shared Obsidian notes."""
    report={"source":str(db_path),"imported":0,"skipped":0,"error":""}
    if not db_path.is_file():
        return report
    try:
        source=sqlite3.connect(f"file:{db_path.as_posix()}?mode=ro",uri=True)
        source.row_factory=sqlite3.Row
        tables={r[0] for r in source.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        if not {"products","brands"}.issubset(tables):
            source.close(); return report
        legacy=source.execute("SELECT p.*,b.name brand_name FROM products p JOIN brands b ON b.id=p.brand_id").fetchall()
        source.close()
        for row in legacy:
            item=dict(row)
            if product_repository.get(None,item["id"]):
                report["skipped"]+=1; continue
            device_type=canonical_device_type(item.get("category") or "other")
            category=canonical_product_category(device_type)
            image=item.get("image_path")
            if image and not (IMAGE_DIR/image).is_file():
                candidates=[root/image for root in (image_root,db_path.parent/"uploads") if root]
                source_image=next((p for p in candidates if p.is_file()),None)
                if source_image:
                    shutil.copy2(source_image,IMAGE_DIR/image)
                else:
                    image=None
            product={
                "id":item["id"],"brand_id":stable_brand_id(item["brand_name"]),"brand_name":item["brand_name"],
                "name":item["name"],"series":item.get("series") or "","model":item.get("model") or "",
                "category":category,"device_type":device_type,"system_category":"uncategorized","standard_cents":item.get("standard_cents") or 0,
                "cost_cents":item.get("cost_cents"),"unit":item.get("unit") or "个","origin":item.get("origin") or "",
                "params":item.get("params") or "","description":item.get("description") or "","image_path":image,
                "enabled":int(item.get("enabled",1)),"updated_at":item.get("updated_at") or now(),
            }
            saved=product_repository.save(product)
            with connect() as target:
                ensure_shadow_product(target,saved)
            report["imported"]+=1
    except Exception as error:
        report["error"]=str(error)
    return report


def migrate_known_legacy_catalogs():
    # Local data/ is authoritative. Legacy caches must not add
    # records to it, even when an old SQLite file was copied with the package.
    # Do not move product notes automatically before a data backup exists.
    if PATHS.local_data_mode:
        return
    candidates=[(DB,UPLOADS)]
    configured=os.environ.get("POINT_QUOTE_LEGACY_DB")
    if configured:
        legacy=Path(configured).expanduser().resolve()
        if legacy != DB.resolve(): candidates.append((legacy,legacy.parent/"uploads"))
    for path,images in candidates:
        report=migrate_legacy_catalog(path,images)
        if report["imported"] or report["error"]:
            MIGRATION_REPORT.append(report)
    try:
        normalized=product_repository.normalize_layout()
        if normalized:
            MIGRATION_REPORT.append({"source":str(PRODUCT_DIR),"normalized":normalized})
    except OSError as error:
        error_type=record_exception("vault_startup",error,operation="scan",filePath=str(PRODUCT_DIR))
        MIGRATION_REPORT.append({"source":str(PRODUCT_DIR),"errorType":error_type,"error":str(error)})


def product_values(data: ProductIn, *, pid: str, brand: dict, image_path: str | None = None):
    category = canonical_category(data.category)
    raw_device_type = data.category if data.device_type == "other" and data.category not in PRODUCT_CATEGORIES else data.device_type
    device_type = canonical_device_type(raw_device_type, category)
    return {"id": pid, "brand_id": brand["id"], "brand_name": brand["name"], "name": data.name.strip(), "series": data.series.strip(), "model": data.model.strip(),
            "category": category, "device_type": device_type, "system_category": canonical_system_category(data.system_category), "standard_cents": money(data.standard_price),
            "cost_cents": money(data.cost_price) if data.cost_price not in (None, "") else None,
            "unit": data.unit.strip(), "origin": data.origin.strip(), "params": data.params, "description": data.description,
            "image_path": image_path, "enabled": int(data.enabled), "updated_at": now()}


migrate_known_legacy_catalogs()


def initialize_local_data():
    """Restore records exclusively from the current tool folder's data/."""
    with connect() as db:
        knowledge_repository.ensure_tables(db)
        try:
            if not PRODUCT_DIR.is_dir() or not BRAND_DIR.is_dir() or not PROJECT_DIR.is_dir():
                raise FileNotFoundError("本地 data/ 目录结构不完整")
            restored = knowledge_repository.restore_local_records(db)
            for product in product_repository.list(db, include_disabled=True):
                ensure_shadow_product(db, product)
            restored["projects"] = knowledge_repository.restore_local_projects(db)
            return {"conflicts": [], **restored, "metrics": {"local_data_restore": True}}
        except OSError as error:
            error_type = record_exception("local_data_startup", error, operation="restore_local_data",
                                          filePath=getattr(error, "filename", None))
            return {"conflicts": [], "brands": 0, "projects": 0,
                    "metrics": {"local_data_restore": True, "startup_restore_failed": error_type}}


LOCAL_DATA_RESTORE_REPORT = initialize_local_data()

def _persist_local_project(project_id: str) -> None:
    try:
        with connect() as db:
            snapshot = knowledge_repository.snapshot_project(db, project_id)
            if snapshot:
                knowledge_repository.write_project(snapshot)
    except Exception:
        # A failed local Markdown write must not leave SQLite showing an
        # unsaved project state that would disappear on the next launch.
        path = knowledge_repository.project_path(project_id)
        snapshot = knowledge_repository.read_project(path) if path else None
        with connect() as db:
            if snapshot:
                knowledge_repository.restore_project(db, snapshot)
            else:
                db.execute("DELETE FROM projects WHERE id=?", (project_id,))
        raise


async def _affected_local_project(request) -> str | None:
    if request.method not in {"POST", "PATCH", "DELETE"}:
        return None
    project_match = re.match(r"/api/projects/([a-f0-9]{32})(?:/(.*))?$", request.url.path)
    if project_match:
        project_id, suffix = project_match.groups()
        local_only = (
            (suffix is None and request.method == "PATCH")
            or suffix == "signed"
            or (suffix == "quote" and request.method == "PATCH")
            or (suffix == "quotation-v2" and request.method == "PATCH")
            or (suffix and (suffix.startswith("quote-items") or suffix == "fees"
                            or suffix.startswith("quotation-v2/items")
                            or suffix in {"quotation-v2/reorder"} or suffix.startswith("quotation-v2/versions")))
            or (suffix and re.fullmatch(r"drawings/[a-f0-9]{32}", suffix) and request.method == "PATCH")
        )
        if local_only:
            return project_id
    match = re.fullmatch(r"/api/projects/([a-f0-9]{32})/points", request.url.path)
    if match and request.method == "POST":
        return match.group(1)
    match = re.fullmatch(r"/api/points/([a-f0-9]{32})(?:/duplicate|/restore-price)?", request.url.path)
    if match:
        with connect() as db:
            row = db.execute("SELECT project_id FROM points WHERE id=?", (match.group(1),)).fetchone()
            return row[0] if row else None
    if request.method == "POST" and request.url.path == "/api/points/restore":
        try:
            payload = await request.json()
            project_id = payload.get("project_id")
            return project_id if isinstance(project_id, str) and re.fullmatch(r"[a-f0-9]{32}", project_id) else None
        except (ValueError, AttributeError):
            return None
    if request.method == "POST" and request.url.path == "/api/points/assign-product":
        try:
            payload = await request.json()
            ids = payload.get("point_ids", [])
            if not isinstance(ids, list) or not ids:
                return None
            placeholders = ",".join("?" for _ in ids)
            with connect() as db:
                projects = {row[0] for row in db.execute(
                    f"SELECT DISTINCT project_id FROM points WHERE id IN ({placeholders})", ids
                )}
            return next(iter(projects)) if len(projects) == 1 else None
        except (ValueError, AttributeError):
            return None
    match = re.fullmatch(r"/api/quotation-systems/([a-f0-9]{32})", request.url.path)
    if match:
        with connect() as db:
            row = db.execute("SELECT q.project_id FROM quotation_systems s JOIN quotes q ON q.id=s.quote_id WHERE s.id=?", (match.group(1),)).fetchone()
            return row[0] if row else None
    match = re.fullmatch(r"/api/quotation-items/([a-f0-9]{32})(?:/restore-source-quantity)?", request.url.path)
    if match and not request.url.path.endswith("/image"):
        with connect() as db:
            row = db.execute("SELECT q.project_id FROM quotation_items i JOIN quotes q ON q.id=i.quote_id WHERE i.id=?", (match.group(1),)).fetchone()
            return row[0] if row else None
    return None


@app.middleware("http")
async def local_data_requests(request, call_next):
    """Persist successful business changes only inside the current tool's data/."""
    if not request.url.path.startswith("/api/"):
        return await call_next(request)
    request_id = request.headers.get("X-Request-ID") or f"REQ-{uuid.uuid4().hex[:12]}"
    token = set_request_id(request_id)
    started = time.perf_counter()
    try:
        affected_project_id = await _affected_local_project(request)
        response = await call_next(request)
        if request.method not in {"GET", "HEAD", "OPTIONS"} and response.status_code < 400:
            project_match = re.match(r"/api/projects/([a-f0-9]{32})(?:/|$)", request.url.path)
            project_id = project_match.group(1) if project_match else affected_project_id
            if project_id and not (request.method == "DELETE" and request.url.path == f"/api/projects/{project_id}"):
                _persist_local_project(project_id)
        response.headers["X-Request-ID"] = request_id
        record_event("api_request", operation="write" if request.method not in {"GET", "HEAD", "OPTIONS"} else "read",
                     durationMs=round((time.perf_counter() - started) * 1000, 2),
                     success=response.status_code < 500, statusCode=response.status_code,
                     route=request.url.path, method=request.method)
        return response
    except Exception as error:
        error_type = record_exception("api_request", error, operation="write",
            filePath=getattr(error, "filename", None), route=request.url.path, method=request.method)
        return JSONResponse(status_code=500, headers={"X-Request-ID": request_id},
                            content={"detail":"本地保存失败，诊断日志已记录。","error_code":error_type,"request_id":request_id})
    finally:
        reset_request_id(token)


@app.get("/api/knowledge-sync/status")
def knowledge_sync_status():
    with connect() as db:
        return {"brand_notes": db.execute("SELECT count(*) FROM brands").fetchone()[0],
                "project_notes": db.execute("SELECT count(*) FROM projects").fetchone()[0],
                "conflicts": [], "sync_status": "local", "pending_count": 0,
                "failure_count": 0, "recent_error": "", "error_type": ""}


@app.post("/api/knowledge-sync/retry")
def retry_knowledge_sync():
    # Retained for older clients; it never contacts a remote knowledge base.
    return knowledge_sync_status()


@app.get("/api/diagnostics/status")
def diagnostics_status():
    return {"log_path": str(LOG_FILE), "exists": LOG_FILE.is_file(), "size_bytes": LOG_FILE.stat().st_size if LOG_FILE.is_file() else 0}


@app.post("/api/products", status_code=201)
def create_product(data: ProductIn):
    with connect() as db:
        brand = resolve_brand(db, data)
        saved = product_repository.save(product_values(data, pid=uid(), brand=brand))
        ensure_shadow_product(db, saved)
        return product_out(saved,db)


@app.get("/api/products/{product_id}")
def get_product(product_id: str):
    with connect() as db:
        p=product_repository.get(db,product_id)
        if not p: raise HTTPException(404,"产品不存在")
        return product_out(p,db)


@app.put("/api/products/{product_id}")
def update_product(product_id: str, data: ProductIn):
    with connect() as db:
        old = product_repository.get(db, product_id)
        if not old: raise HTTPException(404, "产品笔记不存在")
        brand = resolve_brand(db, data)
        saved = product_repository.save(product_values(data, pid=product_id, brand=brand, image_path=old["image_path"]))
        ensure_shadow_product(db, saved)
        return product_out(saved,db)


@app.delete("/api/products/{product_id}")
def delete_product(product_id: str):
    with connect() as db:
        old = product_repository.get(db, product_id)
        if not old: raise HTTPException(404, "产品笔记不存在")
        try:
            product_repository.delete(product_id)
        except OSError as error:
            raise HTTPException(500, f"无法删除知识库产品文件，请检查产品目录写入权限：{error}")
        db.execute("DELETE FROM products WHERE id=?", (product_id,))
    return {"ok":True}


@app.post("/api/products/{product_id}/image")
async def upload_product_image(product_id: str, file: UploadFile = File(...)):
    with connect() as db:
        old = product_repository.get(db, product_id)
        if not old: raise HTTPException(404, "产品笔记不存在")
    content = await file.read()
    if len(content)>5*1024*1024: raise HTTPException(413,"图片不能超过 5 MB")
    ext=Path(file.filename or "").suffix.lower()
    if ext not in (".jpg",".jpeg",".png",".webp"): raise HTTPException(422,"仅支持 JPG、PNG、WebP")
    try:
        with Image.open(io.BytesIO(content)) as im:
            im.verify()
        with Image.open(io.BytesIO(content)) as im:
            if im.width*im.height>25000000: raise ValueError()
    except Exception:
        raise HTTPException(422,"图片无法读取")
    # 原始文件按原格式保存；透明 PNG/WebP 不做 RGB 转换，显示与导出时再铺白底。
    path=IMAGE_DIR/f"{uid()}{ext}"
    temporary=IMAGE_DIR/f".{path.name}.uploading"
    try:
        temporary.write_bytes(content)
        temporary.replace(path)
    except OSError as error:
        try:
            temporary.unlink(missing_ok=True)
        except OSError:
            pass
        raise HTTPException(500, f"无法把产品图片写入知识库图片目录，请检查 09_附件/产品图片 的写入权限：{error}")
    try:
        with connect() as db:
            saved = product_repository.save({**old,"image_path":path.name,"updated_at":now()})
            ensure_shadow_product(db,saved)
            return product_out(saved,db)
    except OSError as error:
        try:
            path.unlink(missing_ok=True)
        except OSError:
            pass
        raise HTTPException(500, f"图片已收到，但无法更新知识库产品笔记，请检查 01_产品库 的修改权限：{error}")


@app.patch("/api/products/batch")
def batch_products(data: ProductBatchIn):
    if data.action not in ("enable", "disable", "delete", "category", "system_category"):
        raise HTTPException(422, "批量操作不存在")
    if data.action == "category" and not (data.category or "").strip():
        raise HTTPException(422, "请选择分类")
    if data.action == "system_category" and canonical_system_category(data.system_category) == "uncategorized" and data.system_category != "uncategorized":
        raise HTTPException(422, "请选择所属系统")
    updated = 0
    with connect() as db:
        for product_id in dict.fromkeys(data.ids):
            old = product_repository.get(db, product_id)
            if not old:
                continue
            if data.action == "delete":
                try:
                    product_repository.delete(product_id)
                except OSError as error:
                    raise HTTPException(500, f"无法删除知识库产品文件，请检查产品目录写入权限：{error}")
                db.execute("DELETE FROM products WHERE id=?", (product_id,))
                updated += 1
                continue
            values = {**old, "updated_at": now()}
            if data.action == "enable": values["enabled"] = 1
            elif data.action == "disable": values["enabled"] = 0
            elif data.action == "category":
                values["category"] = canonical_category(data.category or "")
                forced={"影音":"av","灯光":"light","网络":"network"}.get(values["category"])
                values["device_type"] = forced or old.get("device_type") or "other"
            else:
                values["system_category"] = canonical_system_category(data.system_category)
            saved = product_repository.save(values)
            ensure_shadow_product(db, saved)
            updated += 1
    return {"updated": updated}


@app.get("/api/products/export.xlsx")
def export_products():
    with connect() as db:
        data=product_repository.list(db,include_disabled=True)
    wb=Workbook(); ws=wb.active; ws.title="产品库"
    headers=["产品ID","品牌","产品名称","型号","分类","点位类型","所属系统","售价","进货价","产地","单位","参数","说明","启用","笔记路径"]
    ws.append(headers)
    for p in data: ws.append([p["id"],p["brand_name"],p["name"],p.get("model") or "",p["category"],p["device_type"],p["system_category"],p["standard_cents"]/100,p["cost_cents"]/100 if p["cost_cents"] is not None else None,p["origin"],p["unit"],p["params"],p["description"],"是" if p["enabled"] else "否",p["note_path"]])
    for cell in ws[1]: cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="26324A")
    for col in "ABCDEFGHIJKLMNO": ws.column_dimensions[col].width=22 if col not in "ACDM" else 36
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
    output=io.BytesIO(); wb.save(output); output.seek(0)
    return StreamingResponse(output,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":"attachment; filename=products.xlsx"})


@app.post("/api/products/import")
async def import_products(file: UploadFile = File(...)):
    content=await file.read(); ext=Path(file.filename or "").suffix.lower()
    if len(content)>10*1024*1024 or ext not in (".csv",".xlsx"): raise HTTPException(422,"请上传 10MB 内的 CSV 或 XLSX")
    try:
        if ext==".csv":
            records=list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
        else:
            wb=load_workbook(io.BytesIO(content),read_only=True,data_only=True); vals=list(wb.active.values)
            records=[dict(zip(vals[0],r)) for r in vals[1:] if any(v is not None for v in r)]
    except Exception:
        raise HTTPException(422,"导入文件无法解析")
    if len(records)>5000: raise HTTPException(422,"单次最多导入 5000 行")
    required={"品牌","产品名称","分类"}
    if records and not required.issubset(records[0]): raise HTTPException(422,"缺少必需列：品牌、产品名称、分类")
    if records and "标准价格" not in records[0] and "售价" not in records[0]: raise HTTPException(422,"缺少售价或标准价格列")
    prepared=[]
    for i,r in enumerate(records,2):
        try:
            brand=str(r.get("品牌") or "").strip(); name=str(r.get("产品名称") or "").strip(); category=str(r.get("分类") or "").strip(); model=str(r.get("型号") or "").strip()
            if not brand or not name: raise ValueError()
            cents=money(r.get("售价") if r.get("售价") is not None else r.get("标准价格"))
            cost=money(r.get("进货价")) if r.get("进货价") not in (None,"") else None
            category=canonical_category(category)
            raw_device_type=str(r.get("点位类型") or r.get("设备类型") or str(r.get("分类") or ""))
            prepared.append({"id":uid(),"brand_name":brand,"name":name,"series":"","model":model,"category":category,"device_type":canonical_device_type(raw_device_type,category),"system_category":canonical_system_category(r.get("所属系统")),"standard_cents":cents,"cost_cents":cost,"origin":str(r.get("产地") or ""),"unit":str(r.get("单位") or "个"),"params":str(r.get("参数") or ""),"description":str(r.get("说明") or ""),"image_path":None,"enabled":1,"updated_at":now()})
        except Exception:
            raise HTTPException(422,f"第 {i} 行数据有误；未导入任何产品")
    with connect() as db:
        for p in prepared:
            saved=product_repository.save(p)
            ensure_shadow_product(db,saved)
    return {"imported":len(prepared)}


def next_number(db, project_id, type_id):
    prefix=TYPE_MAP[type_id][2]
    row=db.execute("SELECT last_number FROM point_counters WHERE project_id=? AND type_id=?",(project_id,type_id)).fetchone()
    last=row[0] if row else 0
    if not row:
        existing=[int(r[0][len(prefix):]) for r in db.execute("SELECT number FROM points WHERE project_id=? AND type_id=?",(project_id,type_id)) if r[0].startswith(prefix) and r[0][len(prefix):].isdigit()]
        last=max(existing,default=0)
    next_value=last+1
    db.execute("INSERT INTO point_counters(project_id,type_id,last_number) VALUES(?,?,?) ON CONFLICT(project_id,type_id) DO UPDATE SET last_number=excluded.last_number",(project_id,type_id,next_value))
    return f"{prefix}{next_value:02d}"


@app.post("/api/projects/{pid}/points", status_code=201)
def create_point(pid: str, data: PointIn, drawing_id: str):
    started = time.perf_counter()
    if data.type_id not in TYPE_MAP: raise HTTPException(422,"设备类型不存在")
    with connect() as db:
        drawing=one(db,"SELECT id,floor_id FROM drawings WHERE id=? AND project_id=?",(drawing_id,pid))
        if data.layer_id:
            one(db,"SELECT id FROM layers WHERE id=? AND floor_id=?",(data.layer_id,drawing["floor_id"]))
            layer_id=data.layer_id
        else:
            layer_id=one(db,"SELECT id FROM layers WHERE floor_id=? ORDER BY sort_order LIMIT 1",(drawing["floor_id"],))["id"]
        if one(db,"SELECT locked FROM layers WHERE id=?",(layer_id,))["locked"]: raise HTTPException(409,"当前图层已锁定")
        product=None; snapshot=None
        if data.product_id:
            # The product catalog is loaded into the process index at startup/page
            # load. Do not stat or recursively scan the SMB catalog on a canvas click.
            product=product_repository.get_cached(data.product_id)
            if not product or not product["enabled"]: raise HTTPException(404,"产品不存在或已停用")
            if data.type_id != "other" and product.get("device_type")!=data.type_id: raise HTTPException(422,"产品的设备点位与当前工具不匹配")
            ensure_shadow_product(db,product)
            snapshot={k:product.get(k) for k in ("id","name","brand_id","brand_name","series","model","category","device_type","system_category","standard_cents","unit","origin","params","image_path")}
        number=next_number(db,pid,data.type_id); point_id=uid()
        shape = point_shape(data.type_id)
        x2 = data.x2 if data.x2 is not None else (min(1, data.x + .12) if shape != "point" else None)
        y2 = data.y2 if data.y2 is not None else (data.y if shape in {"line","sector"} else min(1, data.y + .08) if shape == "area" else None)
        standard=product["standard_cents"] if product else None
        db.execute("INSERT INTO points(id,project_id,drawing_id,type_id,number,name,room,quantity,note,x,y,size,color,locked,product_id,product_snapshot,standard_cents,actual_cents,created_at,updated_at,shape,x2,y2,show_motor,fov_angle,curtain_style,curtain_length) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",(point_id,pid,drawing_id,data.type_id,number,data.name or number,data.room,data.quantity,data.note,data.x,data.y,data.size,data.color or TYPE_MAP[data.type_id][3],int(data.locked),product["id"] if product else None,json.dumps(snapshot,ensure_ascii=False) if snapshot else None,standard,standard,now(),now(),shape,x2,y2,int(data.show_motor),data.fov_angle,data.curtain_style,data.curtain_length))
        db.execute("UPDATE points SET layer_id=?,line_width=? WHERE id=?",(layer_id,data.line_width,point_id))
        sync_line(db,point_id)
        db.execute("UPDATE projects SET updated_at=? WHERE id=?",(now(),pid))
        result = point_out(one(db,"SELECT * FROM points WHERE id=?",(point_id,)))
    record_event("addPoint.save", operation="local_commit", projectId=pid,
        pointId=point_id, durationMs=round((time.perf_counter()-started)*1000, 2), success=True)
    return result


@app.patch("/api/points/{point_id}")
def update_point(point_id: str, data: PointPatch):
    updates=data.model_dump(exclude_unset=True); price=updates.pop("actual_price",None)
    with connect() as db:
        old=one(db,"SELECT * FROM points WHERE id=?",(point_id,))
        if old.get("layer_id") and one(db,"SELECT locked FROM layers WHERE id=?",(old["layer_id"],))["locked"] and updates:
            raise HTTPException(409,"点位所在图层已锁定")
        if "layer_id" in updates:
            floor=one(db,"SELECT floor_id FROM drawings WHERE id=?",(old["drawing_id"],))["floor_id"]
            one(db,"SELECT id FROM layers WHERE id=? AND floor_id=?",(updates["layer_id"],floor))
        if old["locked"] and any(k in updates for k in ("x","y","x2","y2")): raise HTTPException(409,"点位已锁定")
        if "type_id" in updates and updates["type_id"] not in TYPE_MAP: raise HTTPException(422,"设备类型不存在")
        if "color" in updates and not re.fullmatch(r"#[0-9a-fA-F]{6}",updates["color"] or ""): raise HTTPException(422,"颜色格式错误")
        if "type_id" in updates and updates["type_id"]!=old["type_id"]:
            updates["number"]=next_number(db,old["project_id"],updates["type_id"])
            if old["name"]==old["number"]: updates["name"]=updates["number"]
            updates["color"]=TYPE_MAP[updates["type_id"]][3]
            updates["shape"] = point_shape(updates["type_id"])
            if updates["shape"]=="point": updates.update(x2=None,y2=None)
            elif updates["shape"] in {"line","sector"}: updates.update(x2=min(1,old["x"]+.12),y2=old["y"])
            else: updates.update(x2=min(1,old["x"]+.12),y2=min(1,old["y"]+.08))
            updates.update(product_id=None,product_snapshot=None,standard_cents=None,actual_cents=None)
            db.execute("DELETE FROM quote_lines WHERE point_id=?",(point_id,))
        if price is not None:
            if not old["product_snapshot"] and "product_snapshot" not in updates: raise HTTPException(422,"请先配置产品")
            updates["actual_cents"]=money(price)
        updates["updated_at"]=now()
        if updates:
            db.execute("UPDATE points SET "+",".join(f"{k}=?" for k in updates)+" WHERE id=?",[*[int(v) if isinstance(v,bool) else v for v in updates.values()],point_id])
        sync_line(db,point_id)
        db.execute("UPDATE projects SET updated_at=? WHERE id=?",(now(),old["project_id"]))
        return point_out(one(db,"SELECT * FROM points WHERE id=?",(point_id,)))


@app.delete("/api/points/{point_id}")
def delete_point(point_id: str):
    with connect() as db:
        p=one(db,"SELECT * FROM points WHERE id=?",(point_id,))
        if p.get("layer_id") and one(db,"SELECT locked FROM layers WHERE id=?",(p["layer_id"],))["locked"]: raise HTTPException(409,"点位所在图层已锁定")
        db.execute("DELETE FROM points WHERE id=?",(point_id,))
        db.execute("UPDATE projects SET updated_at=? WHERE id=?",(now(),p["project_id"]))
    return {"ok":True}


@app.post("/api/points/restore")
def restore_point(data: PointRestore):
    if data.type_id not in TYPE_MAP or not re.fullmatch(r"#[0-9a-fA-F]{6}",data.color): raise HTTPException(422,"点位数据无效")
    with connect() as db:
        one(db,"SELECT id FROM drawings WHERE id=? AND project_id=?",(data.drawing_id,data.project_id))
        floor=one(db,"SELECT floor_id FROM drawings WHERE id=?",(data.drawing_id,))["floor_id"]
        layer_id=data.layer_id or one(db,"SELECT id FROM layers WHERE floor_id=? ORDER BY sort_order LIMIT 1",(floor,))["id"]
        one(db,"SELECT id FROM layers WHERE id=? AND floor_id=?",(layer_id,floor))
        old=db.execute("SELECT created_at FROM points WHERE id=?",(data.id,)).fetchone()
        values=(data.id,data.project_id,data.drawing_id,data.type_id,data.number,data.name,data.room,data.quantity,data.note,data.x,data.y,data.size,data.color,data.locked,data.product_id,json.dumps(data.product_snapshot,ensure_ascii=False) if data.product_snapshot else None,money(data.standard_price) if data.standard_price is not None else None,money(data.actual_price) if data.actual_price is not None else None,old[0] if old else now(),now(),data.shape,data.x2,data.y2,int(data.show_motor),data.fov_angle,data.curtain_style,data.curtain_length)
        db.execute("INSERT OR REPLACE INTO points(id,project_id,drawing_id,type_id,number,name,room,quantity,note,x,y,size,color,locked,product_id,product_snapshot,standard_cents,actual_cents,created_at,updated_at,shape,x2,y2,show_motor,fov_angle,curtain_style,curtain_length) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",values)
        db.execute("UPDATE points SET layer_id=?,line_width=? WHERE id=?",(layer_id,data.line_width,data.id))
        sync_line(db,data.id)
        db.execute("UPDATE projects SET updated_at=? WHERE id=?",(now(),data.project_id))
        return point_out(one(db,"SELECT * FROM points WHERE id=?",(data.id,)))


@app.post("/api/points/paste")
def paste_point(data: PointRestore):
    with connect() as db:
        one(db, "SELECT id FROM drawings WHERE id=? AND project_id=?", (data.drawing_id,data.project_id))
        number = next_number(db,data.project_id,data.type_id)
    copied = data.model_copy(update={"id":uid(), "number":number, "name":number if data.name==data.number else data.name})
    result = restore_point(copied)
    return result


@app.post("/api/points/{point_id}/duplicate")
def duplicate_point(point_id: str):
    with connect() as db:
        p=one(db,"SELECT * FROM points WHERE id=?",(point_id,))
        new_id=uid(); number=next_number(db,p["project_id"],p["type_id"])
        dx=.03; dy=.03
        db.execute("INSERT INTO points(id,project_id,drawing_id,type_id,number,name,room,quantity,note,x,y,size,color,locked,product_id,product_snapshot,standard_cents,actual_cents,created_at,updated_at,shape,x2,y2,show_motor,fov_angle,curtain_style,curtain_length) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",(new_id,p["project_id"],p["drawing_id"],p["type_id"],number,number if p["name"]==p["number"] else p["name"],p["room"],p["quantity"],p["note"],min(1,p["x"]+dx),min(1,p["y"]+dy),p["size"],p["color"],0,p["product_id"],p["product_snapshot"],p["standard_cents"],p["actual_cents"],now(),now(),p.get("shape") or "point",min(1,p["x2"]+dx) if p.get("x2") is not None else None,min(1,p["y2"]+dy) if p.get("y2") is not None else None,p.get("show_motor",1),p.get("fov_angle",60),p.get("curtain_style","single"),p.get("curtain_length",0)))
        db.execute("UPDATE points SET layer_id=?,line_width=? WHERE id=?",(p["layer_id"],p["line_width"],new_id))
        sync_line(db,new_id)
        return point_out(one(db,"SELECT * FROM points WHERE id=?",(new_id,)))


def sync_line(db, point_id):
    p=one(db,"SELECT * FROM points WHERE id=?",(point_id,))
    if p["type_id"] in NON_PRODUCT_TYPES or not p["product_snapshot"]:
        db.execute("DELETE FROM quote_lines WHERE point_id=?",(point_id,)); return
    quote=one(db,"SELECT id FROM quotes WHERE project_id=?",(p["project_id"],))
    existing=db.execute("SELECT id FROM quote_lines WHERE point_id=?",(point_id,)).fetchone()
    if existing:
        db.execute("UPDATE quote_lines SET quantity=?,actual_cents=?,subtotal_cents=?,snapshot=?,updated_at=? WHERE point_id=?",(p["quantity"],p["actual_cents"],p["quantity"]*p["actual_cents"],p["product_snapshot"],now(),point_id))
    else:
        db.execute("INSERT INTO quote_lines VALUES(?,?,?,?,?,?,?,?)",(uid(),quote["id"],point_id,p["quantity"],p["actual_cents"],p["quantity"]*p["actual_cents"],p["product_snapshot"],now()))


@app.post("/api/points/assign-product")
def assign_product(data: AssignIn):
    with connect() as db:
        product=product_repository.get_cached(data.product_id)
        if not product or not product["enabled"]: raise HTTPException(404,"产品不存在或已停用")
        ensure_shadow_product(db,product)
        points=[one(db,"SELECT * FROM points WHERE id=?",(pid,)) for pid in data.point_ids]
        if any(p["type_id"] in NON_PRODUCT_TYPES for p in points): raise HTTPException(422,"非智能区域不能配置产品或进入报价")
        if len({p["type_id"] for p in points})!=1: raise HTTPException(422,"批量配置只支持同类型点位")
        snapshot={k:product.get(k) for k in ("id","name","brand_id","brand_name","series","model","category","device_type","system_category","standard_cents","unit","origin","params","image_path")}
        for p in points:
            db.execute("UPDATE points SET product_id=?,product_snapshot=?,standard_cents=?,actual_cents=?,updated_at=? WHERE id=?",(product["id"],json.dumps(snapshot,ensure_ascii=False),product["standard_cents"],product["standard_cents"],now(),p["id"]))
            sync_line(db,p["id"])
        return [point_out(one(db,"SELECT * FROM points WHERE id=?",(p["id"],))) for p in points]


@app.post("/api/points/{point_id}/restore-price")
def restore_price(point_id: str):
    with connect() as db:
        p=one(db,"SELECT * FROM points WHERE id=?",(point_id,))
        if p["standard_cents"] is None: raise HTTPException(422,"点位未配置产品")
        db.execute("UPDATE points SET actual_cents=?,updated_at=? WHERE id=?",(p["standard_cents"],now(),point_id)); sync_line(db,point_id)
        return point_out(one(db,"SELECT * FROM points WHERE id=?",(point_id,)))


def quote_data(db,pid):
    project=one(db,"SELECT * FROM projects WHERE id=?",(pid,)); q=one(db,"SELECT * FROM quotes WHERE project_id=?",(pid,))
    records=rows(db.execute("SELECT p.*,ql.id line_id,ql.quantity line_quantity,ql.actual_cents line_actual_cents FROM points p LEFT JOIN quote_lines ql ON ql.point_id=p.id WHERE p.project_id=? ORDER BY p.created_at",(pid,)))
    detail=[]; summary={}; missing=[]
    for p in records:
        if p["type_id"] in NON_PRODUCT_TYPES: continue
        if not p["product_snapshot"]: missing.append(p["number"]); continue
        s=json.loads(p["product_snapshot"]); qty=p["line_quantity"] or p["quantity"]; actual=p["line_actual_cents"] if p["line_actual_cents"] is not None else p["actual_cents"]
        d={"line_id":p["line_id"],"source":"point","point_id":p["id"],"number":p["number"],"room":p["room"],"quantity":qty,"note":p["note"],"actual_price":f"{actual/100:.2f}","standard_price":f'{p["standard_cents"]/100:.2f}',"subtotal":f"{qty*actual/100:.2f}","product":s}
        detail.append(d)
        key=(s["id"],actual)
        if key not in summary: summary[key]={"product":s,"quantity":0,"point_numbers":[],"actual_price":d["actual_price"],"standard_price":d["standard_price"],"subtotal_cents":0}
        summary[key]["quantity"]+=qty; summary[key]["point_numbers"].append(p["number"]); summary[key]["subtotal_cents"]+=qty*actual
    manual=rows(db.execute("SELECT * FROM manual_quote_lines WHERE quote_id=? ORDER BY created_at",(q["id"],)))
    for item in manual:
        s=json.loads(item["snapshot"])
        d={"line_id":item["id"],"source":"manual","point_id":None,"number":"清单加项","room":"不关联图纸","quantity":item["quantity"],"note":item["note"],"actual_price":f'{item["actual_cents"]/100:.2f}',"standard_price":f'{item["standard_cents"]/100:.2f}',"subtotal":f'{item["subtotal_cents"]/100:.2f}',"product":s}
        detail.append(d)
        key=(s["id"],item["actual_cents"])
        if key not in summary: summary[key]={"product":s,"quantity":0,"point_numbers":[],"actual_price":d["actual_price"],"standard_price":d["standard_price"],"subtotal_cents":0}
        summary[key]["quantity"]+=item["quantity"]; summary[key]["point_numbers"].append("清单加项"); summary[key]["subtotal_cents"]+=item["subtotal_cents"]
    summaries=[]
    for s in summary.values(): s["subtotal"]=f'{s.pop("subtotal_cents")/100:.2f}'; summaries.append(s)
    fees=rows(db.execute("SELECT * FROM extra_fees WHERE quote_id=? ORDER BY rowid",(q["id"],)))
    line_total=sum(int(Decimal(d["subtotal"])*100) for d in detail); fee_total=sum(f["amount_cents"] for f in fees)
    return {"project":project,"quote":{"id":q["id"],"note":q["note"]},"detail":detail,"summary":summaries,"missing":missing,"fees":[{**f,"amount":f'{f["amount_cents"]/100:.2f}'} for f in fees],"line_total":f"{line_total/100:.2f}","total":f'{max(0,line_total+fee_total)/100:.2f}'}


@app.get("/api/projects/{pid}/quote")
def get_quote(pid: str):
    with connect() as db: return quote_data(db,pid)


@app.patch("/api/projects/{pid}/quote")
def update_quote(pid: str,data: QuotePatch):
    with connect() as db:
        one(db,"SELECT id FROM quotes WHERE project_id=?",(pid,))
        db.execute("UPDATE quotes SET discount_cents=0,note=?,updated_at=? WHERE project_id=?",(data.note,now(),pid))
        return quote_data(db,pid)


@app.post("/api/projects/{pid}/quote-items", status_code=201)
def add_manual_quote_line(pid: str, data: ManualQuoteLineIn):
    with connect() as db:
        quote=one(db,"SELECT id FROM quotes WHERE project_id=?",(pid,))
        product=product_repository.get(db,data.product_id)
        if not product or not product["enabled"]: raise HTTPException(404,"产品不存在或已停用")
        snapshot={k:product.get(k) for k in ("id","name","brand_id","brand_name","series","model","category","device_type","system_category","standard_cents","unit","origin","params","image_path")}
        actual=money(data.actual_price) if data.actual_price not in (None,"") else product["standard_cents"]
        timestamp=now()
        db.execute("INSERT INTO manual_quote_lines VALUES(?,?,?,?,?,?,?,?,?,?,?)",(uid(),quote["id"],product["id"],data.quantity,product["standard_cents"],actual,data.quantity*actual,json.dumps(snapshot,ensure_ascii=False),data.note,timestamp,timestamp))
        db.execute("UPDATE projects SET updated_at=? WHERE id=?",(timestamp,pid))
        return quote_data(db,pid)


@app.post("/api/projects/{pid}/quote-items/batch", status_code=201)
def add_manual_quote_lines(pid: str, data: ManualQuoteBatchIn):
    product_ids=list(dict.fromkeys(data.product_ids))
    with connect() as db:
        quote=one(db,"SELECT id FROM quotes WHERE project_id=?",(pid,))
        products=[]
        for product_id in product_ids:
            product=product_repository.get(db,product_id)
            if not product or not product["enabled"]: raise HTTPException(404,"部分产品不存在或已停用")
            products.append(product)
        timestamp=now()
        for product in products:
            snapshot={k:product.get(k) for k in ("id","name","brand_id","brand_name","series","model","category","device_type","system_category","standard_cents","unit","origin","params","image_path")}
            actual=product["standard_cents"]
            db.execute("INSERT INTO manual_quote_lines VALUES(?,?,?,?,?,?,?,?,?,?,?)",(uid(),quote["id"],product["id"],data.quantity,product["standard_cents"],actual,data.quantity*actual,json.dumps(snapshot,ensure_ascii=False),data.note,timestamp,timestamp))
        db.execute("UPDATE projects SET updated_at=? WHERE id=?",(timestamp,pid))
        return quote_data(db,pid)


@app.patch("/api/quote-items/{line_id}")
def update_manual_quote_line(line_id: str, data: ManualQuoteLinePatch):
    with connect() as db:
        line=one(db,"SELECT ml.*,q.project_id FROM manual_quote_lines ml JOIN quotes q ON q.id=ml.quote_id WHERE ml.id=?",(line_id,))
        quantity=data.quantity if data.quantity is not None else line["quantity"]
        actual=money(data.actual_price) if data.actual_price is not None else line["actual_cents"]
        note=data.note if data.note is not None else line["note"]
        timestamp=now()
        db.execute("UPDATE manual_quote_lines SET quantity=?,actual_cents=?,subtotal_cents=?,note=?,updated_at=? WHERE id=?",(quantity,actual,quantity*actual,note,timestamp,line_id))
        db.execute("UPDATE projects SET updated_at=? WHERE id=?",(timestamp,line["project_id"]))
        return quote_data(db,line["project_id"])


@app.post("/api/quote-items/{line_id}/restore-price")
def restore_manual_quote_line_price(line_id: str):
    with connect() as db:
        line=one(db,"SELECT ml.*,q.project_id FROM manual_quote_lines ml JOIN quotes q ON q.id=ml.quote_id WHERE ml.id=?",(line_id,))
        db.execute("UPDATE manual_quote_lines SET actual_cents=standard_cents,subtotal_cents=quantity*standard_cents,updated_at=? WHERE id=?",(now(),line_id))
        return quote_data(db,line["project_id"])


@app.delete("/api/quote-items/{line_id}")
def delete_manual_quote_line(line_id: str):
    with connect() as db:
        line=one(db,"SELECT ml.id,q.project_id FROM manual_quote_lines ml JOIN quotes q ON q.id=ml.quote_id WHERE ml.id=?",(line_id,))
        db.execute("DELETE FROM manual_quote_lines WHERE id=?",(line_id,))
        db.execute("UPDATE projects SET updated_at=? WHERE id=?",(now(),line["project_id"]))
    return {"ok":True}


@app.post("/api/projects/{pid}/fees")
def add_fee(pid: str,data: FeeIn):
    with connect() as db:
        q=one(db,"SELECT id FROM quotes WHERE project_id=?",(pid,))
        db.execute("INSERT INTO extra_fees VALUES(?,?,?,?,?)",(uid(),q["id"],data.name,money(data.amount),data.kind))
        return quote_data(db,pid)


@app.delete("/api/fees/{fee_id}")
def delete_fee(fee_id: str):
    with connect() as db:
        one(db,"SELECT id FROM extra_fees WHERE id=?",(fee_id,)); db.execute("DELETE FROM extra_fees WHERE id=?",(fee_id,))
    return {"ok":True}


def ensure_quotation_systems(db, quote_id: str):
    timestamp = now()
    for category in (*SYSTEM_CATEGORIES, UNCATEGORIZED):
        db.execute(
            "INSERT OR IGNORE INTO quotation_systems(id,quote_id,category_id,sort_order,created_at,updated_at) VALUES(?,?,?,?,?,?)",
            (uid(), quote_id, category["id"], category["order"], timestamp, timestamp),
        )


CONSTRUCTION_FEE_ITEMS = (
    {"id": "strong_power", "name": "强电施工"},
    {"id": "weak_power", "name": "弱电施工"},
    {"id": "equipment_installation", "name": "设备安装调试"},
    {"id": "line_material", "name": "线路材料费"},
    {"id": "system_training", "name": "系统调试及培训"},
)


def quotation_snapshot(source: dict, product: dict | None = None) -> dict:
    snapshot = dict(source or {})
    if product:
        snapshot = {**snapshot, **{key: product.get(key) for key in (
            "id", "name", "model", "brand_name", "origin", "unit", "params",
            "image_path", "standard_cents", "system_category",
        ) if product.get(key) not in (None, "")}}
    return snapshot


def sync_quotation_items(db, quote: dict):
    """Copy legacy point/manual sources into editable V2 rows without overwriting edits."""
    quote_id = quote["id"]
    ensure_quotation_systems(db, quote_id)
    timestamp = now()
    expected: set[tuple[str, str]] = set()
    grouped: dict[str, dict] = {}
    point_rows = rows(db.execute(
        "SELECT ql.*,p.product_id,p.product_snapshot,p.standard_cents,p.note point_note "
        "FROM quote_lines ql JOIN points p ON p.id=ql.point_id WHERE ql.quote_id=?",
        (quote_id,),
    ))
    for row in point_rows:
        try:
            raw = json.loads(row.get("snapshot") or row.get("product_snapshot") or "{}")
        except (TypeError, json.JSONDecodeError):
            raw = {}
        product = product_repository.get(db, row.get("product_id")) if row.get("product_id") else None
        snapshot = quotation_snapshot(raw, product)
        product_id = str(snapshot.get("id") or row.get("product_id") or "")
        source_key = f"{product_id}:{int(row['actual_cents'])}"
        entry = grouped.setdefault(source_key, {"quantity": 0, "row": row, "snapshot": snapshot})
        entry["quantity"] += int(row["quantity"])
    for source_key, entry in grouped.items():
        source_type = "points"; expected.add((source_type, source_key))
        snapshot = entry["snapshot"]; row = entry["row"]
        category = canonical_system_category(snapshot.get("system_category"))
        existing = db.execute(
            "SELECT id,quantity_override,system_category FROM quotation_items WHERE quote_id=? AND source_type=? AND source_key=?",
            (quote_id, source_type, source_key),
        ).fetchone()
        if existing:
            updates = ["source_quantity=?", "updated_at=?"]
            values: list[object] = [entry["quantity"], timestamp]
            if not existing["quantity_override"]:
                updates.append("quantity=?"); values.append(entry["quantity"])
            if existing["system_category"] == "uncategorized" and category != "uncategorized":
                updates.append("system_category=?"); values.append(category)
            values.append(existing["id"])
            db.execute(f"UPDATE quotation_items SET {','.join(updates)} WHERE id=?", values)
        else:
            next_order = db.execute("SELECT COALESCE(MAX(sort_order),-1)+1 FROM quotation_items WHERE quote_id=? AND system_category=?", (quote_id, category)).fetchone()[0]
            db.execute(
                """INSERT INTO quotation_items(id,quote_id,source_type,source_key,product_id,system_category,product_name,product_code,image_path,brand,origin,quantity,source_quantity,unit_price_cents,standard_cents,unit,params,remark,sort_order,visible,created_at,updated_at)
                VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (uid(), quote_id, source_type, source_key, snapshot.get("id"), category,
                 normalize_quotation_text(snapshot.get("name") or "未命名产品"), normalize_quotation_text(snapshot.get("model")), snapshot.get("image_path"),
                 normalize_quotation_text(snapshot.get("brand_name")), normalize_quotation_text(snapshot.get("origin")), entry["quantity"], entry["quantity"],
                 int(row["actual_cents"]), int(snapshot.get("standard_cents") or row.get("standard_cents") or row["actual_cents"]),
                 normalize_quotation_text(snapshot.get("unit") or "个"), normalize_quotation_text(snapshot.get("params"), multiline=True),
                 normalize_quotation_text(row.get("point_note"), multiline=True), next_order, 1, timestamp, timestamp),
            )
    manual_rows = rows(db.execute("SELECT * FROM manual_quote_lines WHERE quote_id=?", (quote_id,)))
    for row in manual_rows:
        source_type = "manual"; source_key = row["id"]; expected.add((source_type, source_key))
        try: raw = json.loads(row.get("snapshot") or "{}")
        except (TypeError, json.JSONDecodeError): raw = {}
        product = product_repository.get(db, row.get("product_id")) if row.get("product_id") else None
        snapshot = quotation_snapshot(raw, product); category = canonical_system_category(snapshot.get("system_category"))
        existing = db.execute("SELECT id,quantity_override,system_category FROM quotation_items WHERE quote_id=? AND source_type=? AND source_key=?", (quote_id, source_type, source_key)).fetchone()
        if existing:
            updates=["source_quantity=?","updated_at=?"]; values=[row["quantity"],timestamp]
            if not existing["quantity_override"]: updates.append("quantity=?"); values.append(row["quantity"])
            if existing["system_category"] == "uncategorized" and category != "uncategorized": updates.append("system_category=?"); values.append(category)
            values.append(existing["id"]); db.execute(f"UPDATE quotation_items SET {','.join(updates)} WHERE id=?",values)
        else:
            next_order=db.execute("SELECT COALESCE(MAX(sort_order),-1)+1 FROM quotation_items WHERE quote_id=? AND system_category=?",(quote_id,category)).fetchone()[0]
            db.execute("""INSERT INTO quotation_items(id,quote_id,source_type,source_key,product_id,system_category,product_name,product_code,image_path,brand,origin,quantity,source_quantity,unit_price_cents,standard_cents,unit,params,remark,sort_order,visible,created_at,updated_at)
                VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (uid(),quote_id,source_type,source_key,snapshot.get("id") or row.get("product_id"),category,
                 normalize_quotation_text(snapshot.get("name") or "未命名产品"),normalize_quotation_text(snapshot.get("model")),snapshot.get("image_path"),
                 normalize_quotation_text(snapshot.get("brand_name")),normalize_quotation_text(snapshot.get("origin")),row["quantity"],row["quantity"],row["actual_cents"],
                 snapshot.get("standard_cents") or row["standard_cents"],normalize_quotation_text(snapshot.get("unit") or "个"),normalize_quotation_text(snapshot.get("params"),multiline=True),
                 normalize_quotation_text(row.get("note"),multiline=True),next_order,1,timestamp,timestamp))
    for item in rows(db.execute("SELECT id,source_type,source_key FROM quotation_items WHERE quote_id=? AND source_type IN ('points','manual')", (quote_id,))):
        if (item["source_type"], item["source_key"]) not in expected:
            db.execute("DELETE FROM quotation_items WHERE id=?", (item["id"],))


def quotation_v2_data(db, project_id: str) -> dict:
    project = one(db, "SELECT * FROM projects WHERE id=?", (project_id,))
    quote = one(db, "SELECT * FROM quotes WHERE project_id=?", (project_id,))
    sync_quotation_items(db, quote)
    for fee in CONSTRUCTION_FEE_ITEMS:
        db.execute("INSERT OR IGNORE INTO quotation_construction_fees(quote_id,fee_key,amount_cents,updated_at) VALUES(?,?,0,?)", (quote["id"], fee["id"], now()))
    fee_rows = {row["fee_key"]: int(row["amount_cents"] or 0) for row in rows(db.execute("SELECT fee_key,amount_cents FROM quotation_construction_fees WHERE quote_id=?", (quote["id"],)))}
    construction_fees = [{**fee, "amount_cents": fee_rows.get(fee["id"], 0)} for fee in CONSTRUCTION_FEE_ITEMS]
    system_rows = {row["category_id"]: row for row in rows(db.execute("SELECT * FROM quotation_systems WHERE quote_id=?", (quote["id"],)))}
    items = rows(db.execute("SELECT * FROM quotation_items WHERE quote_id=? ORDER BY system_category,sort_order,created_at", (quote["id"],)))
    systems=[]
    for category in (*SYSTEM_CATEGORIES, UNCATEGORIZED):
        row=system_rows[category["id"]]
        system={**category,**row,"items":[]}
        for item in (x for x in items if x["system_category"] == category["id"]):
            item["visible"] = bool(item["visible"]); item["quantity_override"] = bool(item["quantity_override"])
            item["image_url"] = f'/api/files/{item["image_path"]}' if item.get("image_path") else None
            product = product_repository.get_cached(item.get("product_id")) if item.get("product_id") else None
            brand = find_brand(db, brand_id_value=product.get("brand_id", "") if product else "", name=item.get("brand", ""))
            item["brand_id"] = brand["id"] if brand else None
            item["brand_logo_url"] = brand_out(brand)["logo_url"] if brand else None
            system["items"].append(item)
        systems.append(system)
    settings={key:quote[key] for key in ("project_name","customer_name","show_empty_categories","decimal_places","tax_rate_basis","discount_type","discount_value","template_id","company_info","quote_date","logo_path")}
    settings["show_empty_categories"] = bool(settings["show_empty_categories"])
    calc_started = time.perf_counter()
    totals=calculate_quotation(systems,settings,sum(fee["amount_cents"] for fee in construction_fees))
    record_event("quote.calculate", operation="calculate", projectId=project_id,
        itemCount=sum(len(system["items"]) for system in systems),
        durationMs=round((time.perf_counter()-calc_started)*1000, 2), success=True)
    versions=rows(db.execute("SELECT id,version,total_cents,created_by,note,created_at FROM quotation_versions WHERE quote_id=? ORDER BY version DESC",(quote["id"],)))
    return {"project":project,"quote_id":quote["id"],"version":quote["version"],"settings":settings,"systems":systems,"construction_fees":construction_fees,"totals":totals,"versions":versions,"system_categories":[*SYSTEM_CATEGORIES,UNCATEGORIZED]}


@app.get("/api/quotation-systems")
def quotation_system_categories():
    return [*SYSTEM_CATEGORIES, UNCATEGORIZED]


@app.get("/api/projects/{pid}/quotation-v2")
def get_quotation_v2(pid: str, response: Response):
    response.headers["Cache-Control"] = "no-store, max-age=0"
    with connect() as db:
        return quotation_v2_data(db,pid)


@app.patch("/api/projects/{pid}/quotation-v2")
def update_quotation_v2(pid: str, data: QuotationSettingsPatch):
    with connect() as db:
        quote=one(db,"SELECT * FROM quotes WHERE project_id=?",(pid,)); values=data.model_dump(exclude_unset=True)
        assignments=[]; args=[]
        for key,value in values.items():
            if key == "tax_rate": key="tax_rate_basis"; value=round(Decimal(value)*100)
            elif key == "discount_value":
                value = round(Decimal(value)*100) if values.get("discount_type",quote["discount_type"]) == "fixed" else round(Decimal(value)*100)
            elif key == "show_empty_categories": value=int(value)
            if key == "discount_type" and value not in ("none","percentage","fixed"): raise HTTPException(422,"折扣方式错误")
            assignments.append(f"{key}=?"); args.append(value)
        if assignments:
            assignments.append("updated_at=?"); args.extend([now(),quote["id"]]); db.execute(f"UPDATE quotes SET {','.join(assignments)} WHERE id=?",args)
        return quotation_v2_data(db,pid)


@app.post("/api/projects/{pid}/quotation-v2/logo")
async def upload_quotation_logo(pid: str,file: UploadFile=File(...)):
    content=await file.read()
    if len(content)>5*1024*1024: raise HTTPException(413,"Logo 不能超过 5 MB")
    ext=Path(file.filename or "").suffix.lower()
    if ext not in (".jpg",".jpeg",".png",".webp"): raise HTTPException(422,"Logo 仅支持 JPG、PNG、WebP")
    try:
        with Image.open(io.BytesIO(content)) as image: image.verify()
    except Exception: raise HTTPException(422,"Logo 图片无法读取")
    path=UPLOADS/f"{uid()}{ext}"; path.write_bytes(content)
    with connect() as db:
        quote=one(db,"SELECT * FROM quotes WHERE project_id=?",(pid,))
        old=UPLOADS/quote["logo_path"] if quote.get("logo_path") else None
        db.execute("UPDATE quotes SET logo_path=?,updated_at=? WHERE id=?",(path.name,now(),quote["id"]))
        result=quotation_v2_data(db,pid)
    if old and old.is_file() and old!=path:
        try: old.unlink()
        except OSError: pass
    return result


@app.patch("/api/projects/{pid}/quotation-v2/service-fees")
def apply_all_service_fees(pid: str, data: QuotationSystemPatch):
    if data.service_type not in ("none", "percentage", "fixed"):
        raise HTTPException(422, "服务费方式错误")
    with connect() as db:
        quotation_v2_data(db, pid)  # Ensure all configured categories exist.
        quote = one(db, "SELECT id FROM quotes WHERE project_id=?", (pid,))
        rate = round(Decimal(data.service_rate or "0") * 100)
        amount = money(data.service_amount or "0")
        db.execute("UPDATE quotation_systems SET service_type=?,service_rate_basis=?,service_amount_cents=?,service_name=?,updated_at=? WHERE quote_id=?",
                   (data.service_type,rate,amount,normalize_quotation_text(data.service_name),now(),quote["id"]))
        return quotation_v2_data(db,pid)


@app.patch("/api/projects/{pid}/quotation-v2/construction-fees")
def update_construction_fees(pid: str, data: ConstructionFeesPatch):
    valid_keys = {fee["id"] for fee in CONSTRUCTION_FEE_ITEMS}
    if set(data.fees) - valid_keys:
        raise HTTPException(422, "施工调试费项目无效")
    with connect() as db:
        quote = one(db, "SELECT id FROM quotes WHERE project_id=?", (pid,))
        for key, value in data.fees.items():
            amount = max(0, money(value or "0"))
            db.execute(
                "INSERT INTO quotation_construction_fees(quote_id,fee_key,amount_cents,updated_at) VALUES(?,?,?,?) "
                "ON CONFLICT(quote_id,fee_key) DO UPDATE SET amount_cents=excluded.amount_cents,updated_at=excluded.updated_at",
                (quote["id"], key, amount, now()),
            )
        return quotation_v2_data(db, pid)


@app.patch("/api/quotation-systems/{system_id}")
def update_quotation_system(system_id: str, data: QuotationSystemPatch):
    if data.service_type not in ("none","percentage","fixed"): raise HTTPException(422,"服务费方式错误")
    with connect() as db:
        row=one(db,"SELECT qs.*,q.project_id FROM quotation_systems qs JOIN quotes q ON q.id=qs.quote_id WHERE qs.id=?",(system_id,))
        rate=round(Decimal(data.service_rate or "0")*100); amount=money(data.service_amount or "0")
        db.execute("UPDATE quotation_systems SET service_type=?,service_rate_basis=?,service_amount_cents=?,service_name=?,updated_at=? WHERE id=?",
                   (data.service_type,rate,amount,normalize_quotation_text(data.service_name),now(),system_id))
        return quotation_v2_data(db,row["project_id"])


@app.patch("/api/quotation-items/{item_id}")
def update_quotation_item(item_id: str, data: QuotationItemPatch):
    with connect() as db:
        row=one(db,"SELECT qi.*,q.project_id FROM quotation_items qi JOIN quotes q ON q.id=qi.quote_id WHERE qi.id=?",(item_id,)); values=data.model_dump(exclude_unset=True)
        updates=[]; args=[]
        text_fields={"product_name","product_code","brand","origin","unit","params","remark"}
        for key,value in values.items():
            if key == "system_category": value=canonical_system_category(value)
            elif key == "unit_price": key="unit_price_cents"; value=money(value)
            elif key in text_fields: value=normalize_quotation_text(value,multiline=key in {"params","remark"})
            elif key == "visible": value=int(value)
            elif key == "quantity": updates.append("quantity_override=1")
            updates.append(f"{key}=?"); args.append(value)
        if updates:
            updates.append("updated_at=?"); args.extend([now(),item_id]); db.execute(f"UPDATE quotation_items SET {','.join(updates)} WHERE id=?",args)
        return quotation_v2_data(db,row["project_id"])


@app.post("/api/quotation-items/{item_id}/restore-source-quantity")
def restore_quotation_item_quantity(item_id: str):
    with connect() as db:
        row=one(db,"SELECT qi.*,q.project_id FROM quotation_items qi JOIN quotes q ON q.id=qi.quote_id WHERE qi.id=?",(item_id,))
        db.execute("UPDATE quotation_items SET quantity=source_quantity,quantity_override=0,updated_at=? WHERE id=?",(now(),item_id))
        return quotation_v2_data(db,row["project_id"])


@app.delete("/api/quotation-items/{item_id}")
def delete_quotation_item(item_id: str):
    with connect() as db:
        row=one(db,"SELECT qi.*,q.project_id FROM quotation_items qi JOIN quotes q ON q.id=qi.quote_id WHERE qi.id=?",(item_id,))
        if row["source_type"] == "temp": db.execute("DELETE FROM quotation_items WHERE id=?",(item_id,))
        else: db.execute("UPDATE quotation_items SET visible=0,updated_at=? WHERE id=?",(now(),item_id))
        return quotation_v2_data(db,row["project_id"])


@app.post("/api/projects/{pid}/quotation-v2/items",status_code=201)
def add_temporary_quotation_item(pid: str,data: TemporaryQuotationItemIn):
    category=canonical_system_category(data.system_category)
    with connect() as db:
        quote=one(db,"SELECT * FROM quotes WHERE project_id=?",(pid,)); product=None
        if data.product_id: product=product_repository.get(db,data.product_id)
        product_id=data.product_id; image_path=product.get("image_path") if product else None; standard=product.get("standard_cents") if product else money(data.unit_price)
        if data.save_to_library and not product:
            saved=product_repository.save({"id":uid(),"brand_name":data.brand or "未填写","name":data.product_name,"series":"","model":data.product_code,
                "category":"智能","device_type":"other","system_category":category,"standard_cents":money(data.unit_price),"cost_cents":None,"unit":data.unit,
                "origin":data.origin,"params":data.params,"description":data.remark,"image_path":None,"enabled":1,"updated_at":now()})
            ensure_shadow_product(db,saved); product_id=saved["id"]
        order=db.execute("SELECT COALESCE(MAX(sort_order),-1)+1 FROM quotation_items WHERE quote_id=? AND system_category=?",(quote["id"],category)).fetchone()[0]
        item_id=uid(); timestamp=now()
        db.execute("""INSERT INTO quotation_items(id,quote_id,source_type,source_key,product_id,system_category,product_name,product_code,image_path,brand,origin,quantity,source_quantity,quantity_override,unit_price_cents,standard_cents,unit,params,remark,sort_order,visible,created_at,updated_at)
          VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
          (item_id,quote["id"],"temp",item_id,product_id,category,normalize_quotation_text(data.product_name),normalize_quotation_text(data.product_code),image_path,
           normalize_quotation_text(data.brand),normalize_quotation_text(data.origin),data.quantity,0,1,money(data.unit_price),standard,normalize_quotation_text(data.unit),
           normalize_quotation_text(data.params,multiline=True),normalize_quotation_text(data.remark,multiline=True),order,1,timestamp,timestamp))
        return quotation_v2_data(db,pid)


@app.post("/api/quotation-items/{item_id}/replace")
def replace_quotation_item(item_id: str, data: QuotationReplaceIn):
    with connect() as db:
        row = one(db, "SELECT qi.*,q.project_id FROM quotation_items qi JOIN quotes q ON q.id=qi.quote_id WHERE qi.id=?", (item_id,))
        product = product_repository.get(db, data.product_id)
        if not product or not product.get("enabled", 1):
            raise HTTPException(404, "所选产品不存在或已停用")
        category = canonical_system_category(product.get("system_category"))
        sort_order = row["sort_order"] if row["system_category"] == category else db.execute(
            "SELECT COALESCE(MAX(sort_order),-1)+1 FROM quotation_items WHERE quote_id=? AND system_category=?",
            (row["quote_id"], category),
        ).fetchone()[0]
        db.execute(
            """UPDATE quotation_items SET product_id=?,system_category=?,product_name=?,product_code=?,image_path=?,brand=?,origin=?,
               quantity=1,source_quantity=0,quantity_override=1,unit_price_cents=?,standard_cents=?,unit=?,params=?,remark='',sort_order=?,updated_at=? WHERE id=?""",
            (product["id"], category, normalize_quotation_text(product.get("name") or "未命名产品"),
             normalize_quotation_text(product.get("model")), product.get("image_path"), normalize_quotation_text(product.get("brand_name")),
             normalize_quotation_text(product.get("origin")), int(product.get("standard_cents") or 0), int(product.get("standard_cents") or 0),
             normalize_quotation_text(product.get("unit") or "个"), normalize_quotation_text(product.get("params"), multiline=True), sort_order, now(), item_id),
        )
        return quotation_v2_data(db, row["project_id"])


@app.post("/api/quotation-items/{item_id}/image")
async def upload_quotation_item_image(item_id: str,file: UploadFile=File(...)):
    content=await file.read()
    if len(content)>5*1024*1024: raise HTTPException(413,"图片不能超过 5 MB")
    ext=Path(file.filename or "").suffix.lower()
    if ext not in (".jpg",".jpeg",".png",".webp"): raise HTTPException(422,"仅支持 JPG、PNG、WebP")
    try:
        with Image.open(io.BytesIO(content)) as image: image.verify()
    except Exception: raise HTTPException(422,"图片无法读取")
    path=UPLOADS/f"{uid()}{ext}"; path.write_bytes(content)
    with connect() as db:
        row=one(db,"SELECT qi.*,q.project_id FROM quotation_items qi JOIN quotes q ON q.id=qi.quote_id WHERE qi.id=?",(item_id,))
        db.execute("UPDATE quotation_items SET image_path=?,updated_at=? WHERE id=?",(path.name,now(),item_id))
        return quotation_v2_data(db,row["project_id"])


@app.post("/api/projects/{pid}/quotation-v2/reorder")
def reorder_quotation_items(pid: str,data: QuotationReorderIn):
    with connect() as db:
        quote=one(db,"SELECT id FROM quotes WHERE project_id=?",(pid,))
        first = db.execute("SELECT system_category FROM quotation_items WHERE id=? AND quote_id=?", (data.item_ids[0], quote["id"])).fetchone()
        if not first:
            raise HTTPException(404, "报价产品不存在")
        category = canonical_system_category(data.system_category or first["system_category"])
        if any(db.execute("SELECT 1 FROM quotation_items WHERE id=? AND quote_id=? AND system_category=?", (item_id, quote["id"], category)).fetchone() is None for item_id in data.item_ids):
            raise HTTPException(422, "只能在同一报价系统内调整产品顺序")
        timestamp = now()
        for index,item_id in enumerate(dict.fromkeys(data.item_ids)):
            db.execute("UPDATE quotation_items SET sort_order=?,updated_at=? WHERE id=? AND quote_id=? AND system_category=?",(index,timestamp,item_id,quote["id"],category))
        return quotation_v2_data(db,pid)


@app.post("/api/projects/{pid}/quotation-v2/versions",status_code=201)
def save_quotation_version(pid: str,data: QuotationVersionIn):
    with connect() as db:
        quotation=quotation_v2_data(db,pid); saved_max=db.execute("SELECT COALESCE(MAX(version),0) FROM quotation_versions WHERE quote_id=?",(quotation["quote_id"],)).fetchone()[0]; version=max(int(quotation["version"]),int(saved_max))+1
        timestamp=now(); db.execute("INSERT INTO quotation_versions VALUES(?,?,?,?,?,?,?,?)",(uid(),quotation["quote_id"],version,quotation["totals"]["grand_total_cents"],json.dumps(quotation,ensure_ascii=False),data.created_by,normalize_quotation_text(data.note,multiline=True),timestamp))
        db.execute("UPDATE quotes SET version=?,updated_at=? WHERE id=?",(version,timestamp,quotation["quote_id"]))
        return quotation_v2_data(db,pid)


@app.post("/api/projects/{pid}/quotation-v2/versions/{version_id}/restore")
def restore_quotation_version(pid: str, version_id: str):
    """Apply a saved quotation snapshot to the live quotation without changing history."""
    with connect() as db:
        quote = one(db, "SELECT * FROM quotes WHERE project_id=?", (pid,))
        saved = db.execute(
            "SELECT * FROM quotation_versions WHERE id=? AND quote_id=?",
            (version_id, quote["id"]),
        ).fetchone()
        if not saved:
            raise HTTPException(404, "报价版本不存在")
        try:
            snapshot = json.loads(saved["snapshot_json"])
            settings = snapshot["settings"]
            systems = snapshot["systems"]
            fees = snapshot["construction_fees"]
            if not isinstance(settings, dict) or not isinstance(systems, list) or not isinstance(fees, list):
                raise ValueError("snapshot shape")
        except (TypeError, ValueError, KeyError, json.JSONDecodeError):
            raise HTTPException(409, "该历史版本数据不完整，无法读取")

        timestamp = now()
        # Keep the live version counter monotonic; restoring an older snapshot must not
        # make a later save reuse an existing version number.
        saved_max = db.execute(
            "SELECT COALESCE(MAX(version),0) FROM quotation_versions WHERE quote_id=?", (quote["id"],)
        ).fetchone()[0]
        quote_updates = {key: settings[key] for key in (
            "project_name", "customer_name", "show_empty_categories", "decimal_places",
            "tax_rate_basis", "discount_type", "discount_value", "template_id",
            "company_info", "quote_date", "logo_path",
        ) if key in settings}
        quote_updates["updated_at"] = timestamp
        quote_updates["version"] = max(int(quote.get("version", 1)), int(saved_max))
        db.execute(
            f"UPDATE quotes SET {','.join(f'{key}=?' for key in quote_updates)} WHERE id=?",
            (*quote_updates.values(), quote["id"]),
        )

        # Apply per-system service settings and exact line-item snapshots.
        for system in systems:
            category_id = system.get("category_id") or system.get("id")
            if not category_id:
                continue
            db.execute(
                "UPDATE quotation_systems SET service_type=?,service_rate_basis=?,service_amount_cents=?,service_name=?,updated_at=? WHERE quote_id=? AND category_id=?",
                (system.get("service_type", "none"), int(system.get("service_rate_basis") or 0),
                 int(system.get("service_amount_cents") or 0), system.get("service_name", ""),
                 timestamp, quote["id"], category_id),
            )
        db.execute("DELETE FROM quotation_items WHERE quote_id=?", (quote["id"],))
        for system in systems:
            category = system.get("category_id") or system.get("id")
            if not category:
                continue
            for item in system.get("items", []):
                db.execute(
                    """INSERT INTO quotation_items(
                        id,quote_id,source_type,source_key,product_id,system_category,product_name,product_code,
                        image_path,brand,origin,quantity,source_quantity,quantity_override,unit_price_cents,
                        standard_cents,unit,params,remark,sort_order,visible,created_at,updated_at
                    ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    # Historical product rows become self-contained quote rows so the
                    # normal source synchronizer cannot overwrite their saved details.
                    (uid(), quote["id"], "temp", uid(),
                     item.get("product_id"), item.get("system_category") or category,
                     item.get("product_name") or "未命名产品", item.get("product_code") or "", item.get("image_path"),
                     item.get("brand") or "", item.get("origin") or "", int(item.get("quantity") or 1),
                     int(item.get("source_quantity") or 0), int(bool(item.get("quantity_override"))),
                     int(item.get("unit_price_cents") or 0), int(item.get("standard_cents") or 0),
                     item.get("unit") or "个", item.get("params") or "", item.get("remark") or "",
                     int(item.get("sort_order") or 0), int(bool(item.get("visible", True))), timestamp, timestamp),
                )
        for fee in fees:
            fee_key = fee.get("id")
            if fee_key:
                db.execute(
                    "INSERT INTO quotation_construction_fees(quote_id,fee_key,amount_cents,updated_at) VALUES(?,?,?,?) "
                    "ON CONFLICT(quote_id,fee_key) DO UPDATE SET amount_cents=excluded.amount_cents,updated_at=excluded.updated_at",
                    (quote["id"], fee_key, int(fee.get("amount_cents") or 0), timestamp),
                )
        return quotation_v2_data(db, pid)


@app.delete("/api/projects/{pid}/quotation-v2/versions/{version_id}")
def delete_quotation_version(pid: str, version_id: str):
    with connect() as db:
        quote = one(db, "SELECT id FROM quotes WHERE project_id=?", (pid,))
        result = db.execute("DELETE FROM quotation_versions WHERE id=? AND quote_id=?", (version_id, quote["id"]))
        if result.rowcount == 0:
            raise HTTPException(404, "报价版本不存在")
        return quotation_v2_data(db, pid)


PROJECT_STATUS_DEFINITIONS = [
    ("resolved", "🟢", "已解决"),
    ("developing", "🟡", "开发中"),
    ("optimizing", "🔵", "待优化"),
    ("pending", "⚪", "待开发"),
    ("problem", "🔴", "有问题"),
]


def project_status_value(value: str):
    for key, emoji, label in PROJECT_STATUS_DEFINITIONS:
        if label in value:
            return {"key": key, "emoji": emoji, "label": label, "display": f"{emoji} {label}"}
    return {"key": "unknown", "emoji": "", "label": value.strip() or "未标记", "display": value.strip() or "未标记"}


def markdown_h2_sections(text: str):
    matches = list(re.finditer(r"^##\s+(.+?)\s*$", text, re.MULTILINE))
    return [
        (match.group(1).strip(), text[match.end():matches[index + 1].start()] if index + 1 < len(matches) else text[match.end():])
        for index, match in enumerate(matches)
    ]


def markdown_bullets(text: str, limit: int = 8):
    items = []
    for line in text.splitlines():
        value = re.sub(r"^[-*]\s+", "", line.strip())
        if value != line.strip() and value and not value.startswith("状态："):
            items.append(value.replace("`", ""))
        if len(items) >= limit:
            break
    return items


def parse_requirements_document(text: str):
    overview = next((body for title, body in markdown_h2_sections(text) if title == "总览"), "")
    details = {title.split()[0]: body for title, body in markdown_h2_sections(text) if re.match(r"^RQ-\d", title)}
    result = []
    for line in overview.splitlines():
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if len(cells) < 4 or not re.fullmatch(r"RQ-\d+(?:[A-Z])?", cells[0]):
            continue
        status = project_status_value(cells[2])
        result.append({
            "id": cells[0], "title": cells[1], "status": status,
            "next_action": cells[3].replace("`", ""),
            "details": markdown_bullets(details.get(cells[0], ""), 4),
        })
    return result


def parse_issues_document(text: str):
    result = []
    for heading, body in markdown_h2_sections(text):
        match = re.match(r"^(ISSUE-\d+)\s+(.+)$", heading)
        if not match:
            continue
        fields = {}
        for line in body.splitlines():
            field = re.match(r"^-\s+([^：]+)：(.*)$", line.strip())
            if field:
                fields[field.group(1).strip()] = field.group(2).strip().replace("`", "")
        solution = fields.get("解决方案", "")
        if not solution:
            solution = " ".join(value for value in (fields.get("已完成方案", ""), fields.get("待完成方案", "")) if value)
        result.append({
            "id": match.group(1), "title": match.group(2),
            "requirement_id": fields.get("对应需求", ""),
            "status": project_status_value(fields.get("状态", "")),
            "problem": fields.get("问题", ""), "cause": fields.get("原因", ""),
            "solution": solution, "test_result": fields.get("测试结果", ""),
            "new_issue": fields.get("新问题", ""), "last_updated": fields.get("最后更新", ""),
        })
    return result


def parse_changelog_document(text: str):
    entries = []
    for heading, body in markdown_h2_sections(text):
        match = re.match(r"^(\d{4}-\d{2}-\d{2})｜(.+)$", heading)
        if not match:
            continue
        subsections = {}
        headings = list(re.finditer(r"^###\s+(.+?)\s*$", body, re.MULTILINE))
        for index, section in enumerate(headings):
            section_body = body[section.end():headings[index + 1].start()] if index + 1 < len(headings) else body[section.end():]
            subsections[section.group(1).strip()] = markdown_bullets(section_body)
        preferred = next((subsections[name] for name in ("功能变化", "原因与修复", "整理内容", "状态变化") if subsections.get(name)), [])
        if not preferred:
            preferred = markdown_bullets(body)
        entries.append({"date": match.group(1), "title": match.group(2), "items": preferred[:8]})
    return entries


@app.get("/api/project-status")
def project_status(response: Response):
    docs = {
        "requirements": ROOT / "docs" / "REQUIREMENTS.md",
        "issues": ROOT / "docs" / "ISSUES.md",
        "changelog": ROOT / "docs" / "CHANGELOG.md",
    }
    missing = [path.name for path in docs.values() if not path.is_file()]
    if missing:
        raise HTTPException(503, f"项目状态文档缺失：{', '.join(missing)}")
    requirements = parse_requirements_document(docs["requirements"].read_text(encoding="utf-8"))
    issues = parse_issues_document(docs["issues"].read_text(encoding="utf-8"))
    changelog = parse_changelog_document(docs["changelog"].read_text(encoding="utf-8"))
    status_definitions = [{"key": key, "emoji": emoji, "label": label} for key, emoji, label in PROJECT_STATUS_DEFINITIONS]
    counts = {}
    for name, records in (("requirements", requirements), ("issues", issues), ("total", requirements + issues)):
        counts[name] = {definition["key"]: sum(item["status"]["key"] == definition["key"] for item in records) for definition in status_definitions}
    response.headers["Cache-Control"] = "no-store"
    return {
        "source": {name: f"docs/{path.name}" for name, path in docs.items()},
        "status_definitions": status_definitions, "counts": counts,
        "requirements": requirements, "issues": issues, "changelog": changelog,
    }


@app.get("/api/feedback")
def list_feedback():
    with connect() as db:
        return rows(db.execute("SELECT * FROM feedback ORDER BY created_at DESC"))


@app.post("/api/feedback", status_code=201)
def create_feedback(data: FeedbackIn):
    item={"id":uid(),"title":data.title.strip(),"description":data.description.strip(),"module":data.module.strip(),"screenshot_path":None,"created_at":now(),"status":"待处理"}
    with connect() as db:
        db.execute("INSERT INTO feedback VALUES(:id,:title,:description,:module,:screenshot_path,:created_at,:status)",item)
    return item


@app.post("/api/feedback/{feedback_id}/screenshot")
async def upload_feedback_screenshot(feedback_id: str, file: UploadFile = File(...)):
    ext=Path(file.filename or "").suffix.lower()
    if ext not in (".jpg",".jpeg",".png",".webp"): raise HTTPException(422,"截图仅支持 JPG、PNG、WebP")
    content=await file.read()
    if len(content)>8*1024*1024: raise HTTPException(413,"截图不能超过 8 MB")
    try:
        with Image.open(io.BytesIO(content)) as image:
            image.verify()
    except Exception:
        raise HTTPException(422,"截图无法读取")
    filename=f"{uid()}{ext}"; (UPLOADS/filename).write_bytes(content)
    with connect() as db:
        one(db,"SELECT id FROM feedback WHERE id=?",(feedback_id,))
        db.execute("UPDATE feedback SET screenshot_path=? WHERE id=?",(filename,feedback_id))
    return {"screenshot_path":filename}


def font(size):
    if FONT_PATH is None:
        raise HTTPException(500, "未找到中文字体。请安装微软雅黑，或设置 POINT_QUOTE_FONT_PATH。")
    return ImageFont.truetype(str(FONT_PATH),size)


def safe_name(name):
    return re.sub(r"[^\w\u4e00-\u9fff-]+","_",name)[:50]


def draw_device_symbol(draw, type_id, center, size, color, *, width=3, shadow=True):
    """White, borderless device badge and compact glyph matching the approved canvas icons."""
    x, y = center
    r = size / 2
    w = max(1, round(width))
    square = type_id in {"switch", "wall_ap"}
    box = (x-r, y-r, x+r, y+r)
    shadow_box = (x-r, y-r+max(1, w//2), x+r, y+r+max(1, w//2))
    if square:
        if shadow: draw.rounded_rectangle(shadow_box, radius=max(2, round(r*.48)), fill="#e7ebf1")
        draw.rounded_rectangle(box, radius=max(2, round(r*.48)), fill="white")
    else:
        if shadow: draw.ellipse(shadow_box, fill="#e7ebf1")
        draw.ellipse(box, fill="white")
    def pt(u, v): return (x+(u-24)*r/24, y+(v-24)*r/24)
    def line(*coords): draw.line([pt(*pair) for pair in coords], fill=color, width=w, joint="curve")
    def ellipse(cx, cy, rx, ry=None, fill=None):
        ry = rx if ry is None else ry
        px, py = pt(cx, cy)
        draw.ellipse((px-rx*r/24, py-ry*r/24, px+rx*r/24, py+ry*r/24), outline=None if fill else color, fill=fill, width=w)
    def rect(l,t,rr,bb,radius=0):
        x1,y1=pt(l,t); x2,y2=pt(rr,bb)
        draw.rounded_rectangle((x1,y1,x2,y2),radius=max(1,round(radius*r/24)),outline=color,width=w)
    def arc(l,t,rr,bb,start,end):
        x1,y1=pt(l,t); x2,y2=pt(rr,bb)
        draw.arc((x1,y1,x2,y2),start,end,fill=color,width=w)
    if type_id == "switch":
        rect(7,10,41,38,5); line((24,11),(24,37))
        for yy in (17,31):
            line((13,yy),(18,yy)); line((30,yy),(35,yy))
    elif type_id == "curtain":
        line((5,9),(43,9))
        line((8,14),(22,14),(22,41),(6,41),(8,33),(9,23),(8,14))
        line((40,14),(26,14),(26,41),(42,41),(40,33),(39,23),(40,14))
        line((14,15),(16,24),(14,40)); line((34,15),(32,24),(34,40))
    elif type_id == "sensor":
        arc(12,12,36,36,180,360); arc(18,18,30,30,180,360)
        ellipse(24,28,4); line((18,39),(30,39))
    elif type_id == "thermostat":
        line((20,9),(20,31)); arc(16,28,32,44,0,180); line((28,31),(28,9))
        arc(20,5,28,13,180,360); line((24,16),(24,35)); ellipse(24,36,3,fill=color)
        line((33,14),(40,14)); line((33,21),(38,21))
    elif type_id == "light":
        ellipse(24,21,9)
        for coords in [((19,31),(29,31)),((20,36),(28,36)),((22,40),(26,40)),((24,3),(24,8)),((8,21),(13,21)),((35,21),(40,21)),((12,9),(16,13)),((36,9),(32,13))]: line(*coords)
    elif type_id == "lock":
        rect(11,7,35,41); line((17,19),(28,19)); line((17,24),(25,24)); ellipse(30,31,2)
        line((35,24),(42,24),(42,38),(35,38))
    elif type_id == "camera":
        line((6,16),(35,16),(43,24),(35,32),(6,32),(6,16))
        line((9,21),(26,21)); ellipse(34,24,4)
        line((13,32),(13,38),(32,38)); line((23,38),(23,43))
    elif type_id in {"ceiling_ap", "wall_ap", "network"}:
        rise = 0 if type_id == "wall_ap" else 4 if type_id == "network" else 5
        arc(8,5+rise,40,37+rise,205,335)
        arc(14,12+rise,34,32+rise,205,335)
        arc(20,20+rise,28,28+rise,205,335)
        ellipse(24,37+rise//2,2,fill=color)
    elif type_id == "water":
        line((24,5),(13,23),(12,30),(15,37),(20,41),(28,41),(34,37),(36,30),(35,23),(24,5))
        line((19,31),(21,35),(26,36))
    elif type_id == "av":
        line((9,13),(21,13),(32,6),(32,42),(21,35),(9,35),(9,13))
        arc(33,17,43,31,290,70); arc(33,12,47,36,290,70)
    elif type_id == "system":
        rect(13,13,35,35,3); rect(19,19,29,29)
        for v in (18,30):
            line((v,5),(v,13)); line((v,35),(v,43)); line((5,v),(13,v)); line((35,v),(43,v))
    elif type_id == "non_smart_area":
        line((7,8),(41,8),(41,40),(7,40),(7,8))
        line((13,32),(31,14)); line((21,36),(37,20))
    else:
        ellipse(24,24,7)
        for start,end in [((24,5),(24,14)),((24,34),(24,43)),((5,24),(14,24)),((34,24),(43,24)),((11,11),(17,17)),((31,31),(37,37)),((37,11),(31,17)),((17,31),(11,37))]: line(start,end)


def draw_export(db,pid,did,show_name=True,show_number=True,layer_id=None,layer_ids=None):
    project=one(db,"SELECT * FROM projects WHERE id=?",(pid,)); drawing=one(db,"SELECT * FROM drawings WHERE id=? AND project_id=?",(did,pid))
    if layer_ids:
        layer_ids=[item for item in layer_ids.split(",") if item]
        if any(not db.execute("SELECT 1 FROM layers WHERE id=? AND floor_id=?",(item,drawing["floor_id"])).fetchone() for item in layer_ids): raise HTTPException(422,"图层不属于当前楼层")
    elif layer_id:
        one(db,"SELECT id FROM layers WHERE id=? AND floor_id=?",(layer_id,drawing["floor_id"]))
    original=Image.open(UPLOADS/drawing["background_path"]).convert("RGB")
    ow,oh=original.size
    fixed=drawing.get("canvas_mode")=="fixed"
    design_plot=(690,211,1170,658) if fixed else (690,55,1170,970)
    scale=1.0 if fixed else max(1.0,ow/design_plot[2],oh/design_plot[3])
    W,H=round(1920*scale),round(1080*scale); k=W/1920
    canvas=Image.new("RGB",(W,H),"white"); draw=ImageDraw.Draw(canvas)
    plot_x,plot_y,plot_w,plot_h=[round(v*k) for v in design_plot]
    if fixed:
        cw,ch=drawing.get("canvas_width") or 1600,drawing.get("canvas_height") or 900
        canvas_scale=min(plot_w/cw,plot_h/ch)
        frame_x,frame_y=drawing.get("image_x",0),drawing.get("image_y",0)
        frame_w,frame_h=drawing.get("image_width",cw),drawing.get("image_height",ch)
        image_scale=canvas_scale*frame_w/ow
        pw,ph=round(frame_w*canvas_scale),round(frame_h*canvas_scale)
        left=plot_x+round(frame_x*canvas_scale); top=plot_y+round(frame_y*canvas_scale)
        image=original.resize((max(1,pw),max(1,ph)),Image.Resampling.LANCZOS)
        angle=float(drawing.get("image_rotation") or 0)
        if angle: image=image.rotate(-angle,expand=True,resample=Image.Resampling.BICUBIC)
        canvas.paste(image,(left+(pw-image.width)//2,top+(ph-image.height)//2))
        def point_xy(x,y):
            px=x*frame_w; py=y*frame_h; dx=px-frame_w/2; dy=py-frame_h/2; rad=math.radians(angle)
            return (plot_x+round((frame_x+frame_w/2+dx*math.cos(rad)-dy*math.sin(rad))*canvas_scale),plot_y+round((frame_y+frame_h/2+dx*math.sin(rad)+dy*math.cos(rad))*canvas_scale))
    else:
        image_scale=min(plot_w/ow,plot_h/oh); pw,ph=round(ow*image_scale),round(oh*image_scale)
        left=plot_x+(plot_w-pw)//2; top=plot_y+(plot_h-ph)//2
        canvas.paste(original.resize((pw,ph),Image.Resampling.LANCZOS),(left,top))
        def point_xy(x,y): return left+round(x*pw),top+round(y*ph)
    title_font=font(max(28,round(48*k))); sub_font=font(max(16,round(23*k))); label_font=font(max(14,round(18*k)))
    title=drawing.get("name") or Path(drawing["filename"]).stem
    draw.text((75*k,60*k),title,font=title_font,fill="#626674")
    draw.text((78*k,130*k),project["name"],font=sub_font,fill="#8b8f99")
    if project.get("address"): draw.text((78*k,165*k),project["address"][:32],font=label_font,fill="#9a9da6")
    if layer_ids is not None:
        if layer_ids:
            marks=",".join("?" for _ in layer_ids)
            points=rows(db.execute(f"SELECT * FROM points WHERE drawing_id=? AND layer_id IN ({marks}) ORDER BY created_at",(did,*layer_ids)))
        else: points=[]
    elif layer_id:
        points=rows(db.execute("SELECT * FROM points WHERE drawing_id=? AND layer_id=? ORDER BY created_at",(did,layer_id)))
    else:
        points=rows(db.execute("SELECT p.* FROM points p JOIN layers l ON l.id=p.layer_id WHERE p.drawing_id=? AND l.visible=1 ORDER BY p.created_at",(did,)))
    # Area rectangles remain individually editable in storage, while exports draw their
    # geometric union once so overlapping fills and borders do not darken or double up.
    area_groups={}
    for point in points:
        if point.get("shape")!="area": continue
        area_groups.setdefault((point["type_id"],point["color"]),[]).append(point)
    for (area_type,area_color),items in area_groups.items():
        mask=Image.new("L",canvas.size,0); md=ImageDraw.Draw(mask)
        for point in items:
            ax,ay=point_xy(point["x"],point["y"]); bx,by=point_xy(point.get("x2") if point.get("x2") is not None else point["x"],point.get("y2") if point.get("y2") is not None else point["y"])
            md.rectangle((min(ax,bx),min(ay,by),max(ax,bx),max(ay,by)),fill=255)
        fill=Image.new("RGBA",canvas.size,area_color+("24" if area_type=="non_smart_area" else "35")); canvas.paste(fill,(0,0),mask)
        if area_type=="non_smart_area":
            hatch=Image.new("RGBA",canvas.size,(0,0,0,0)); hd=ImageDraw.Draw(hatch); step=max(16,round(18*k))
            for offset in range(-canvas.height,canvas.width+canvas.height,step): hd.line((offset,0,offset-canvas.height,canvas.height),fill=area_color+"35",width=max(1,round(k)))
            canvas.paste(hatch,(0,0),ImageChops.multiply(mask,hatch.getchannel("A")))
        edge=ImageChops.subtract(mask,mask.filter(ImageFilter.MinFilter(3)))
        outline=Image.new("RGBA",canvas.size,area_color+"ff"); canvas.paste(outline,(0,0),edge)
    draw=ImageDraw.Draw(canvas)
    used=set()
    for p in points:
        x,y=point_xy(p["x"],p["y"]); radius=max(round(12*k),round(p["size"]*image_scale/2)); color=p["color"]
        shape=p.get("shape") or "point"; x2,y2=point_xy(p.get("x2") if p.get("x2") is not None else p["x"],p.get("y2") if p.get("y2") is not None else p["y"])
        strength=int(project.get("shadow_strength") or 0)
        if strength and shape in {"point","sector"}:
            spread=max(3,round(5*k)); side=2*(radius+spread)
            shadow=Image.new("RGBA",(side,side),(0,0,0,0)); shadow_draw=ImageDraw.Draw(shadow)
            shadow_draw.ellipse((spread,spread+round(2*k),side-spread,side-spread+round(2*k)),fill=(23,35,61,min(110,round(strength*.8))))
            shadow=shadow.filter(ImageFilter.GaussianBlur(max(1,round(3*k))))
            canvas.paste(shadow,(round(x-side/2),round(y-side/2)),shadow); draw=ImageDraw.Draw(canvas)
        if shape=="sector":
            dx,dy=x2-x,y2-y; distance=max(1,(dx*dx+dy*dy)**.5); direction=math.degrees(math.atan2(dy,dx)); half=float(p.get("fov_angle") or 60)/2
            overlay=Image.new("RGBA",canvas.size,(0,0,0,0)); od=ImageDraw.Draw(overlay); sector_box=(x-distance,y-distance,x+distance,y+distance)
            od.pieslice(sector_box,start=direction-half,end=direction+half,fill=color+"2e")
            for boundary in (direction-half,direction+half):
                radians=math.radians(boundary); od.line((x,y,x+math.cos(radians)*distance,y+math.sin(radians)*distance),fill=color+"bb",width=max(2,round(2*k)))
            canvas.paste(overlay,(0,0),overlay); draw=ImageDraw.Draw(canvas)
            draw_device_symbol(draw,"camera",(x,y),radius*2,color,width=max(2,round(2*k)),shadow=bool(strength))
        elif shape=="line":
            line_width=max(round(4*k),round(float(p.get("line_width") or 3)*k))
            dx,dy=x2-x,y2-y; length=max(1,(dx*dx+dy*dy)**.5); ux,uy=dx/length,dy/length; px,py=-uy,ux
            arrow=max(round(14*k),line_width*2); half=arrow*.42
            style=p.get("curtain_style") or "single"; split=style in {"double","double_layer_double"}
            mid_x,mid_y=(x+x2)/2,(y+y2)/2; gap=min(round(8*k),length*.08)
            left_tip=(mid_x-ux*gap,mid_y-uy*gap); right_tip=(mid_x+ux*gap,mid_y+uy*gap)
            offset=max(round(20*k),line_width*3); amplitude=max(2,round(8*k))
            def arrow_line(start,end):
                sx,sy=start; ex,ey=end; vx,vy=ex-sx,ey-sy; seg_len=max(1,math.hypot(vx,vy)); avx,avy=vx/seg_len,vy/seg_len
                draw.line((sx,sy,ex,ey),fill=color,width=line_width)
                draw.polygon(((ex,ey),(ex-avx*arrow-avy*half,ey-avy*arrow+avx*half),(ex-avx*arrow+avy*half,ey-avy*arrow-avx*half)),fill=color)
            def wave_line(start,end):
                sx,sy=start; ex,ey=end; vx,vy=ex-sx,ey-sy; seg_len=max(1,math.hypot(vx,vy)); cycles=max(1,round(seg_len/max(1,40*k))); samples=max(24,round(seg_len/max(1,3*k))); wave=[]
                for step in range(samples+1):
                    t=step/samples; displacement=amplitude*math.sin(math.pi*t)*math.sin(2*math.pi*cycles*t)
                    wave.append((sx+vx*t+px*displacement,sy+vy*t+py*displacement))
                draw.line(wave,fill=color,width=line_width,joint="curve")
                arrow_line(wave[-2],wave[-1])
            upper_start=(x+px*offset,y+py*offset); upper_end=(x2+px*offset,y2+py*offset)
            lower_start=(x-px*offset,y-py*offset); lower_end=(x2-px*offset,y2-py*offset)
            upper_left_tip=(left_tip[0]+px*offset,left_tip[1]+py*offset); upper_right_tip=(right_tip[0]+px*offset,right_tip[1]+py*offset)
            lower_left_tip=(left_tip[0]-px*offset,left_tip[1]-py*offset); lower_right_tip=(right_tip[0]-px*offset,right_tip[1]-py*offset)
            has_motor=bool(p.get("show_motor",1)); motor_scale=max(.6,float(p.get("size") or 20)/20); motor_w=max(round(28*k*motor_scale),line_width*2); motor_h=max(round(50*k*motor_scale),line_width*4)
            def track_start(normal_offset=0):
                depth=0
                if has_motor and abs(normal_offset)<motor_h/2:
                    depth=(motor_w/2)*math.sqrt(1-(normal_offset/(motor_h/2))**2)
                return (x+ux*depth+px*normal_offset,y+uy*depth+py*normal_offset)
            upper_start=track_start(offset); lower_start=track_start(-offset)
            if style=="single": arrow_line(track_start(),(x2,y2))
            elif style=="double": arrow_line(track_start(),left_tip); arrow_line((x2,y2),right_tip)
            elif style=="double_layer_single": wave_line(upper_start,upper_end); arrow_line(lower_start,lower_end)
            elif style=="double_layer_double":
                wave_line(upper_start,upper_left_tip); wave_line(upper_end,upper_right_tip)
                arrow_line(lower_start,lower_left_tip); arrow_line(lower_end,lower_right_tip)
            else: arrow_line(track_start(),(x2,y2))
            if has_motor:
                side=math.ceil(math.hypot(motor_w,motor_h))+4; motor_layer=Image.new("RGBA",(side,side),(0,0,0,0)); motor_draw=ImageDraw.Draw(motor_layer)
                motor_draw.rounded_rectangle((side/2-motor_w/2,side/2-motor_h/2,side/2+motor_w/2,side/2+motor_h/2),radius=max(1,round(2*k)),fill="white",outline=color,width=max(1,round(2*k)))
                motor_layer=motor_layer.rotate(-math.degrees(math.atan2(dy,dx)),resample=Image.Resampling.BICUBIC)
                canvas.paste(motor_layer,(round(x-side/2),round(y-side/2)),motor_layer); draw=ImageDraw.Draw(canvas)
        elif shape=="area":
            pass
        else:
            draw_device_symbol(draw,p["type_id"],(x,y),radius*2,color,width=max(2,round(2*k)),shadow=bool(strength))
        label=("非智能区域"+(f" · {p['number']}" if show_number else "")) if p["type_id"]=="non_smart_area" else " · ".join(v for v in (p["number"] if show_number else "",p["name"] if show_name and p["name"]!=p["number"] else "") if v)
        if label:
            bbox=draw.textbbox((0,0),label,font=label_font); tw=bbox[2]-bbox[0]; tx=min(max(left,x+radius+7*k),left+pw-tw-8*k); ty=min(max(top,y-radius-5*k),top+ph-(bbox[3]-bbox[1])-5*k)
            draw.rounded_rectangle((tx-3*k,ty-2*k,tx+tw+3*k,ty+bbox[3]-bbox[1]+3*k),radius=max(3,round(3*k)),fill="white")
            draw.text((tx,ty-bbox[1]),label,font=label_font,fill="#1f2937")
        used.add(p["type_id"])
    legend_types=[]
    for item in TYPES:
        if item[0] not in used: continue
        if item[0]=="other":
            names=sorted({p["name"] if p["name"] and p["name"]!=p["number"] else item[1] for p in points if p["type_id"]=="other"})
            legend_types.extend(("other::"+name,name,item[2],item[3]) for name in names)
        else: legend_types.append(item)
    legend_positions=json.loads(drawing.get("legend_items") or "{}")
    for index,t in enumerate(legend_types):
        item_position=legend_positions.get(t[0]) or [float(drawing.get("legend_x") or .025),float(drawing.get("legend_y") or .025)+index*52/max(1,ph)]
        symbol_x,symbol_y=point_xy(float(item_position[0]),float(item_position[1]))
        symbol_x+=round(22*k); symbol_y+=round(16*k)
        draw_device_symbol(draw,"other" if t[0].startswith("other::") else t[0],(symbol_x,symbol_y),round(28*k),t[3],width=max(2,round(2*k)))
        draw.text((symbol_x+round(28*k),symbol_y-round(10*k)),t[1],font=label_font,fill="#3f4756")
    stamp=f"导出时间：{now()[:19].replace('T',' ')}"; draw.text((78*k,H-64*k),stamp,font=label_font,fill="#8c9099")
    return canvas,project


@app.get("/api/projects/{pid}/drawings/{did}/export.{format}")
def export_drawing(pid: str,did: str,format: str,show_name: bool=True,show_number: bool=True,layer_id: Optional[str]=None,layer_ids: Optional[str]=None):
    if format not in ("png","pdf"): raise HTTPException(422,"格式错误")
    with connect() as db: image,project=draw_export(db,pid,did,show_name,show_number,layer_id,layer_ids)
    path=EXPORTS/f'{safe_name(project["name"])}_点位图_{datetime.now():%Y%m%d_%H%M%S}.{format}'
    if format=="png": image.save(path)
    else: image.save(path,"PDF",resolution=200)
    return FileResponse(path,filename=path.name,headers={"Cache-Control":"no-store, max-age=0"})


@app.get("/api/projects/{pid}/quote/export.xlsx")
def export_quote_excel(pid: str):
    with connect() as db: data=quote_data(db,pid)
    wb=Workbook(); ws=wb.active; ws.title="产品清单"
    ws.append([data["project"]["name"]+" 报价清单"]); ws.merge_cells("A1:J1")
    ws.append(["客户",data["project"]["customer"],"地址",data["project"]["address"]]); ws.merge_cells("D2:J2"); ws.append([])
    headers=["产品名称","图片","产品编号","产品参数","品牌","产地","数量","单价","单位","合计"]
    ws.append(headers)
    for i,d in enumerate(data["detail"],1):
        s=d["product"]; row=ws.max_row+1
        ws.append([s["name"],"",s.get("model") or "",s.get("params") or "",s["brand_name"],s.get("origin") or "",d["quantity"],float(d["actual_price"]),s["unit"],float(d["subtotal"])])
        ws.row_dimensions[row].height=48
        img_path=product_image_file(s["image_path"])
        if img_path:
            im=Image.open(img_path).convert("RGBA"); im.thumbnail((55,55)); white=Image.new("RGBA",im.size,"white"); white.alpha_composite(im); b=io.BytesIO(); white.convert("RGB").save(b,"PNG"); b.seek(0); xi=XImage(b); xi.anchor=f"B{row}"; ws.add_image(xi)
    ws.append([]); ws.append(["合计",None,None,None,None,None,None,None,None,float(data["line_total"])]); ws.merge_cells(start_row=ws.max_row,start_column=1,end_row=ws.max_row,end_column=9)
    for f in data["fees"]: ws.append([f"{f['kind']} · {f['name']}",None,None,None,None,None,None,None,None,float(f["amount"])]); ws.merge_cells(start_row=ws.max_row,start_column=1,end_row=ws.max_row,end_column=9)
    ws.append(["总计",None,None,None,None,None,None,None,None,float(data["total"])]); ws.merge_cells(start_row=ws.max_row,start_column=1,end_row=ws.max_row,end_column=9)
    groups=quote_groups(data)
    summary=wb.create_sheet("项目汇总"); summary.append(["项目类别","项目合计"])
    for group in groups: summary.append([group["label"],float(group["subtotal"])])
    if data["fees"]: summary.append(["服务及其他费用",sum(float(f["amount"]) for f in data["fees"])])
    summary.append(["合计",float(data["total"])])
    for sheet,header_row in ((ws,4),(summary,1)):
        for c in sheet[header_row]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="000000"); c.alignment=Alignment(horizontal="center",vertical="center")
        sheet.freeze_panes=f"A{header_row+1}"; sheet.sheet_view.showGridLines=False
        for row in sheet.iter_rows():
            for cell in row: cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    ws["A1"].font=Font(bold=True,color="FFFFFF",size=18); ws["A1"].fill=PatternFill("solid",fgColor="000000"); ws["A1"].alignment=Alignment(horizontal="center",vertical="center"); ws.row_dimensions[1].height=44
    widths=[20,14,16,42,18,14,10,16,10,18]
    for i,width in enumerate(widths,1): ws.column_dimensions[chr(64+i)].width=width
    summary.column_dimensions["A"].width=28; summary.column_dimensions["B"].width=20
    for row in ws.iter_rows(min_row=5):
        for col in (8,10): row[col-1].number_format='¥#,##0.00'
    for row in summary.iter_rows(min_row=2): row[1].number_format='¥#,##0.00'
    path=EXPORTS/f'{safe_name(data["project"]["name"])}_产品清单_{datetime.now():%Y%m%d_%H%M%S}.xlsx'; wb.save(path)
    return FileResponse(path,filename=path.name,headers={"Cache-Control":"no-store, max-age=0"})


QUOTE_LABELS={"智能":"智能系统","影音":"影音系统","灯光":"灯光系统","网络":"网络系统"}


def quote_groups(data):
    grouped={}
    for item in data["summary"]:
        category=canonical_category(item["product"].get("category") or "智能")
        if category not in grouped: grouped[category]={"category":category,"label":QUOTE_LABELS.get(category,category),"items":[],"subtotal_cents":0}
        grouped[category]["items"].append(item); grouped[category]["subtotal_cents"]+=round(Decimal(item["subtotal"])*100)
    result=[]
    for category in PRODUCT_CATEGORIES:
        if category in grouped:
            group=grouped.pop(category); group["subtotal"]=f'{group.pop("subtotal_cents")/100:.2f}'; result.append(group)
    for group in grouped.values(): group["subtotal"]=f'{group.pop("subtotal_cents")/100:.2f}'; result.append(group)
    return result


@app.get("/api/projects/{pid}/quote/export.pptx")
def export_quote_pptx_file(pid: str):
    with connect() as db:
        data = quote_data(db, pid)
    path = EXPORTS / f'{safe_name(data["project"]["name"])}_可编辑报价_{datetime.now():%Y%m%d_%H%M%S}.pptx'
    export_quote_pptx(data, path, product_image_file)
    return FileResponse(path, filename=path.name, media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation", headers={"Cache-Control": "no-store, max-age=0"})


def quote_page(data,title,items,page_label,show_totals=False,show_grand=False,category_total="0.00"):
    S=2; W,H=1920,1080; im=Image.new("RGB",(W*S,H*S),"white"); d=ImageDraw.Draw(im)
    def box(coords,fill=None,outline="#000000",width=2): d.rectangle(tuple(round(v*S) for v in coords),fill=fill,outline=outline,width=max(1,round(width*S)))
    def txt(x,y,value,size=22,color="#111111",anchor="mm"):
        d.text((x*S,y*S),str(value),font=font(size*S),fill=color,anchor=anchor)
    def wrapped(value, max_chars, max_lines=3):
        value=str(value or "-").replace("\r\n","\n").replace("\r","\n")
        lines=[]
        for paragraph in value.split("\n"):
            lines.extend(textwrap.wrap(paragraph, width=max_chars, break_long_words=True, break_on_hyphens=False) or [""])
        if len(lines)>max_lines:
            lines=lines[:max_lines]
            lines[-1]=(lines[-1][:-1] if lines[-1] else "")+"…"
        return "\n".join(lines)
    def wrapped_txt(x,y,value,size,max_chars,max_lines=3):
        d.multiline_text((x*S,y*S),wrapped(value,max_chars,max_lines),font=font(size*S),fill="#111111",anchor="mm",align="center",spacing=2*S)
    box((132,115,1788,220),fill="#000000"); txt(960,168,title,34,"#ffffff")
    txt(134,92,data["project"]["name"],18,"#60646d",anchor="lm"); txt(1786,92,page_label,16,"#777b84",anchor="rm")
    x=[132,380,500,690,1040,1190,1290,1400,1540,1630,1788]
    headers=["产品名称","图片","产品编号","产品参数","品牌","产地","数量","单价","单位","合计"]
    top=285; header_h=44; row_h=82
    for i,h in enumerate(headers): box((x[i],top,x[i+1],top+header_h),fill="#000000"); txt((x[i]+x[i+1])/2,top+header_h/2,h,17,"#ffffff")
    for row_index,item in enumerate(items):
        y=top+header_h+row_index*row_h; s=item["product"]
        for i in range(len(headers)): box((x[i],y,x[i+1],y+row_h),fill="#ffffff")
        wrapped_txt((x[0]+x[1])/2,y+row_h/2,s["name"],16,14,2)
        img_path=product_image_file(s.get("image_path"))
        if img_path:
            source=Image.open(img_path).convert("RGBA"); source.thumbnail((110*S,58*S),Image.Resampling.LANCZOS)
            thumb=Image.new("RGBA",(110*S,58*S),"white"); thumb.alpha_composite(source,((thumb.width-source.width)//2,(thumb.height-source.height)//2)); im.paste(thumb.convert("RGB"),(round((x[1]+10)*S),round((y+6)*S)))
        wrapped_txt((x[2]+x[3])/2,y+row_h/2,s.get("model") or "-",15,12,2)
        wrapped_txt((x[3]+x[4])/2,y+row_h/2,s.get("params") or "-",13,22,3)
        wrapped_txt((x[4]+x[5])/2,y+row_h/2,s.get("brand_name") or "-",15,9,2)
        wrapped_txt((x[5]+x[6])/2,y+row_h/2,s.get("origin") or "-",15,7,2)
        txt((x[6]+x[7])/2,y+row_h/2,item["quantity"],17)
        txt((x[7]+x[8])/2,y+row_h/2,"¥"+item["actual_price"],16)
        txt((x[8]+x[9])/2,y+row_h/2,s.get("unit") or "个",17)
        txt((x[9]+x[10])/2,y+row_h/2,"¥"+item["subtotal"],16)
    bottom=top+header_h+max(1,len(items))*row_h
    if not items:
        for i in range(len(headers)): box((x[i],top+header_h,x[i+1],bottom),fill="#ffffff")
        txt(960,top+header_h+row_h/2,"暂无已配置产品",18,"#777777")
    box((132,bottom,1788,bottom+36),fill="#000000")
    if show_totals:
        box((132,bottom+36,1695,bottom+91),fill="#ffffff"); box((1695,bottom+36,1788,bottom+91),fill="#ffffff"); txt(913.5,bottom+63.5,"本类合计",18); txt(1741.5,bottom+63.5,"¥"+category_total,18)
        fee_y=bottom+91
        if show_grand and data["fees"]:
            fee_names="、".join(f["name"] for f in data["fees"]); fee_total=sum(Decimal(f["amount"]) for f in data["fees"])
            box((132,fee_y,1695,fee_y+55),fill="#ffffff"); box((1695,fee_y,1788,fee_y+55),fill="#ffffff"); txt(913.5,fee_y+27.5,"（服务）"+fee_names[:55],16); txt(1741.5,fee_y+27.5,f"¥{fee_total:.2f}",17); fee_y+=55
        if show_grand:
            box((132,fee_y,1695,fee_y+55),fill="#ffffff"); box((1695,fee_y,1788,fee_y+55),fill="#ffffff"); txt(913.5,fee_y+27.5,"总计",19); txt(1741.5,fee_y+27.5,"¥"+data["total"],19)
    txt(132,1035,"导出时间 "+now()[:19].replace("T"," "),14,"#777b84",anchor="lm")
    return im


def quote_summary_page(data,groups,page_no,total_pages):
    S=2; im=Image.new("RGB",(1920*S,1080*S),"white"); d=ImageDraw.Draw(im)
    def box(c,fill=None,width=2): d.rectangle(tuple(round(v*S) for v in c),fill=fill,outline="#000000",width=max(1,round(width*S)))
    def txt(x,y,v,size=22,color="#111111",anchor="mm"): d.text((x*S,y*S),str(v),font=font(size*S),fill=color,anchor=anchor)
    # Match the full-width detail table on the preceding pages.
    box((132,115,1788,220),fill="#000000"); txt(960,168,"项目汇总",34,"#ffffff")
    txt(134,92,data["project"]["name"],18,"#60646d",anchor="lm"); txt(1786,92,f"第 {page_no} / {total_pages} 页",16,"#777b84",anchor="rm")
    y=285; header_h=64; table_bottom=930
    box((132,y,960,y+header_h),fill="#000000"); box((960,y,1788,y+header_h),fill="#000000"); txt(546,y+header_h/2,"项目类别",23,"#ffffff"); txt(1374,y+header_h/2,"项目合计",23,"#ffffff"); y+=header_h
    row_count=len(groups)+(1 if data["fees"] else 0)+1
    row_h=max(58,(table_bottom-y)//max(1,row_count))
    for group in groups:
        box((132,y,960,y+row_h),fill="#ffffff"); box((960,y,1788,y+row_h),fill="#ffffff"); txt(546,y+row_h/2,group["label"],24); txt(1374,y+row_h/2,"¥"+group["subtotal"],24); y+=row_h
    if data["fees"]:
        fee_total=sum(Decimal(f["amount"]) for f in data["fees"]); box((132,y,960,y+row_h),fill="#ffffff"); box((960,y,1788,y+row_h),fill="#ffffff"); txt(546,y+row_h/2,"服务及其他费用",24); txt(1374,y+row_h/2,f"¥{fee_total:.2f}",24); y+=row_h
    box((132,y,960,y+row_h),fill="#ffffff"); box((960,y,1788,y+row_h),fill="#ffffff"); txt(546,y+row_h/2,"合计",25); txt(1374,y+row_h/2,"¥"+data["total"],25)
    txt(132,1035,"客户："+(data["project"].get("customer") or "-")+"    地址："+(data["project"].get("address") or "-"),14,"#6d717a",anchor="lm")
    return im


@app.get("/api/projects/{pid}/quote/export.pdf")
def export_quote_pdf(pid: str):
    with connect() as db: data=quote_data(db,pid)
    groups=quote_groups(data); pages=[]; chunks=[]; per_page=7
    if not groups: groups=[{"category":"智能","label":"产品","items":[],"subtotal":"0.00"}]
    for group in groups:
        parts=[group["items"][i:i+per_page] for i in range(0,len(group["items"]),per_page)] or [[]]
        for part_index,part in enumerate(parts): chunks.append((group,part,part_index,len(parts)))
    total_pages=len(chunks)+1
    for page_index,(group,part,part_index,part_count) in enumerate(chunks,1):
        suffix=f"（{part_index+1}/{part_count}）" if part_count>1 else ""
        title=group["label"]+"报价清单"+suffix
        is_group_last=part_index==part_count-1; is_final=page_index==len(chunks)
        pages.append(quote_page(data,title,part,f"第 {page_index} / {total_pages} 页",show_totals=is_group_last,show_grand=is_final,category_total=group["subtotal"]))
    pages.append(quote_summary_page(data,quote_groups(data),total_pages,total_pages))
    path=EXPORTS/f'{safe_name(data["project"]["name"])}_客户报价_{datetime.now():%Y%m%d_%H%M%S}.pdf'
    pages[0].save(path,"PDF",resolution=200,save_all=True,append_images=pages[1:])
    return FileResponse(path,filename=path.name,headers={"Cache-Control":"no-store, max-age=0"})


@app.get("/api/projects/{pid}/quotation-v2/export.pdf")
def export_quotation_v2_pdf(pid: str):
    with connect() as db:
        quotation=quotation_v2_data(db,pid)
    frontend_port=os.environ.get("POINT_QUOTE_WEB_PORT","5173")
    url=f"http://127.0.0.1:{frontend_port}/print/quotation/{pid}"
    filename=f'{safe_name(quotation["settings"]["project_name"])}_智能家居报价_V{quotation["version"]}_{datetime.now():%Y-%m-%d}.pdf'
    path=EXPORTS/filename
    try:
        result=subprocess.run(["node",str(ROOT/"backend"/"render_quotation_pdf.mjs"),url,str(path)],cwd=ROOT,capture_output=True,text=True,timeout=180)
    except (OSError,subprocess.TimeoutExpired) as error:
        raise HTTPException(500,f"PDF 打印引擎启动失败：{error}")
    if result.returncode or not path.is_file():
        detail=(result.stderr or result.stdout or "未知错误").strip()[-1200:]
        raise HTTPException(500,f"PDF 生成失败，请确认网页服务和 Chromium 已安装：{detail}")
    return FileResponse(path,filename=filename,media_type="application/pdf",headers={"Cache-Control":"no-store, max-age=0"})
